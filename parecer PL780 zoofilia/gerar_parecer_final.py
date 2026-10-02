# -*- coding: utf-8 -*-
"""Parecer DGAV sobre o PL n.º 780/XVII/2.ª — versão final, texto do utilizador."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.abspath(__file__))
PRETO = RGBColor(0, 0, 0)

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin = sec.right_margin = Cm(2.5)

est = doc.styles['Normal']
est.font.name = 'Calibri'
est.font.size = Pt(11)
est.font.color.rgb = PRETO


def p(txt, tamanho=11, antes=0, depois=8, alinhar=WD_ALIGN_PARAGRAPH.JUSTIFY):
    par = doc.add_paragraph()
    par.alignment = alinhar
    pf = par.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    pf.line_spacing = 1.15
    r = par.add_run(txt)
    r.font.size = Pt(tamanho)
    r.font.color.rgb = PRETO
    return par


def titulo(txt):
    p(txt, tamanho=11, antes=14, depois=6, alinhar=WD_ALIGN_PARAGRAPH.LEFT)


p('DIREÇÃO-GERAL DE ALIMENTAÇÃO E VETERINÁRIA', depois=2, alinhar=WD_ALIGN_PARAGRAPH.LEFT)
p('Bem-estar dos animais de companhia', depois=2, alinhar=WD_ALIGN_PARAGRAPH.LEFT)
p('Processo n.º ______   Data ______', depois=16, alinhar=WD_ALIGN_PARAGRAPH.LEFT)

p('1. Projeto de Lei que criminaliza a prática da zoofilia reforçando a proteção animal e a tutela '
  'penal e contraordenacional.')
p('2. Altera o artigo 387.º do Código Penal, e adita a alínea h) ao n.º 3 do artigo 1.º da Lei '
  'n.º 92/95, de 12 de setembro.')

titulo('Quanto ao Artigo 387.º do Código Penal:')
p('3. O n.º 3 do artigo 387.º passa a ter duas alíneas. A alínea a) reproduz o crime de maus-tratos '
  'vigente, e a alínea b) acrescenta «Ofender sexualmente um animal de companhia através de cópula, '
  'coito anal, coito oral ou a introdução vaginal, anal ou oral de partes do corpo ou de objetos».')
p('4. A nova alínea está sujeita ao preâmbulo do artigo que exceciona motivos legítimos - por exemplo, '
  'atos médico-veterinários -, permitindo que a ofensa sexual possa ser punível sem prova de dor, '
  'sofrimento ou lesão.')
p('5. A moldura penal mantém-se.')

titulo('Quanto ao artigo 1.º da Lei n.º 92/95:')
p('6. A Lei 92/95 acolhe a mesma descrição na alínea h) do n.º 3 do artigo 1.º, passando a conduta a '
  'constituir também contraordenação, aplicável a qualquer animal.')
p('7. Juridicamente não se afasta a dúvida sobre se o artigo 3.º acolhe as exclusões do art.º 1.º '
  'quanto a «violências injustificadas», mas presume-se que as causas de exclusão da ilicitude se '
  'aplicam por via do Regime Geral das Contra-Ordenações.')

titulo('Parecer')
p('8. No presente, o n.º 3 do artigo 7.º do Decreto-Lei n.º 276/2001, de 17 de outubro, estabelece que '
  'são proibidas «todas as violências contra animais, considerando-se como tais os atos consistentes '
  'em, sem necessidade, se infligir a morte, o sofrimento ou lesões a um animal», norma punida como '
  'contraordenação económica muito grave, pela alínea d) do n.º 2 do artigo 68.º.')
p('9. O artigo 6.º, que fixa o dever especial de cuidado, só é sancionado por remissão direta quando o '
  'perigo recaia sobre outro animal, pela alínea j) do n.º 1, ou sobre outrem, pela alínea b) do n.º 2.')
p('10. No aspeto penal, o n.º 3 do artigo 387.º pune apenas quem inflija dor, sofrimento ou maus-tratos '
  'físicos a animal de companhia.')
p('11. Assim, o ato sexual do qual não resulte sofrimento ou lesão demonstráveis não se encontra '
  'sancionável. Esta lacuna pode ser coberta pela Proposta, como a própria exposição de motivos '
  'identifica.')
p('12. Resulta relevante para a premência desta alteração o impacto destes comportamentos no bem-estar '
  'animal, bem como, os riscos higiossanitários e zoonóticos associados. Está documentada transmissão '
  'por esta via de agentes de doenças como a brucelose. Há também casuística de lesão em animais por '
  'esta via, em número não residual. Num levantamento de 448 casos de lesão não acidental publicado em '
  'revista de clínica veterinária, 6% eram de natureza sexual, parte das quais sem dano externo '
  'evidente.')
p('13. Em conclusão, as ofensas sexuais de que não resulte sofrimento nem lesão persistentes e '
  'demonstráveis não têm hoje sanção penal nem contraordenacional. Se a formulação jurídica for '
  'validada, não se encontra objeções técnicas, pelo que o parecer é favorável.')

out = os.path.join(BASE, 'Parecer_DGAV_PL780_XVII_final.docx')
doc.save(out)
print('gerado:', os.path.basename(out))
