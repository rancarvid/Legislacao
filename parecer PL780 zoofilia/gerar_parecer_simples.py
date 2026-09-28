# -*- coding: utf-8 -*-
"""Parecer DGAV sobre o PL n.º 780/XVII/2.ª — texto simples, sem realces."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.abspath(__file__))
PRETO = RGBColor(0, 0, 0)

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin = sec.right_margin = Cm(2.5)

est = doc.styles['Normal']
est.font.name = 'Calibri'
est.font.size = Pt(11)
est.font.color.rgb = PRETO


def p(txt, tamanho=11, antes=0, depois=8, alinhar=WD_ALIGN_PARAGRAPH.JUSTIFY):
    par = doc.add_paragraph()
    par.alignment = alinhar
    pf = par.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    pf.line_spacing = 1.15
    r = par.add_run(txt)
    r.font.size = Pt(tamanho)
    r.font.color.rgb = PRETO
    return par


def titulo(txt):
    p(txt, tamanho=12, antes=16, depois=6, alinhar=WD_ALIGN_PARAGRAPH.LEFT)


p('DIREÇÃO-GERAL DE ALIMENTAÇÃO E VETERINÁRIA', depois=2, alinhar=WD_ALIGN_PARAGRAPH.LEFT)
p('Bem-estar dos animais de companhia', depois=2, alinhar=WD_ALIGN_PARAGRAPH.LEFT)
p('Processo n.º ______   Data ______', depois=14, alinhar=WD_ALIGN_PARAGRAPH.LEFT)

p('Parecer sobre o Projeto de Lei n.º 780/XVII/2.ª (PAN), que criminaliza a prática da zoofilia '
  'reforçando a proteção animal e a tutela penal e contraordenacional', tamanho=12, depois=12,
  alinhar=WD_ALIGN_PARAGRAPH.LEFT)

p('O projeto altera o artigo 387.º do Código Penal e adita a alínea h) ao n.º 3 do artigo 1.º da Lei '
  'n.º 92/95, de 12 de setembro. Este parecer incide sobre o bem-estar dos animais de companhia. As '
  'observações sobre as restantes matérias são sinalizações, sem prejuízo da pronúncia das unidades '
  'competentes.')
p('Sentido: favorável, condicionado à correção referida no ponto 3.')

titulo('1. O que a iniciativa acrescenta')
p('O n.º 3 do artigo 387.º passa a ter duas alíneas, ambas sob o elemento «sem motivo legítimo». A '
  'alínea a) reproduz o crime de maus tratos vigente. A alínea b) acrescenta «Ofender sexualmente um '
  'animal de companhia através de cópula, coito anal, coito oral ou a introdução vaginal, anal ou oral '
  'de partes do corpo ou de objetos». A moldura penal mantém-se.')
p('Na Lei n.º 92/95 adita-se a alínea h), com fórmula equivalente mas referida a «animais», sem '
  'qualificação de espécie.')
p('O efeito útil é um só. A conduta passa a ser punível sem prova de dor, sofrimento ou lesão. É o que '
  'a legislação vigente não permite.')

titulo('2. A insuficiência que a iniciativa supre')
p('Toda a tutela hoje disponível é de resultado. O n.º 3 do artigo 387.º exige dor, sofrimento ou maus '
  'tratos físicos provados. O n.º 1 do artigo 1.º da Lei n.º 92/95 exige a morte, o sofrimento cruel e '
  'prolongado ou graves lesões.')
p('Há um ponto que releva diretamente para esta divisão. O artigo 6.º do Decreto-Lei n.º 276/2001, de '
  '17 de outubro, impõe ao detentor o dever de cuidar do animal «de forma a não pôr em causa os '
  'parâmetros de bem-estar». Mas as contraordenações que sancionam esse dever são apenas duas. A alínea '
  'j) do n.º 1 do artigo 68.º pune a violação que crie perigo para outro animal. A alínea b) do n.º 2 '
  'pune a que crie perigo para outrem. Violar o dever de cuidado pondo em perigo o próprio animal não é '
  'contraordenação. A ofensa sexual situa-se nesse espaço.')
p('Acresce que a alínea c) do n.º 5 do artigo 387.º já qualifica como especialmente censurável o crime '
  'determinado «para excitação». A agravante existe. Falta o tipo que ela pressupõe.')
p('Sinalização fora da competência desta divisão. A conduta envolve contacto direto de mucosas entre '
  'espécies e tem relevância higiossanitária e zoonótica documentada. Um estudo caso-controlo '
  'multicêntrico apurou associação epidemiológica (Zequi SC, et al., J Sex Med. 2012;9(7):1860-67). '
  'Está descrita transmissão sexual de brucelose (Li N, et al., IDCases. 2020;21:e00871). É matéria de '
  'saúde pública veterinária. Sugere-se que seja levada à fundamentação da iniciativa.')

titulo('3. Correção premente: sobreinclusão sobre atos médico-veterinários')
p('A alínea h) proposta abrange «a introdução vaginal, anal ou oral de partes do corpo ou de objetos». '
  'A norma penal subordina ambas as alíneas ao elemento «sem motivo legítimo». A norma '
  'contraordenacional não tem elemento limitador nenhum.')
p('O n.º 1 do mesmo artigo não resolve a falta. É autodefinido e o n.º 3 abre com «São também '
  'proibidos», acrescentando e não especificando. A demonstração está no próprio artigo: as alíneas a), '
  'b), c), e) e f) do n.º 3 têm limitadores próprios, que seriam redundantes se o n.º 1 se transmitisse.')
p('Ficam abrangidos, na letra, atos correntes da prática clínica e zootécnica: termometria retal, '
  'palpação e ecografia transretais, inseminação artificial, sondagem, enemas, exploração obstétrica.')
p('Convém delimitar o alcance da objeção. O artigo 32.º do Regime Geral das Contra-Ordenações manda '
  'aplicar as normas do Código Penal, incluindo as causas de exclusão da ilicitude do artigo 31.º. O ato '
  'praticado no exercício da atividade médico-veterinária estaria justificado. Não se sustenta que um '
  'médico veterinário viesse a ser sancionado.')
p('A objeção é de estrutura. A norma fica sobreinclusiva na sua face e a licitude passa a discutir-se '
  'por causa de justificação não escrita. O artigo 11.º da Lei n.º 92/95 atribui a fiscalização a nove '
  'entidades, incluindo autoridades policiais.')
p('Propõe-se introduzir na alínea h) um elemento que restrinja a conduta à finalidade sexual. Em '
  'alternativa, ressalva expressa dos atos médico-veterinários e zootécnicos legalmente praticados. A '
  'primeira solução é preferível. Com elemento finalístico o ato clínico não chega a preencher a norma, '
  'o que dispensa a ressalva e a sua atualização.')

titulo('4. Observações de aperfeiçoamento')
p('Taxatividade. A lista é fechada. Ficam fora a masturbação do animal, o contacto oral-genital sem '
  'coito e a imposição de monta, incluindo quando a desproporção de porte possa causar lesões. A '
  'exposição de motivos refere condutas que o articulado não abrange.')
p('Conteúdos audiovisuais. A exposição de motivos anuncia a criminalização da produção e difusão de '
  'conteúdos. O articulado nada dispõe. A alínea e) do n.º 3 do artigo 1.º já cobre parte da produção. '
  'Nada cobre a difusão.')
p('Dever de comunicação. Deverá prever-se dever de comunicação a cargo do médico veterinário que detete '
  'indícios compatíveis com ofensa sexual? E como se articula esse dever com o sigilo do Código '
  'Deontológico Médico-Veterinário? Sugere-se a audição da Ordem dos Médicos Veterinários.')

titulo('5. Conclusão')
p('A insuficiência invocada verifica-se. A tutela vigente é toda de resultado e a violação do dever de '
  'cuidado que ponha em perigo o próprio animal não é contraordenável. A iniciativa tem utilidade '
  'própria por deslocar a tutela do resultado para o ato.')
p('A redação proposta para a Lei n.º 92/95 é sobreinclusiva quanto a atos médico-veterinários e '
  'zootécnicos. A correção do ponto 3 é premente. Nada mais se opõe, em termos técnicos, à aprovação da '
  'iniciativa.')

out = os.path.join(BASE, 'Parecer_DGAV_PL780_XVII_simples.docx')
doc.save(out)
print('gerado:', os.path.basename(out))
