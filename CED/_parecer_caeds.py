# -*- coding: utf-8 -*-
"""Gera o Parecer CED - CAEDS, V.2."""
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

PRETO = RGBColor(0, 0, 0)
BASE = os.path.dirname(os.path.abspath(__file__))


def doc_novo():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(3)
    s.top_margin = s.bottom_margin = Cm(2.5)
    n = d.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(11)
    n.font.color.rgb = PRETO
    n.paragraph_format.space_after = Pt(8)
    n.paragraph_format.line_spacing = 1.15
    return d


def p(d, txt, negrito=False, estilo=None, justificar=True):
    q = d.add_paragraph(style=estilo)
    if justificar and estilo is None:
        q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = q.add_run(txt)
    r.font.bold = negrito
    r.font.color.rgb = PRETO
    return q


def titulo(d, txt):
    q = d.add_paragraph()
    q.paragraph_format.space_before = Pt(14)
    r = q.add_run(txt)
    r.font.bold = True
    r.font.color.rgb = PRETO
    return q


d = doc_novo()

p(d, "Cara Maria,")
p(d, "Segue o meu parecer.")

p(d, "Em resumo: o animal inserido num programa CED é um animal errante e deve ser registado no "
     "Sistema de Informação de Animais de Companhia em nome do município. O voluntário que "
     "cuida da colónia é detentor, não titular.")

p(d, "A Lei n.º 27/2016, de 23 de agosto, fez do CED uma política pública. O seu art.º 4.º estabelece "
     "que «o Estado, por razões de saúde pública, assegura, por intermédio dos centros de recolha "
     "oficial de animais, a captura, vacinação e esterilização dos animais errantes sempre que "
     "necessário, assim como a concretização de programas captura, esterilização, devolução (CED) "
     "para gatos».")

p(d, "É o Estado que assegura o CED, e fá-lo através dos centros de recolha oficial, que são "
     "equipamentos municipais. A lei não atribui o programa a particulares nem menciona a figura do "
     "cuidador. A Portaria n.º 146/2017, de 26 de abril, que regulamenta a matéria, também não "
     "define cuidador.")

p(d, "A Portaria confirma a natureza administrativa do programa. O n.º 1 do seu art.º 9.º dispõe: "
     "«Como forma de gestão da população de gatos errantes e nos casos em que tal se justifique, "
     "podem as câmaras municipais, sob parecer do médico veterinário municipal, autorizar a "
     "manutenção, em locais especialmente designados para o efeito, de colónias de gatos, no âmbito "
     "de programas de captura, esterilização e devolução (CED) ao local de origem».")

p(d, "A colónia existe, por isso, por autorização da câmara municipal. O n.º 2 do mesmo artigo "
     "permite que a câmara atribua a gestão do programa a uma organização de proteção animal, "
     "mediante proposta desta. Atribuir a gestão não é transferir a titularidade dos animais.")

p(d, "Uma precisão de linguagem, que importa para o registo. A titularidade é do município, que é a "
     "pessoa coletiva. A câmara municipal é o órgão que exerce as competências. É por isso que o "
     "registo no SIAC se faz em nome do município, e não do órgão.")

p(d, "Os animais capturados passam pelo centro de recolha oficial antes de integrarem a colónia. "
     "A al. d) do n.º 4 do art.º 9.º da Portaria exige «Que os animais capturados, antes de "
     "integrarem a colónia, são entregues nos CRO para verificação da sua aptidão».")

p(d, "Quanto ao registo, o n.º 5 do art.º 11.º do Decreto-Lei n.º 82/2019, de 27 de junho, dispõe: "
     "«Os animais que sejam recolhidos num Centro de Recolha Oficial (CRO) e que não sejam "
     "reclamados pelos seus proprietários devem ser registados no SIAC em nome do titular desse CRO, "
     "após o período de 15 dias previsto no n.º 4 do artigo 8.º da Portaria n.º 146/2017, de 26 de "
     "abril». Esse n.º 4 não fixa o prazo: remete-o para o n.º 1 do art.º 3.º da Lei n.º 27/2016. "
     "Toda a cadeia assenta, por isso, na mesma lei que comete ao Estado, através dos centros de "
     "recolha oficial, a execução do CED.")

p(d, "A norma não foi escrita a pensar no CED, e a Portaria não diz expressamente em nome de quem se "
     "registam os gatos de colónia. Mas o critério da lei é um só: o animal errante que entra num "
     "centro de recolha oficial e não tem proprietário conhecido é registado em nome do titular desse "
     "centro. A Lei n.º 27/2016 prevê dois destinos para esse animal, a adoção, no n.º 1 do art.º 3.º, "
     "e a devolução ao local de origem, no art.º 4.º. O prazo de 15 dias distingue destinos. Não "
     "distingue titulares.")

p(d, "A cadeia que conduz a um gato inserido num programa CED é, assim, a seguinte:")
for i, item in enumerate((
        "É um animal errante, cuja captura e recolha compete às câmaras municipais, nos termos do "
        "n.º 1 do art.º 7.º da Portaria n.º 146/2017;",
        "É entregue no centro de recolha oficial antes de integrar a colónia, nos termos da al. d) "
        "do n.º 4 do art.º 9.º da mesma Portaria;",
        "Não havendo proprietário conhecido, é registado em nome do titular desse centro, que é o "
        "município."), 1):
    q = p(d, item, estilo="List Number")
    q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

p(d, "O Decreto-Lei n.º 82/2019 define, na al. f) do art.º 3.º, «Titular de animal de companhia» "
     "como «o proprietário ou o possuidor, quer se trate de pessoa singular ou coletiva, que seja "
     "responsável pelo animal de companhia, independentemente da finalidade com que o detém, e cuja "
     "posse faça presumir a propriedade e em cujo nome deve efetuar-se o registo da titularidade do "
     "animal de companhia no SIAC». E define, na al. a) do mesmo artigo, «Detentor» como a pessoa "
     "«que se encontre na situação de possuidor precário, nos termos previstos no artigo 1253.º do "
     "Código Civil». É esta a posição de quem alimenta e acompanha a colónia: tem o animal a seu "
     "cuidado, mas não o tem como seu.")

p(d, "A titularidade não exige que o titular tenha o animal consigo. O proprietário de um cão que se "
     "perde continua titular, e é por isso que a al. d) do n.º 2 do art.º 13.º do mesmo decreto-lei o "
     "obriga a comunicar o desaparecimento ao SIAC. O n.º 1 do art.º 14.º distingue o titular do "
     "«simples detentor». Nada impede, por isso, que o município seja titular de animais que vivem "
     "em liberdade.")

p(d, "Por último, o n.º 9 do art.º 9.º da Portaria permite à câmara municipal, verificado o "
     "incumprimento de qualquer dos requisitos do n.º 4, determinar medidas corretivas ou a suspensão "
     "do programa e «proceder à recolha dos animais para o CRO». Se os animais fossem dos cuidadores, "
     "a câmara estaria a retirar-lhes animais próprios, sem procedimento e sem indemnização. A norma "
     "só se compreende se os animais forem do município.")

titulo(d, "Respostas às questões colocadas")

p(d, "1. Quando os gatos de rua forem esterilizados no âmbito do Programa CED, em nome de quem "
     "deverão ser registados os respetivos microchips no SIAC?", negrito=True)
p(d, "Em nome do município.")

p(d, "2. Poderão estes animais ser registados em nome da Câmara Municipal de Silves ou da Junta de "
     "Freguesia territorialmente competente, mantendo os voluntários a responsabilidade pelo "
     "acompanhamento das colónias?", negrito=True)
p(d, "Sim. O registo faz-se em nome do Município de Silves. Não em nome da junta de freguesia. A "
     "competência para "
     "autorizar as colónias é da câmara municipal, sob parecer do médico veterinário municipal, e o "
     "animal passa pelo centro de recolha oficial, que é municipal. A junta de freguesia pode "
     "colaborar, inclusive por delegação de competências do município nos termos da Lei n.º 75/2013, "
     "de 12 de setembro, mas essa colaboração não transfere a titularidade dos animais, pela mesma "
     "razão que a atribuição da gestão a uma organização de proteção animal também não transfere.")
p(d, "Os voluntários mantêm o acompanhamento das colónias, na qualidade de detentores, e devem "
     "constar do plano de gestão da colónia, nos termos da al. a) do n.º 4 do art.º 9.º da Portaria "
     "n.º 146/2017.")

p(d, "3. Existe alguma disposição legal que obrigue os voluntários a registar em seu nome pessoal os "
     "gatos de rua que levam ao veterinário para esterilização?", negrito=True)
p(d, "Não. Nenhuma disposição legal o impõe. A indicação em sentido contrário não tem suporte legal.")

p(d, "4. Qual é o procedimento legalmente adequado para permitir que as esterilizações continuem "
     "através do Programa CED, sem que os voluntários tenham de assumir individualmente a "
     "titularidade de dezenas ou centenas de animais que vivem em colónias?", negrito=True)
p(d, "Requerer à Câmara Municipal de Silves a autorização das colónias ao abrigo do n.º 1 do art.º 9.º "
     "da Portaria n.º 146/2017, com identificação dos locais e dos voluntários responsáveis pela "
     "execução no plano de gestão de cada colónia. Autorizado o programa, os animais são "
     "identificados, esterilizados e registados em nome do município.")
p(d, "Não é exigida a constituição de associação para que os voluntários colaborem no programa. A "
     "constituição de associação e o protocolo previsto no n.º 2 do art.º 9.º só são necessários se "
     "pretenderem assumir a gestão do programa.")

titulo(d, "Nota sobre a norma citada no pedido")

p(d, "A norma invocada pela requerente existe, mas consta do regime aplicável na Região Autónoma dos "
     "Açores: n.º 7 do art.º 6.º e n.º 13 do art.º 6.º-B do Decreto Legislativo Regional n.º 12/2016/A, "
     "de 8 de julho, na redação do Decreto Legislativo Regional n.º 13/2023/A, de 14 de abril. Ali "
     "determina-se que os gatos do Programa CED que estejam sob responsabilidade de associações de "
     "proteção animal «são registados em nome do município com jurisdição territorial». Não é "
     "aplicável no continente, mas vai no sentido do entendimento acima.")

out = os.path.join(BASE, "Parecer_CED_CAEDS_V2.docx")
d.save(out)
print("gerado:", os.path.basename(out))
