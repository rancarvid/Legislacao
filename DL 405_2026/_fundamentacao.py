# -*- coding: utf-8 -*-
"""Fundamentação detalhada das observações ao DL das atividades económicas (DL 405/XXV/2026)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

AZUL = RGBColor(0x1F, 0x3B, 0x63)
CINZA = RGBColor(0x55, 0x55, 0x55)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = s.bottom_margin = Cm(2.2)

st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def p(txt='', bold=False, italic=False, size=10.5, cor=None, space=6, left=0):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(space)
    if left:
        par.paragraph_format.left_indent = Cm(left)
    r = par.add_run(txt)
    r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if cor is not None:
        r.font.color.rgb = cor
    return par


def h1(txt):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(18); par.paragraph_format.space_after = Pt(6)
    r = par.add_run(txt); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = AZUL


def h2(txt):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(12); par.paragraph_format.space_after = Pt(4)
    r = par.add_run(txt); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = AZUL


def cit(txt, fonte):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.8)
    par.paragraph_format.right_indent = Cm(0.4)
    par.paragraph_format.space_after = Pt(2)
    r = par.add_run(txt); r.italic = True; r.font.size = Pt(9.5)
    f = doc.add_paragraph()
    f.paragraph_format.left_indent = Cm(0.8)
    f.paragraph_format.space_after = Pt(8)
    rf = f.add_run(fonte); rf.font.size = Pt(8.5); rf.font.color.rgb = CINZA


def bullet(txt):
    par = doc.add_paragraph(style='List Bullet')
    par.paragraph_format.space_after = Pt(3)
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = par.add_run(txt); r.font.size = Pt(10.5)


def tabela(cab, linhas, larguras=None):
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cab):
        cell = t.rows[0].cells[i]; cell.text = ''
        par = cell.paragraphs[0]; par.paragraph_format.space_after = Pt(2)
        r = par.add_run(c); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = AZUL
    for linha in linhas:
        cells = t.add_row().cells
        for i, c in enumerate(linha):
            cells[i].text = ''
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = par.add_run(c); r.font.size = Pt(9)
    if larguras:
        for row in t.rows:
            for i, w in enumerate(larguras):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ───────────────────────────── Capa ─────────────────────────────
p('DIREÇÃO-GERAL DE ALIMENTAÇÃO E VETERINÁRIA', bold=True, size=9.5, cor=CINZA, space=0)
p('Direção de Serviços de Bem-Estar Animal', size=9.5, cor=CINZA, space=14)

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(2)
r = par.add_run('FUNDAMENTAÇÃO DAS OBSERVAÇÕES'); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = AZUL

p('Projeto de DL das atividades económicas — DL 405/XXV/2026', bold=True, size=12, space=2)
p('Documento de apoio interno. Acompanha a série de observações enviadas ponto a ponto.',
  size=10.5, cor=CINZA, space=14)

h1('Nota de utilização')
p('Cada ponto da série de observações é enviado em texto curto, com um máximo de palavras fixado, para '
  'seguir na resposta à Direção de Serviços. Esse texto é deliberadamente contido: situa a norma, articula com '
  'a legislação vigente e com o Regulamento (UE) 2026/1818 e deixa a decisão à consideração de quem decide.')
p('Este documento guarda o que fica de fora: a demonstração completa, as citações verbatim, a análise das '
  'alterações registadas no ficheiro, as inferências assumidas como tais e as fontes com referência exata. '
  'Serve para sustentar a posição se ela for questionada, e para que outra pessoa possa retomar o trabalho sem '
  'repetir a verificação.')
p('É atualizado à medida que cada ponto é trabalhado. O registo de alterações está no fim.')

h1('Índice da série e numeração')
p('A numeração é a da série de observações, e não a do diploma. Foi fixada à medida que os pontos foram '
  'sendo trabalhados; os pontos 01 e 02 estão fechados, os seguintes estão previstos e a ordem pode mudar.')

tabela(
    ['N.º', 'Tema', 'Norma principal', 'Estado'],
    [
        ['01', 'Estabelecimentos de comércio a retalho de animais de companhia: ausência da DGAV no '
                'procedimento',
         'Art.º 16.º, n.º 4, e Anexo II', 'Fechado'],
        ['02', 'Animais de companhia na definição de atividade pecuária: a exclusão eliminada',
         'Art.º 48.º, als. b) e c)', 'Fechado'],
        ['03', 'Deferimento tácito em atividades com animais vivos', 'Art.º 25.º', 'Previsto'],
        ['04', 'Cláusula de articulação com o regime dos animais de companhia e com o Regulamento',
         'Art.ºs 2.º e 36.º, n.º 3', 'Previsto'],
        ['05', 'Recolha de animais errantes: delimitação às espécies pecuárias',
         'Art.ºs 57.º e 59.º', 'Previsto'],
        ['06', 'Venda em comércio não sedentário, feiras e leilões: falta de articulação',
         'Art.ºs 133.º, n.º 2, e 153.º, n.º 1', 'Previsto'],
        ['07', 'Venda e publicidade em linha', 'Art.º 1.º, n.º 4', 'Previsto'],
        ['08', 'Acesso de animais de companhia a estabelecimentos', 'Art.º 38.º', 'Previsto'],
        ['09', 'SIAC, base de dados setorial e cadastro', 'Art.ºs 198.º a 200.º', 'Previsto'],
        ['10', 'Lapsos formais com efeito material', 'Anexo V; numeração da Secção III', 'Previsto'],
    ],
    larguras=[1.0, 6.2, 4.3, 2.0],
)

p('Nota sobre a renumeração: a ordem seguida na primeira análise global, registada no contributo interno de '
  '5 de outubro de 2026, era diferente. O tema que ali figurava como ponto 6, sobre a fronteira entre detenção '
  'caseira e animal de companhia, passou a ser o ponto 02 desta série e absorveu aquele. Os restantes desceram '
  'uma posição.', size=9.5, cor=CINZA)

# ───────────────────────────── Ponto 01 ─────────────────────────────
h1('Ponto 01 — Estabelecimentos de comércio a retalho de animais de companhia')

h2('1.1. A norma no projeto')
p('A atividade está prevista no n.º 4 do artigo 16.º do anexo, entre as sujeitas a comunicação prévia, com a '
  'designação «A exploração de estabelecimentos de comércio a retalho de animais de companhia e respetivos '
  'alimentos, em estabelecimentos especializados». Corresponde ao CAE 47762, «Comércio a retalho de animais de '
  'companhia e respetivos alimentos», constante da lista II do anexo I.')
p('Advertência quanto à referência: as alíneas do n.º 4 do artigo 16.º estão com a numeração automática '
  'danificada no ficheiro, aparecendo quase todas como «a)». O Anexo II remete para este caso como «alínea f)», '
  'mas essa letra não é verificável no articulado tal como está. Por isso, nas observações enviadas, a norma é '
  'identificada apenas pelo número, sem a alínea.')

h2('1.2. O que o Anexo II estabelece')
p('A linha correspondente do Anexo II, que identifica as entidades intervenientes em cada procedimento, tem o '
  'seguinte conteúdo:')
tabela(
    ['Atividade económica', 'Entidade coordenadora', 'Entidades públicas consultadas', 'Entidades notificadas'],
    [['Exploração de estabelecimentos de comércio a retalho de animais de companhia e respetivos alimentos, '
      'em estabelecimentos especializados (cf. alínea f)',
      'A câmara municipal territorialmente competente', '(em branco)', 'A DGDCCS']],
    larguras=[5.6, 3.3, 3.3, 2.4],
)
p('A comparação com as linhas vizinhas é o que torna o ponto demonstrável. Para a exploração de '
  'estabelecimentos de comércio por grosso e a retalho e de armazéns de alimentos para animais de criação, '
  'constante da al. f) do n.º 2 do mesmo artigo 16.º, o Anexo II indica a câmara municipal como coordenadora e '
  '«A DGAV» como entidade pública consultada. Para a atividade pecuária de classe 1 e de classe 2, a DGAV '
  'figura igualmente entre as consultadas. A DGAV é, portanto, chamada a pronunciar-se sobre a ração e sobre o '
  'efetivo pecuário, e não sobre o estabelecimento que detém animais de companhia vivos para venda.')

h2('1.3. Por que razão a atividade passa a estar aqui')
p('O n.º 1 do artigo 3.º do Decreto-Lei n.º 276/2001, de 17 de outubro, abre com uma ressalva:')
cit('«1 - Sem prejuízo do disposto no Decreto-Lei n.º 10/2015, de 16 de janeiro, quanto aos estabelecimentos '
    'de comércio a retalho de animais de companhia, o exercício da atividade de exploração de alojamentos, bem '
    'como a atividade de criação comercial de animais de companhia depende de: [...]»',
    'n.º 1 do art.º 3.º do Decreto-Lei n.º 276/2001, de 17 de outubro, na redação consolidada.')
p('A al. j) do n.º 1 do artigo 8.º do decreto-lei preambular revoga o Decreto-Lei n.º 10/2015, de 16 de '
  'janeiro. O n.º 4 do mesmo artigo 8.º determina que «As remissões, legais e regulamentares, para normas ora '
  'revogadas consideram-se feitas, com as devidas adaptações, para o presente decreto-lei, salvo se a '
  'interpretação daquelas impuser solução diferente». A ressalva do Decreto-Lei n.º 276/2001 passa, assim, a '
  'apontar para o novo diploma.')
p('Importa delimitar o efeito: o que transita é o título de acesso à atividade, não as normas materiais. Os '
  'requisitos de detenção dos animais à venda continuam no Decreto-Lei n.º 276/2001, que não é revogado — '
  'designadamente o seu Capítulo III, «Normas para os alojamentos de reprodução, criação, manutenção e venda de '
  'animais de companhia», com o art.º 26.º para pequenos roedores e coelhos, o art.º 27.º e o Anexo III para '
  'cães e gatos, o art.º 32.º quanto às instalações para venda, os art.ºs 36.º e 37.º quanto a animais doentes, '
  'fêmeas prenhes e ninhadas, e o art.º 38.º quanto a pessoal auxiliar e assistência médico-veterinária. O '
  'problema não é de normas materiais: é de quem fica a saber que o estabelecimento abriu.')

h2('1.4. A consequência procedimental')
p('O procedimento de comunicação prévia não tem fase de consulta. O n.º 3 do artigo 21.º do anexo descreve-o '
  'em três fases, a apresentação da comunicação prévia na Plataforma dos Licenciamentos, a decisão e a emissão '
  'do título. E o n.º 3 do artigo 29.º dispõe que, nesse procedimento, «o título digital de atividade económica '
  'é automaticamente emitido com o pagamento da taxa devida ou, não sendo devida taxa, com a apresentação da '
  'comunicação prévia».')
p('Daqui decorre a conclusão técnica que sustenta a proposta: a inclusão da DGAV entre as entidades '
  'consultadas é incompatível com a manutenção da atividade no n.º 4 do artigo 16.º, porque não existe aí fase '
  'em que uma entidade consultada possa pronunciar-se. Para haver pronúncia prévia, a atividade tem de passar '
  'ao n.º 3, que é o procedimento de comunicação com prazo. O custo dessa passagem é o prazo de pronúncia, que '
  'no Anexo III corresponde à linha residual «Outras consultas — 20 dias», com deferimento tácito favorável ao '
  'requerente nos termos do n.º 3 do artigo 25.º.')
p('Existe, dentro do próprio diploma, uma formulação aproveitável para dar peso à pronúncia:')
cit('«2 - Sem prejuízo da competência da entidade coordenadora, são vinculativas as pronúncias da DGAV em '
    'matéria de saúde e bem-estar animal, biossegurança, aprovação sanitária de estabelecimentos e validação do '
    'plano de produção»',
    'n.º 2 do art.º 54.º do anexo (condições gerais para o exercício da atividade pecuária), na versão de '
    '25.9.2026.')

h2('1.5. O que o Regulamento (UE) 2026/1818 passa a impor')
p('A loja de animais é, para o Regulamento, um estabelecimento de venda:')
cit('«(q) ‘selling establishment’ means any premises or structure where dogs or cats are kept for sale without '
    'being born there, including pet shops or private homes, as well as any premises or structures for assembly '
    'operations in which dogs or cats from more than one establishment are assembled;»',
    'al. q) do art.º 4.º do Regulamento (UE) 2026/1818 — versão inglesa, JO L, 2026/1818, de 10.8.2026.')
cit('«q) «Estabelecimento de venda», qualquer instalação ou estrutura onde cães ou gatos são detidos para '
    'venda sem aí terem nascido, incluindo lojas de animais de companhia ou casas particulares, bem como '
    'quaisquer instalações ou estruturas destinadas a operações de agrupamento nas quais sejam agrupados cães '
    'ou gatos a partir de mais do que um estabelecimento;»',
    'al. q) do art.º 4.º do Regulamento (UE) 2026/1818 — versão portuguesa, JO L, 2026/1818, de 10.8.2026.')
p('E sobre esse estabelecimento recai uma obrigação de notificação, com um registo a cargo da autoridade '
  'competente:')
cit('«1. Operators shall notify the competent authorities of their activity, providing at least the following '
    'information for each of their establishments: (a) the name, address and contact details of the operator; '
    '(b) the location of the establishment; (c) the type of establishment: breeding establishment, selling '
    'establishment, shelter or foster home; (d) the species and, for breeding establishments, the breeds of the '
    'dogs or cats kept in the establishment; (e) the capacity of the establishment, expressed as the maximum '
    'number of dogs and cats which can be kept in the establishment; (f) for breeding establishments, the '
    'estimated number of litters to be placed on the market per year. [...] 4. The competent authority shall '
    'keep a register of establishments. The competent authority may, for that purpose, use the register '
    'established pursuant to Article 101(1), point (a), of Regulation (EU) 2016/429.»',
    'art.º 9.º, n.ºs 1 e 4, do Regulamento (UE) 2026/1818 — versão inglesa.')
cit('«1. Os operadores notificam as autoridades competentes da sua atividade, facultando pelo menos as '
    'seguintes informações para cada um dos seus estabelecimentos: a) O nome, o endereço e dados de contacto do '
    'operador; b) A localização do estabelecimento; c) O tipo de estabelecimento: estabelecimento de criação, '
    'estabelecimento de venda, abrigo ou lar de acolhimento; d) A espécie e, para os estabelecimentos de '
    'criação, as raças dos cães ou gatos detidos no estabelecimento; e) A capacidade do estabelecimento, '
    'expressa em termos de número máximo de cães e gatos que podem ser detidos no estabelecimento; f) No caso '
    'dos estabelecimentos de criação, o número estimado de ninhadas a colocar no mercado por ano. [...] 4. A '
    'autoridade competente deve manter um registo de estabelecimentos. A autoridade competente pode utilizar, '
    'para esse efeito, o registo criado nos termos do artigo 101.º, n.º 1, alínea a), do Regulamento (UE) '
    '2016/429.»',
    'art.º 9.º, n.ºs 1 e 4, do Regulamento (UE) 2026/1818 — versão portuguesa.')
p('Dois elementos completam o quadro. Primeiro, não há limiares: o considerando 23 esclarece que, «atenta a '
  'natureza exclusivamente comercial dos estabelecimentos de venda, não é necessário fixar limiares», pelo que '
  'os requisitos se aplicam a todos, independentemente do número de cães ou gatos detidos. Segundo, o artigo '
  '9.º é aplicável a partir de 31 de agosto de 2028, data da regra geral do artigo 33.º, ou seja, menos de dois '
  'anos depois da entrada em vigor prevista para o diploma nacional, fixada em 4 de janeiro de 2027 pelo n.º 1 '
  'do artigo 19.º do decreto-lei preambular.')
p('Nota de delimitação: o Regulamento não exige aprovação prévia dos estabelecimentos de venda. A aprovação, '
  'com inspeção no local, está no artigo 10.º e dirige-se aos estabelecimentos de criação acima de cinco '
  'ninhadas por ano civil ou cinco fêmeas reprodutoras, sendo aplicável apenas a partir de 31 de agosto de '
  '2034. Exigir vistoria prévia às lojas seria uma regra nacional mais restritiva, admissível nos termos do '
  'n.º 1 do artigo 30.º mas sujeita à comunicação à Comissão prevista no n.º 2 do mesmo artigo. Foi por isso '
  'que a proposta enviada se ficou pela pronúncia e pela notificação.')

h2('1.6. Fontes do ponto 01')
bullet('Projeto: «DL 405 XXV 2026_PECUÁRIA_GSEA 25_09_v1.docx», anexo à mensagem de 25.9.2026 — art.ºs 16.º, '
       '21.º, 25.º, 29.º e 54.º do anexo; Anexos I, II e III; art.ºs 8.º e 19.º do decreto-lei preambular.')
bullet('Decreto-Lei n.º 276/2001, de 17 de outubro, versão consolidada — art.ºs 1.º, 2.º, 3.º, 3.º-A, 24.º a '
       '27.º, 32.º, 36.º a 38.º e Anexos II e III.')
bullet('Decreto-Lei n.º 10/2015, de 16 de janeiro — diploma revogado pela al. j) do n.º 1 do art.º 8.º do '
       'decreto-lei preambular.')
bullet('Regulamento (UE) 2026/1818 do Parlamento Europeu e do Conselho, de 17 de junho de 2026, JO L, '
       '2026/1818, de 10.8.2026, ELI http://data.europa.eu/eli/reg/2026/1818/oj — considerando 23, art.ºs 4.º, '
       '9.º, 10.º, 30.º e 33.º, nas versões inglesa e portuguesa, ambas autênticas.')

# ───────────────────────────── Ponto 02 ─────────────────────────────
h1('Ponto 02 — Animais de companhia na definição de atividade pecuária')

h2('2.1. A alteração registada no ficheiro')
p('A alínea c) do artigo 48.º do anexo define «atividade pecuária». No ficheiro de 25.9.2026 o seu início '
  'apresenta-se assim, com a alteração registada visível:')
cit('«Atividade pecuária», todas as atividades de reprodução, produção, detenção, comercialização, exposição '
    'e outras relativas a animais das espécies pecuárias, com exceção da apicultura e [eliminado: «animais de '
    'companhia»] em: [...]',
    'al. c) do art.º 48.º do anexo. Eliminação registada em nome de Nádia Silvestre, com data de 21.9.2026.')
p('O texto que hoje se lê no ficheiro, com as alterações aceites, é «com exceção da apicultura eem:». A '
  'sequência «eem» resulta da justaposição do «e» da conjunção, que antecedia a expressão eliminada, ao «em» '
  'que a seguia. É indício de que a frase não foi relida depois da eliminação.')
p('A eliminação integra-se numa revisão de grande amplitude da Secção III. Nessa mesma data e em nome da '
  'mesma pessoa foram registadas 224 inserções e 140 eliminações, que reescreveram a secção da atividade '
  'pecuária, criaram subsecções e introduziram artigos que antes não existiam, entre os quais os artigos '
  '«Detenção caseira», «Regulamentação» e «Registo de estabelecimentos não sujeitos».')

h2('2.2. A exclusão existe no regime em vigor')
p('O que foi eliminado reproduzia uma delimitação do Novo Regime do Exercício da Atividade Pecuária. A al. a) '
  'do n.º 3 do artigo 1.º do Decreto-Lei n.º 81/2013, de 14 de junho, exclui do âmbito do regime, por código '
  'de atividade económica, a apicultura, CAE 01491, e os animais de companhia, CAE 01493. O projeto comprimiu '
  'essa exclusão de âmbito para dentro de uma definição e, nessa passagem, perdeu metade dela.')
p('Verificação complementar: pesquisaram-se no projeto os códigos 01491 e 01493 e nenhum consta. O diploma '
  'não tem lista de códigos de atividade económica para a pecuária, cuja delimitação é feita por limiares de '
  'cabeças normais no artigo 16.º, pelo que a exclusão não reaparece por outra via.')

h2('2.3. A cláusula que permanece, e de onde vem')
p('A alínea b) do mesmo artigo 48.º define «animal de espécie pecuária» e mantém a referência que suscita a '
  'dúvida. Confrontada com a definição do regime em vigor, a redação é igual palavra por palavra, com uma única '
  'diferença, o acrescento dos insetos:')
cit('«c) «Animal de espécie pecuária» qualquer espécimen vivo bovino, suíno, ovino, caprino, equídeo, ave, '
    'leporídeo (coelhos e lebres) ou outra espécie que seja explorada com destino à sua reprodução ou produção '
    'de carne, leite, ovos, lã, seda, pelo, pele ou repovoamento cinegético, bem como a produção pecuária de '
    'animais destinados a animais de companhia, de trabalho ou a atividades culturais ou desportivas;»',
    'al. c) do art.º 2.º do Decreto-Lei n.º 81/2013, de 14 de junho. A redação da al. b) do art.º 48.º do '
    'projeto é idêntica, com o acrescento de «insetos» após «leporídeo (coelhos e lebres)».')
p('Quanto ao sentido da expressão, a leitura é finalística e não de espécie. O paralelo das três hipóteses, '
  'animais destinados a animais de companhia, de trabalho ou a atividades culturais ou desportivas, indica que '
  'o critério é o destino do animal e não aquilo que ele alimenta. A cláusula serve para que a produção de '
  'espécies pecuárias com destino diferente do alimentar, designadamente leporídeos e aves para o mercado dos '
  'animais de companhia, ou equídeos para trabalho ou desporto, não escape ao regime. A redação admite, porém, '
  'a leitura alternativa de animais destinados a alimentar animais de companhia, e essa ambiguidade subsiste '
  'desde 2013.')
p('A origem foi confirmada até à versão original do Decreto-Lei n.º 81/2013, publicada em 14 de junho de 2013. '
  'Não foi possível confirmar, nas fontes consultadas, se a formulação já constava do Decreto-Lei n.º 214/2008, '
  'que aquele revogou. Fica como verificação pendente.')

h2('2.4. Os comentários do ficheiro no mesmo lugar')
p('O parágrafo da definição de «animal de espécie pecuária» tem três comentários, todos sobre o acrescento dos '
  'insetos e nenhum sobre os animais de companhia:')
tabela(
    ['Marca', 'Autor', 'Teor'],
    [['PO54.1', 'Paulo Meireles de Oliveira', '«Confirma-se os insectos na pecuária?»'],
     ['GMECT54.2', 'G-MECT', '«@Pecuária»'],
     ['NS54.3', 'Nádia Silvestre', '«Confirmamos a sua inserção»']],
    larguras=[2.2, 4.5, 8.8],
)
p('O parágrafo onde a eliminação ocorreu, o da definição de «atividade pecuária», tem um único comentário, '
  'também de Nádia Silvestre, e igualmente alheio à questão: «ver com dgadr a sua proposta: Atividade pecuária '
  'em explorações pecuárias, entrepostos e centros de agrupamentos e atividades complementares de gestão de EP, '
  'anexas e autónomas». Não há, no ficheiro, qualquer justificação registada para a eliminação.')

h2('2.5. Leitura sobre o motivo (inferência)')
p('O que segue é inferência desta Direção de Serviços, e não facto documentado. Apresenta-se como tal.',
  italic=True)
p('A eliminação parece estar ligada ao artigo que foi criado na mesma revisão. A «detenção caseira» é definida '
  'como a detenção de um número reduzido de espécies pecuárias não cinegéticas em que «a posse desses animais '
  'tem o objetivo de lazer ou abastecimento do seu detentor», isenta de licenciamento mas sujeita a registo no '
  'SNIRA. Se os animais de companhia continuassem excluídos da definição de atividade pecuária, o universo que '
  'esse artigo visa, o detentor de algumas galinhas, coelhos ou de um equídeo mantidos por lazer, sairia pela '
  'porta das exceções antes de entrar pela porta do registo. A apicultura não colide com esse objetivo e por '
  'isso subsistiu.')
p('Corrobora esta leitura a alteração do preâmbulo registada no mesmo dia e em nome da mesma pessoa. O texto '
  'anterior descrevia a eliminação do controlo prévio da classe 3 e afirmava que essas atividades e a detenção '
  'caseira «continuam sujeitas» ao registo no SNIRA, traduzindo-se isso «na flexibilização do regime jurídico '
  'aplicável aos pequenos detentores, eliminando barreiras burocráticas desajustadas à dimensão da sua '
  'atividade, sem, no entanto, comprometer as exigências de rastreabilidade». O texto resultante diz que «a '
  'detenção caseira passará a estar sujeita ao registo para efeitos de Sistema Nacional de Informação e Registo '
  'Animal (SNIRA)» e que a integração plena «traduz-se na clarificação do regime jurídico aplicável aos '
  'pequenos detentores, salvaguardando as exigências de rastreabilidade em matéria de saúde pública e animal». '
  'Passou-se de uma afirmação de continuidade e de desregulação para a afirmação de uma obrigação nova.')
p('Consequência prática desta leitura: não se pode pedir a reposição da exclusão como se fosse um lapso de '
  'transcrição, porque isso equivaleria a pedir a revogação de um objetivo deliberado e legítimo, que é a '
  'rastreabilidade dos pequenos detentores. O que se pode propor é um critério que preserve esse objetivo sem '
  'arrastar os animais detidos exclusivamente para fins de companhia. Foi essa a formulação escolhida para o '
  'texto enviado.')

h2('2.6. O efeito cruzado com o regime dos animais de companhia')
p('O n.º 2 do artigo 1.º do Decreto-Lei n.º 276/2001 exclui do seu âmbito «as espécies da fauna selvagem '
  'autóctone e exótica e os seus descendentes criados em cativeiro, objeto de regulamentação específica, e os '
  'touros de lide e as espécies de pecuária». Lido à letra, o coelho estaria fora daquele diploma. Mas o mesmo '
  'diploma regula-o expressamente: o art.º 26.º, «Condições particulares para a manutenção de pequenos roedores '
  'e coelhos», exige material de cama adequado e remete as medidas das caixas para o Anexo II, que tem quadros '
  'próprios para pequenos roedores, para pequenos roedores em reprodução e para coelhos em reprodução. A '
  'contradição aparente resolve-se pelo fim da detenção e não pela espécie: o coelho mantido como animal de '
  'companhia é animal de companhia, o coelho de uma cunicultura não é. É esse critério implícito que a '
  'eliminação põe em causa, porque o projeto passa a puxar pela espécie e pela atividade.')
p('A isso acresce o artigo criado na mesma revisão, «Registo de estabelecimentos não sujeitos», que remete a '
  'atividade não sujeita a título apenas para as normas sanitárias, de bem-estar animal, identificação e '
  'registo do Decreto-Lei n.º 142/2006, de 27 de julho, sem qualquer referência ao Decreto-Lei n.º 276/2001.')

h2('2.7. A ordem de grandeza dos limiares')
p('A «detenção caseira» admite capacidade instalada até 3 CN, não podendo ser excedido o limite de 2 CN por '
  'espécie pecuária. O projeto não permite, porém, converter esses limiares em número de animais: o seu Anexo '
  'V tem o título «Equivalências em cabeças normais (CN)», uma remissão para a «alínea e) do Artigo 46.º», que '
  'não corresponde à norma que invoca a tabela, e uma nota de pé que define cabeça normal, mas não tem tabela '
  'alguma.')
p('Recorrendo à tabela do regime que substitui, o Anexo II do Decreto-Lei n.º 81/2013, obtêm-se as seguintes '
  'ordens de grandeza para o limite de 2 CN por espécie:')
tabela(
    ['Espécie e tipo de animal', 'CN por animal', 'Animais correspondentes a 2 CN'],
    [['Coelho ou lebre, de recria ou acabamento', '0,009', 'cerca de 222'],
     ['Coelha ou lebre reprodutora, em aleitamento', '0,04', '50'],
     ['Galinha poedeira', '0,013', 'cerca de 153'],
     ['Frango ou pintada', '0,006', 'cerca de 333']],
    larguras=[7.0, 3.0, 5.5],
)
p('Significa que a detenção de algumas centenas de coelhos «para lazer» cabe, em cabeças normais, dentro da '
  'detenção caseira, ficando isenta de licenciamento e com o único dever de registo no SNIRA. Fica ainda '
  'disponível o argumento de que, tratando-se de atividade pecuária, as medidas mínimas do Anexo II do '
  'Decreto-Lei n.º 276/2001 não lhe são oponíveis. Em matéria de bem-estar, é o resultado inverso do '
  'pretendido.')

h2('2.8. O critério que o RGAC adota')
p('O projeto de Regime Geral do Animal de Companhia, em elaboração nesta Direção-Geral e que é trabalho em '
  'curso e não legislação vigente, resolve a fronteira por critério finalista, abandonando a exclusão das '
  'espécies de pecuária:')
cit('«2 - Para efeitos do presente decreto-lei, são animais de companhia os animais das espécies constantes da '
    'Parte A do Anexo I do Regulamento (UE) 2016/429 do Parlamento Europeu e do Conselho, de 9 de março de '
    '2016, e, quando detidos para fins de companhia, os das espécies constantes da Parte B do mesmo anexo.»',
    'n.º 2 do art.º 2.º do projeto de RGAC, versão de trabalho de 30.6.2026, 18h00.')
p('Na Parte A do Anexo I daquele regulamento estão o cão, o gato e o furão; na Parte B, entre outros, os '
  'coelhos e os pequenos roedores. O critério do RGAC qualifica, portanto, as três primeiras espécies como '
  'animais de companhia sem condição, e as restantes quando a detenção tenha esse fim. É este o critério '
  'equivalente a que o texto enviado alude, sem o nomear, por o RGAC não ser ainda diploma publicado.')

h2('2.9. Efeito por espécie, no estado atual do projeto')
tabela(
    ['Espécie', 'Risco de ser lida como atividade pecuária', 'Normas de bem-estar aplicáveis hoje'],
    [['Cães e gatos',
      'Baixo. Não constam da lista nominativa. A cláusula «produção pecuária de animais destinados a animais '
      'de companhia» descreve, à letra, a criação para venda, e a barreira textual que a afastava era a '
      'exclusão eliminada.',
      'Art.ºs 7.º a 18.º e 27.º do DL 276/2001 e Anexo III, com as dimensões mínimas, a proibição de exposição '
      'à venda antes da 8.ª semana, o limite de 15 dias em gaiola e o exercício diário. DL 315/2009 para raças '
      'potencialmente perigosas. Identificação pelo DL 82/2019, no SIAC. A partir de 2028, Regulamento (UE) '
      '2026/1818; o art.º 15.º e o ponto 2 do Anexo I, sobre alojamento, a partir de 31.8.2031.'],
     ['Furões',
      'Baixo, pela mesma via dos cães e gatos.',
      'Nenhuma norma específica. A pesquisa das palavras «furão» e «furões» no DL 276/2001 não devolve '
      'ocorrências; aplicam-se só as normas gerais. Subsiste dúvida de âmbito, por o n.º 2 do art.º 1.º excluir '
      'a fauna selvagem autóctone e os seus descendentes criados em cativeiro. Identificação obrigatória pelo '
      'DL 82/2019. Fora do Regulamento (UE) 2026/1818, que abrange apenas cães e gatos.'],
     ['Coelhos',
      'Elevado. «Leporídeo (coelhos e lebres)» consta da lista nominativa, e a detenção por lazer é o próprio '
      'critério da detenção caseira.',
      'Art.º 26.º do DL 276/2001 e Anexo II, com medidas mínimas de caixa, incluindo quadro para coelhos em '
      'reprodução. Fora do Regulamento (UE) 2026/1818.'],
     ['Hamsters e porquinhos-da-índia',
      'Intermédio. Não constam da lista nominativa, mas cabem em «outra espécie» conjugada com a cláusula '
      'finalista.',
      'Art.º 26.º do DL 276/2001 e Anexo II, als. a) e b), com medidas mínimas para outros pequenos roedores e '
      'para pequenos roedores em reprodução. Fora do Regulamento (UE) 2026/1818.']],
    larguras=[2.6, 5.6, 7.3],
)

h2('2.10. Fontes do ponto 02')
bullet('Projeto: «DL 405 XXV 2026_PECUÁRIA_GSEA 25_09_v1.docx» — preâmbulo; art.ºs 16.º, 48.º, als. b), c) e '
       'i), 49.º e 51.º do anexo; Anexo V. Alterações e comentários registados no próprio ficheiro, com autor e '
       'data, lidos a partir do respetivo XML.')
bullet('Decreto-Lei n.º 81/2013, de 14 de junho, Novo Regime do Exercício da Atividade Pecuária — art.º 1.º, '
       'n.º 3, al. a), art.º 2.º, al. c), e Anexo II, «Equivalências em cabeças normais (CN)». Anexo verificado '
       'no Diário da República, 1.ª série, n.º 113, de 14 de junho de 2013, pp. 3324 e 3325.')
bullet('Decreto-Lei n.º 276/2001, de 17 de outubro, versão consolidada — art.ºs 1.º, n.º 2, 2.º, 26.º e Anexo '
       'II; e, para a comparação por espécie, art.ºs 7.º a 18.º, 27.º e Anexo III.')
bullet('Decreto-Lei n.º 82/2019, de 27 de junho — identificação obrigatória de cães, gatos e furões.')
bullet('Decreto-Lei n.º 142/2006, de 27 de julho — regime de identificação, registo e circulação de animais, '
       'para que remete o artigo «Registo de estabelecimentos não sujeitos» do projeto.')
bullet('Regulamento (UE) 2016/429 do Parlamento Europeu e do Conselho, de 9 de março de 2016, Anexo I, Partes '
       'A e B.')
bullet('Projeto de RGAC, «RGAC_DAJA_REV. FORMAL_V1_Versão TRABALHO - Revisto 30-06-2026 18h00 Grupo.docx» — '
       'art.º 2.º, n.º 2. Trabalho em curso, citado como tal.')

h2('2.11. Verificações pendentes')
bullet('Confirmar se a formulação «produção pecuária de animais destinados a animais de companhia» já constava '
       'do Decreto-Lei n.º 214/2008, de 10 de novembro, que o Decreto-Lei n.º 81/2013 revogou.')
bullet('Confirmar se o Anexo V do projeto virá a reproduzir a tabela do Anexo II do Decreto-Lei n.º 81/2013 ou '
       'se terá valores próprios, e corrigir a remissão para a «alínea e) do Artigo 46.º».')
bullet('Quantificar o universo afetado: número de estabelecimentos do CAE 47762 e, destes, quantos detêm '
       'animais vivos. Sem este dado não é possível avaliar a exequibilidade de procedimentos mais exigentes.')

# ──────────────────── Secção transversal: Lei da Saúde Animal ────────────────────
h1('Secção transversal — A Lei da Saúde Animal como pano de fundo')

p('Esta secção serve os pontos 01, 02 e 09 e explica por que razão o Regulamento (UE) 2016/429, Lei da Saúde '
  'Animal, é o enquadramento que decide as duas questões já fechadas.')

h2('T.1. O registo de 2028 é o registo da Lei da Saúde Animal')
p('A obrigação de notificação dos estabelecimentos do artigo 9.º do Regulamento (UE) 2026/1818 não cria '
  'sistema novo: encaixa no da Lei da Saúde Animal, e di-lo expressamente.')
cit('«3. Member States shall use the information provided for in accordance with Article 84 of Regulation (EU) '
    '2016/429. Operators shall not be required to notify the information already submitted in accordance with '
    'that Article again. 4. The competent authority shall keep a register of establishments. The competent '
    'authority may, for that purpose, use the register established pursuant to Article 101(1), point (a), of '
    'Regulation (EU) 2016/429.»',
    'art.º 9.º, n.ºs 3 e 4, do Regulamento (UE) 2026/1818 — versão inglesa, JO L, 2026/1818, de 10.8.2026.')
cit('«3. Os Estados-Membros devem utilizar as informações facultadas em conformidade com o artigo 84.º do '
    'Regulamento (UE) 2016/429. Os operadores não são obrigados a notificar novamente as informações já '
    'comunicadas em conformidade com o referido artigo. 4. A autoridade competente deve manter um registo de '
    'estabelecimentos. A autoridade competente pode utilizar, para esse efeito, o registo criado nos termos do '
    'artigo 101.º, n.º 1, alínea a), do Regulamento (UE) 2016/429.»',
    'art.º 9.º, n.ºs 3 e 4, do Regulamento (UE) 2026/1818 — versão portuguesa.')
p('Daqui decorre que a obrigação de registo destes estabelecimentos não nasce em 31 de agosto de 2028. O '
  'artigo 84.º da Lei da Saúde Animal, aplicável desde 21 de abril de 2021, impõe aos operadores de '
  'estabelecimentos que detêm animais terrestres o dever de informar a autoridade competente antes de iniciar '
  'a atividade, para que o estabelecimento seja registado; e o artigo 101.º impõe à autoridade competente a '
  'constituição e manutenção dos registos correspondentes. Nenhum dos dois está limitado às espécies '
  'pecuárias, pelo que o criador e a loja de cães e gatos são, nessa aceção, estabelecimentos. O que o '
  'Regulamento de 2026 acrescenta é conteúdo e publicidade, não a obrigação de base.')

h2('T.2. A fronteira que se pede ao diploma nacional já é europeia')
p('A Lei da Saúde Animal distingue animal de companhia de animal detido em estabelecimento, e distingue-o '
  'pela finalidade da detenção:')
cit('«11) «Animal de companhia», um animal detido das espécies listadas no anexo I, que é detido para fins '
    'privados não comerciais;»',
    'art.º 4.º, ponto 11, do Regulamento (UE) 2016/429.')
p('É o mesmo critério que o projeto de RGAC adota no n.º 2 do seu artigo 2.º, com a distinção entre as '
  'espécies da Parte A do Anexo I, que são sempre animais de companhia, e as da Parte B, que o são quando '
  'detidas para fins de companhia. O critério que as observações 02 pedem não é, por isso, uma invenção '
  'nacional: é o critério do direito da União, de aplicação direta, que o próprio projeto já invoca ao '
  'referir, na al. d) do art.º 2.º do anexo, o regime de identificação, registo e circulação dos animais.')

h2('T.3. A «detenção caseira» funde duas categorias que a Lei da Saúde Animal separa')
p('A al. i) do art.º 48.º do anexo define a detenção caseira pela posse de animais cujo objetivo é «de lazer '
  'ou abastecimento do seu detentor». São duas realidades distintas à luz da Lei da Saúde Animal. A detenção '
  'por lazer, sem comércio, corresponde à detenção de animais de companhia na aceção do art.º 4.º, ponto 11, '
  'e não constitui estabelecimento. A detenção para abastecimento do próprio detentor é produção primária e '
  'cai no âmbito do art.º 84.º, com as derrogações que este e o Regulamento Delegado (UE) 2019/2035 admitem '
  'para estabelecimentos de pequena dimensão.')
p('Sujeitar ambas, indistintamente, ao registo no SNIRA trata como exploração pecuária aquilo que o direito '
  'europeu classifica como detenção de animais de companhia. É esta a observação de fundo do ponto 02, e é '
  'por esta via, e não pela reposição da exclusão eliminada, que ela melhor se sustenta: preserva o objetivo '
  'de rastreabilidade de quem produz e deixa de fora quem detém apenas por companhia.')

h2('T.4. O veículo nacional do registo')
p('Em Portugal, o registo que materializa o art.º 101.º da Lei da Saúde Animal quanto a estabelecimentos de '
  'animais de companhia é, hoje, a lista da DGAV das meras comunicações prévias do art.º 3.º-A do Decreto-Lei '
  'n.º 276/2001, divulgada nos termos do n.º 12 do seu art.º 3.º. O projeto de RGAC consolida essa função no '
  'SIAC, que passa a ser a base de dados oficial de identificação e registo dos animais de companhia, do RNAZ '
  'e dos estabelecimentos, e monta no artigo da divulgação dos estabelecimentos a lista pública com número '
  'único de aprovação, exigida também pelo n.º 3 do art.º 10.º do Regulamento (UE) 2026/1818.')
p('É esta a ligação ao ponto 01: encaminhada a loja de animais para um procedimento puramente municipal, sem '
  'pronúncia nem notificação da DGAV, nada do que o procedimento recolhe alimenta o registo que a autoridade '
  'competente tem de manter. E é esta a ligação ao ponto 09: o cadastro setorial e a base de dados setorial do '
  'anexo não incluem o SIAC entre os sistemas a articular.')

h2('T.5. Fontes e verificações pendentes desta secção')
bullet('Regulamento (UE) 2026/1818 — art.º 9.º, n.ºs 3 e 4, e art.º 10.º, n.º 3, nas versões inglesa e '
       'portuguesa do texto publicado no Jornal Oficial. Confirmado.')
bullet('Regulamento (UE) 2016/429 — art.º 4.º, ponto 11. Confirmado.')
bullet('Regulamento (UE) 2016/429 — art.ºs 84.º, 93.º e 101.º. O teor foi apurado por fonte secundária: o '
       'acesso ao EUR-Lex foi recusado a partir deste ambiente, pelo que o texto não foi lido no Jornal '
       'Oficial. A descrição acima não deve ser usada como citação verbatim sem essa confirmação.')
bullet('Regulamento Delegado (UE) 2019/2035 da Comissão — referido pelo projeto de RGAC, inclusive quanto ao '
       'art.º 71.º-A, em matéria de rastreabilidade de cães, gatos e furões. Não verificado nesta sessão.')
bullet('Verificação pendente: se os estabelecimentos de animais de companhia estão hoje efetivamente '
       'registados ao abrigo do art.º 84.º da Lei da Saúde Animal e em que sistema. É a pergunta que determina '
       'se existe já um incumprimento a corrigir ou apenas um risco a prevenir.')

# ───────────────────────────── Registo ─────────────────────────────
h1('Registo de alterações')
tabela(
    ['Data', 'Alteração'],
    [['6.10.2026', 'Criação do documento. Índice da série e numeração. Fundamentação dos pontos 01 e 02.'],
     ['6.10.2026', 'Secção transversal sobre a Lei da Saúde Animal, Regulamento (UE) 2016/429, que enquadra '
                   'os pontos 01, 02 e 09. Observações 02 em terceira versão, com o alerta reduzido ao '
                   'essencial e o enquadramento no Decreto-Lei n.º 276/2001 e na Lei da Saúde Animal.'],
     ['6.10.2026', 'Observações 02 em segunda versão, em linguagem corrente e em registo de dúvida, a '
                   'pedido. A primeira versão fica no processo. Ambas assentam na fundamentação deste ponto 02, '
                   'que não foi alterada.'],
     ['6.10.2026', 'Correção ao contributo interno de 5.10.2026: o seu ponto 5.4 afirmava que a venda de '
                   'animais em feiras e no comércio não sedentário ficava sem requisitos de bem-estar. O n.º 6 '
                   'do art.º 35.º do Decreto-Lei n.º 276/2001 proíbe a venda ambulante de animais de companhia '
                   'e o mesmo artigo fixa as condições da venda em feiras e mercados, com mera comunicação '
                   'prévia à câmara para vistoria pelo médico veterinário municipal. Não há lacuna, há falta de '
                   'articulação. O tema passou ao ponto 06 desta série.']],
    larguras=[2.4, 13.1],
)

doc.save('/home/user/Legislacao/DL 405_2026/Fundamentacao_das_Observacoes.docx')
print('OK')
