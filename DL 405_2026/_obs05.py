# -*- coding: utf-8 -*-
"""Observações 05 — acesso de animais de companhia a estabelecimentos."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Acesso de animais de companhia a estabelecimentos'
SUB = 'Projeto de DL das atividades económicas, versão de 25.9.2026. À consideração.'

TEXTO = [
 'O artigo permite a permanência de animais de companhia em espaços fechados, mediante autorização da '
 'entidade exploradora manifestada por dístico à entrada, com trela curta ou acondicionamento adequado, sem '
 'circulação livre e com proibição total de permanência nas zonas da área de serviço e junto aos locais onde '
 'estão expostos alimentos para venda. Deixamos quatro notas.',

 'A primeira, nos cães de assistência. O Decreto-Lei n.º 74/2007, de 27 de março, não consagra uma permissão '
 'de permanência: consagra um direito de acesso que, nos termos do n.º 1 do seu artigo 3.º, «prevalece sobre '
 'quaisquer proibições ou limitações que contrariem o disposto no presente decreto-lei, ainda que assinaladas '
 'por placas ou outros sinais distintivos». Os estabelecimentos de comércio e os da restauração e do turismo '
 'constam expressamente das alíneas i) e j) do artigo 2.º. Daí três consequências: o dístico não é oponível ao '
 'cão de assistência; o direito assiste à pessoa com deficiência ou ao treinador habilitado, e não ao '
 '«portador»; e a recusa só é admissível nos termos do n.º 3 do artigo 3.º daquele diploma, por sinais '
 'manifestos de doença, agressividade ou falta de higiene, e não pelo critério do normal funcionamento do '
 'estabelecimento. Sugere-se ressalvar expressamente o Decreto-Lei n.º 74/2007 e excluir os cães de '
 'assistência do regime de autorização e de recusa deste artigo.',

 'A segunda, nos cães perigosos e potencialmente perigosos. O n.º 2 do artigo 13.º do Decreto-Lei n.º 315/2009, '
 'de 29 de outubro, exige, para estes animais, «açaimo funcional que não permita comer nem morder e, neste '
 'caso, devidamente seguro com trela curta até 1 m de comprimento, que deve estar fixa a coleira ou a '
 'peitoral». A referência do projeto a trela curta, sem medida e sem açaimo, fica abaixo desse padrão. '
 'Sugere-se ressalvar o Decreto-Lei n.º 315/2009, para que não se gere dúvida sobre qual das duas normas '
 'prevalece no interior do estabelecimento.',

 'A terceira, na higiene dos alimentos. A proibição está circunscrita às zonas da área de serviço e aos locais '
 'de exposição de alimentos para venda, ficando de fora as zonas de preparação e de armazenagem. O Regulamento '
 '(CE) n.º 852/2004 impõe procedimentos que impeçam o acesso de animais domésticos aos locais onde os alimentos '
 'são preparados, manuseados ou armazenados, salvo autorização da autoridade competente. Convirá alinhar a '
 'alínea com esse âmbito, mais largo. (Citação exata a confirmar no Anexo II, Capítulo IX, do Regulamento.)',

 'A quarta, no que o artigo não exige. A permanência é admitida sem qualquer referência à identificação '
 'eletrónica e ao registo, obrigatórios nos termos do Decreto-Lei n.º 82/2019, de 27 de junho, nem à vacinação '
 'antirrábica válida. São condições que a lei já impõe à circulação do animal e cuja remissão, aqui, seria útil. '
 'Assinala-se também que o artigo está construído na perspetiva do funcionamento do estabelecimento e não '
 'contempla o bem-estar do próprio animal durante a permanência. Ficamos ao dispor para ajudar na redação.',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(12.5)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.italic = True; r.font.size = Pt(9.5)
for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Observacoes_05_Acesso_animais_estabelecimentos.docx')
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
