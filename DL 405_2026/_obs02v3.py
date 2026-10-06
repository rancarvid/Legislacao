# -*- coding: utf-8 -*-
"""Observações 02, v.3 — alerta simples, com o enquadramento no 276/2001 e na Lei da Saúde Animal."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Animais de companhia e atividade pecuária'
SUB = 'Projeto de DL das atividades económicas, versão de 25.9.2026. À consideração.'

TEXTO = [
 'A definição de atividade pecuária, na alínea c) do artigo 48.º do anexo, exclui a apicultura. Até 21 de '
 'setembro excluía também os animais de companhia, e foi retirada, ficando «com exceção da apicultura '
 'eem:».',

 'Assinalamos porque o regime em vigor mantém ambas: a alínea a) do n.º 3 do artigo 1.º do Decreto-Lei '
 'n.º 81/2013 afasta do seu âmbito a apicultura e os animais de companhia.',

 'É no artigo 26.º do Decreto-Lei n.º 276/2001 que estão as medidas das caixas para pequenos roedores e '
 'coelhos, mas esse diploma exclui «as espécies de pecuária». E a Lei da Saúde Animal, no ponto 11 do '
 'artigo 4.º do Regulamento (UE) 2016/429, define animal de companhia pela finalidade: o animal das '
 'espécies do anexo I «detido para fins privados não comerciais».',

 'Compreende-se o objetivo de rastreabilidade dos pequenos detentores. Dá-se à consideração que se esclareça '
 'o alcance pretendido e que se adote aquele critério, ressalvando que o animal detido para fins privados não '
 'comerciais é animal de companhia para efeitos de bem-estar, sem prejuízo da identificação e do registo. E '
 'que se corrija o «eem».',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(12.5)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.italic = True; r.font.size = Pt(9.5)
for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Observacoes_02_Animais_companhia_pecuaria_v3.docx')
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
