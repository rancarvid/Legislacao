# -*- coding: utf-8 -*-
"""Observações 03 — deferimento tácito em atividades com animais vivos."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Deferimento tácito em atividades com animais vivos'
SUB = 'Projeto de DL das atividades económicas, versão de 25.9.2026. À consideração.'

TEXTO = [
 'O preâmbulo assume consagrar como regra o deferimento tácito, e o n.º 3 do artigo 25.º concretiza-o: não '
 'havendo pronúncia no prazo do Anexo III, considera-se favorável à pretensão do requerente. Nas consultas sem '
 'prazo próprio, esse prazo é de 20 dias.',

 'A dúvida que deixamos é se a regra é adequada quando a atividade envolve a detenção de animais vivos.',

 'Na legislação vigente, quando a decisão respeita à detenção de animais, o legislador afastou-a '
 'expressamente. O n.º 2 do artigo 3.º-D do Decreto-Lei n.º 276/2001 dispõe que, decorridos 60 dias sobre o '
 'pedido devidamente instruído, e independentemente da visita de controlo, «não há lugar a deferimento '
 'tácito».',

 'No mesmo sentido, a partir de 31 de agosto de 2034 o artigo 10.º do Regulamento (UE) 2026/1818 só permite '
 'colocar cães ou gatos no mercado após aprovação do estabelecimento de criação pela autoridade competente, '
 'mediante inspeção no local.',

 'Acresce que o artigo 31.º, ao enumerar as causas de indeferimento, não contempla qualquer pronúncia em '
 'matéria de saúde ou de bem-estar animal. Ficamos ao dispor para ajudar na clarificação.',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(12.5)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.italic = True; r.font.size = Pt(9.5)
for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Observacoes_03_Deferimento_tacito.docx')
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
