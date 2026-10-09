# -*- coding: utf-8 -*-
"""Observações 04, v.4 — acrescentada a Lei n.º 38/2026 e a convergência dos dois deveres."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Animais errantes e centros de recolha oficial'
SUB = 'Projeto de DL das atividades económicas, versão de 25.9.2026. À consideração.'

TEXTO = [
 'Sobre a questão de incluir neste diploma a recolha de animais errantes de todas as espécies, deixamos o que '
 'esta área tem a oferecer em concreto e os aspetos que nos parece importante ressalvar.',

 'Uma nota de contexto primeiro, porque muda a dimensão da questão. Desde 8 de agosto está em vigor a Lei '
 'n.º 38/2026, de 3 de agosto, que integrou os animais na proteção civil. A alínea e) do n.º 2 do artigo 2.º '
 'da Lei n.º 65/2007 inclui agora, nos domínios da proteção civil municipal, o planeamento de soluções de '
 'emergência para «a evacuação, alojamento e abastecimento dos animais presentes no município, incluindo a '
 'realização de simulacros». A norma não distingue espécies. No mesmo sentido, a alínea d) do artigo 13.º do '
 'Decreto-Lei n.º 90-A/2022 prevê, no teatro de operações, «uma zona de concentração de acolhimento de '
 'animais».',

 'Ou seja, os municípios já têm, desde agosto, um dever de planeamento que inclui o alojamento de animais de '
 'qualquer espécie em emergência. Se a este acrescer a recolha corrente de animais errantes de espécie '
 'pecuária prevista no artigo 59.º, os dois deveres convergem na mesma necessidade física — um lugar onde '
 'alojar estas espécies sob responsabilidade municipal — e para nenhum deles existe hoje norma de bem-estar '
 'aplicável. É o que nos leva ao ponto seguinte.',

 'O artigo 59.º pressupõe um período de guarda do animal, mas não identifica a instalação nem as condições '
 'em que essa guarda se faz. O n.º 3 manda notificar o detentor para pagar «as despesas realizadas com a '
 'retirada, captura, transporte, alojamento, alimentação, identificação e cuidados médico-veterinários», pelo '
 'que há alojamento e alimentação a cargo de alguém, por um prazo que é de cinco dias úteis quando o detentor '
 'é conhecido e indeterminado quando não é. Note-se que a única exigência de adequação de alojamento, no '
 'n.º 4, recai sobre a exploração pecuária de origem, que só pode receber o animal de volta após verificação '
 'de que «dispõe de condições adequadas de alojamento e contenção». Para o período de guarda não há norma '
 'equivalente.',

 'E o parâmetro não existiria mesmo que o artigo o invocasse, por dupla exclusão: o n.º 2 do artigo 1.º do '
 'Decreto-Lei n.º 276/2001 exclui do seu âmbito «as espécies de pecuária», e o artigo 2.º da Portaria '
 'n.º 146/2017 circunscreve o regime dos centros de recolha às espécies da Parte A do Anexo I do Regulamento '
 '(UE) 2016/429, ou seja, ao cão, ao gato e ao furão. Quando o n.º 5 do artigo 59.º manda determinar o destino '
 'do animal «observadas as normas relativas à identificação, registo, circulação, sanidade e bem-estar '
 'animal», é para normas que, quanto a estas espécies e neste contexto, estão por fixar.',

 'É esta a lacuna que nos parece dever ficar ressalvada, e é também o contributo que podemos dar. Uma norma '
 'técnica de alojamento para estas espécies teria de fixar, no mínimo: áreas e dimensões por espécie e por '
 'fase de desenvolvimento; separação entre espécies incompatíveis, incluindo o isolamento visual, sonoro e '
 'olfativo entre presas e predadores; quarentena e biossegurança à entrada, uma vez que o animal chega sem '
 'identificação e sem exploração de origem conhecida; alimentação e abeberamento próprios da espécie; meios de '
 'contenção e de maneio seguro; e lotação máxima determinada pelos recursos efetivos da instalação. Deixamos '
 'ainda uma nota de dimensionamento: o compartimento para outras espécies previsto nos Avisos de financiamento '
 'tem área mínima de três metros quadrados por animal, medida pensada para espécies de pequeno porte e '
 'insuficiente para um equídeo ou um bovino. O valor não pode ser transposto sem revisão.',

 'A resposta física não seria, em todo o caso, inteiramente nova. Desde o Despacho n.º 3321/2018 até ao Aviso '
 'n.º 1/2025, um compartimento para outras espécies foi requisito mínimo de construção de um CRO municipal e '
 'item financiado, com apoio que evoluiu de 1.300 € para 1 537,02 €. Manteve-se oito anos, primeiro sob o ICNF '
 'e depois sob a DGAV. Não consta do Aviso n.º 1/2026 e pode ser reposto por via administrativa.',

 'A instalação de destino determina o regime aplicável ao animal, e a escolha não é indiferente. Entregue o '
 'animal a um centro de recolha oficial, aplica-se a Lei n.º 27/2016, que é espécie-neutra, uma vez que fala '
 'em «centros de recolha oficial de animais», com a presunção de abandono, a esterilização obrigatória, o '
 'encaminhamento para adoção e a proibição do abate como forma de controlo da população. O n.º 4 do artigo '
 '57.º, para os animais apreendidos em exploração, aponta em sentido oposto: matadouro, se aprovados para '
 'consumo, ou destruição. Como o artigo 59.º não identifica instalação alguma, a consequência fica em aberto. '
 'Parece-nos matéria a esclarecer expressamente.',

 'A qualificação da própria instalação merece atenção. A alínea b) do artigo 48.º inclui o equídeo e o '
 'leporídeo na definição de animal de espécie pecuária, e os limiares assentam em cabeças normais, cuja '
 'tabela de equivalências o Anexo V ainda não contém, pelo que não é possível determinar a partir de que '
 'efetivo uma instalação de acolhimento passaria a exercer atividade pecuária sujeita a título. A questão é '
 'prática: o Projeto de Resolução n.º 82/XIV/1.ª registava que «para se proceder à criação de um santuário de '
 'animais de quinta, é obrigatória a inscrição como exploração de animais de pecuária». Seria útil deixar '
 'claro que a instalação que recebe animais ao abrigo do artigo 59.º não exerce, por esse facto, atividade '
 'pecuária.',

 'Duas notas de execução. As despesas ficam sem titular na hipótese que mais interessa: os n.os 3 e 6 '
 'imputam-nas ao detentor identificado, mas o n.º 5 regula precisamente o caso em que este não é conhecido, e '
 'aí o custo da permanência recai sobre a instalação. E os prazos não coincidem — oito dias de permanência '
 'mínima pelo n.º 1 do artigo 9.º do Decreto-Lei n.º 314/2003, quinze dias para a presunção de abandono pelo '
 'n.º 1 do artigo 3.º da Lei n.º 27/2016, cinco dias úteis pelo n.º 3 do artigo 59.º —, pelo que, recebendo a '
 'mesma instalação animais ao abrigo de regimes diferentes, correm prazos distintos em paralelo.',

 'Não há, por último, base de medida do problema: o Relatório de Atividades dos Centros de Recolha Oficial de '
 '2025 refere que «alguns CRO têm ações realizadas com outras espécies», mas esses dados não foram '
 'considerados para os totais. A recolha dessa informação é algo que podemos assegurar.',

 'Uma nota de precisão: o centro de recolha oficial e o «centro de agrupamento» definido no artigo 48.º são '
 'figuras distintas, com regimes e finalidades próprias.',

 'Ficamos ao dispor para desenvolver qualquer destes pontos, em especial os critérios de alojamento.',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(12.5)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.italic = True; r.font.size = Pt(9.5)
for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Observacoes_04_Animais_errantes_v4.docx')
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
