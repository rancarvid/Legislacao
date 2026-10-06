# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Venda de animais de companhia: a ausência da DGAV no novo DL das atividades económicas'

SUB = 'DL 405/XXV/2026, versão de 25.9.2026. DSBEA, 6 de outubro de 2026.'

TEXTO = [
 'O projeto de DL das atividades económicas revoga, na alínea j) do n.º 1 do artigo 8.º, o '
 'Decreto-Lei n.º 10/2015, de 16 de janeiro. É a esse diploma que remete o n.º 1 do artigo 3.º '
 'do Decreto-Lei n.º 276/2001, de 17 de outubro, quando ressalva «Sem prejuízo do disposto no Decreto-Lei '
 'n.º 10/2015, de 16 de janeiro, quanto aos estabelecimentos de comércio a retalho de animais de companhia». '
 'Revogado o diploma remetido, o acesso à atividade das lojas que vendem animais vivos passa a estar no '
 'novo diploma.',

 'O novo diploma recolhe essa atividade no n.º 4 do artigo 16.º do seu anexo e sujeita-a a comunicação '
 'prévia. '
 'No Anexo II, que identifica as entidades intervenientes, correspondem-lhe a câmara municipal como '
 'entidade coordenadora, a Direção-Geral da Defesa do Consumidor, Comércio e Serviços como entidade '
 'notificada e nenhuma entidade pública consultada. A DGAV não figura em nenhuma das três colunas.',

 'O alerta é este. A DGAV é a autoridade competente em matéria de animais de companhia, nos termos da '
 'alínea x) do n.º 1 do artigo 2.º do Decreto-Lei n.º 276/2001, e cabe-lhe verificar as condições de '
 'alojamento, higiene, saúde e bem-estar dos animais detidos para venda. Na redação proposta, a abertura de '
 'um estabelecimento que detém animais vivos deixa de chegar ao seu conhecimento. A assimetria é difícil de '
 'justificar: para os alimentos para animais de criação, a alínea f) do n.º 2 do artigo 16.º exige vistoria '
 'prévia e aprovação da DGAV; para o animal, nada.',

 'A partir de 31 de agosto de 2028 torna-se aplicável o artigo 9.º do Regulamento (UE) 2026/1818, que '
 'obriga os operadores a notificar a autoridade competente quanto a cada estabelecimento, com indicação da '
 'localização, do tipo, das espécies e da capacidade, e obriga essa autoridade a manter um registo de '
 'estabelecimentos. A loja de animais é, para esse efeito, um «estabelecimento de venda», e o considerando '
 '23 esclarece que não se fixam limiares. Sem a DGAV no procedimento, falta a informação para constituir '
 'esse registo.',

 'Propõe-se o aditamento da DGAV às entidades públicas consultadas do Anexo II, na linha relativa a estes '
 'estabelecimentos, ou, em alternativa mínima, às entidades notificadas. É uma alteração de uma linha e '
 'não acrescenta prazo ao operador.',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph()
par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(13)

par = doc.add_paragraph()
par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.font.size = Pt(9.5); r.italic = True

for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Ponto1_Venda_animais_DGAV.docx')
print('palavras corpo:', sum(len(t.split()) for t in TEXTO))
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
