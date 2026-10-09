# -*- coding: utf-8 -*-
"""Observações 02, v.2 — linguagem corrente, a levantar a dúvida."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Animais de companhia e atividade pecuária'
SUB = 'Observações ao projeto de DL das atividades económicas (versão de 25.9.2026), à consideração. V.2.'

TEXTO = [
 'Ao ler a definição de atividade pecuária, na alínea c) do artigo 48.º do anexo, ficámos com uma dúvida que '
 'preferimos colocar do que deixar passar.',

 'A definição exclui a apicultura. Até à versão de 21 de setembro excluía também os animais de companhia, e '
 'essa parte foi retirada. Ficou a frase «com exceção da apicultura eem:», o que parece indicar que a '
 'eliminação não foi relida.',

 'A dúvida é simples: pretendeu-se mesmo que a criação e a detenção de animais de companhia passem a ser '
 'atividade pecuária? Perguntamos porque o regime em vigor diz o contrário. O Decreto-Lei n.º 81/2013 afasta '
 'do seu âmbito a apicultura e os animais de companhia, e no projeto só a apicultura ficou.',

 'E perguntamos também porque a alínea anterior mantém a referência à produção de animais destinados a '
 'animais de companhia. Sozinha, sem a exclusão que a acompanhava, essa referência pode levar a que um coelho '
 'ou um hamster detidos por gosto sejam tratados como efetivo pecuário, com registo no SNIRA, quando é no '
 'Decreto-Lei n.º 276/2001 que estão as regras de alojamento dessas espécies, incluindo as medidas das caixas.',

 'Se não foi essa a intenção, basta repor a exclusão, ou dizer de outro modo que os animais detidos apenas '
 'para companhia não entram aqui. Ficamos ao dispor para ajudar na redação. E convém, em qualquer caso, '
 'corrigir o «eem».',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(12.5)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.italic = True; r.font.size = Pt(9.5)
for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Observacoes_02_Animais_companhia_pecuaria_v2.docx')
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
