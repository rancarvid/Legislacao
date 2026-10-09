# -*- coding: utf-8 -*-
"""Observações 04 — animais errantes e articulação com os centros de recolha oficial."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Animais errantes e centros de recolha oficial'
SUB = 'Projeto de DL das atividades económicas, versão de 25.9.2026. À consideração.'

TEXTO = [
 'Sobre a questão de incluir neste diploma os animais errantes de todas as espécies, deixamos o que esta área '
 'tem a oferecer em concreto.',

 'A base legal não é obstáculo. A Lei n.º 27/2016 cria uma rede de «centros de recolha oficial de animais», '
 'sem restrição de espécie. A limitação ao cão, ao gato e ao furão foi introduzida pelo artigo 2.º da Portaria '
 'n.º 146/2017, por remissão para a Parte A do Anexo I do Regulamento (UE) 2016/429. Sendo matéria de '
 'portaria, o âmbito pode ser ajustado sem alteração legislativa.',

 'A solução física também não é nova. Desde o Despacho n.º 3321/2018 até ao Aviso n.º 1/2025, um compartimento '
 'para outras espécies foi requisito mínimo de construção de um CRO municipal e item financiado, com apoio que '
 'evoluiu de 1.300 € para 1 537,02 € por três metros quadrados. Manteve-se oito anos, primeiro sob o ICNF e '
 'depois sob a DGAV. Não consta do Aviso n.º 1/2026 e pode ser reposto por via administrativa.',

 'E a prática já existe. O Relatório de Atividades dos CRO de 2025 refere que «alguns CRO têm ações realizadas '
 'com outras espécies», dados que não foram considerados para os totais.',

 'No plano comparado, os Países Baixos e a Bélgica consagram um dever de recolha redigido para «um animal», '
 'qualquer que seja a espécie. A Lei Regional da Calábria n.º 45/2023 admite a instalação única com estruturas '
 '«separadas, física e funcionalmente» e serviços comuns. O código de boas práticas galês de 2020 é o único '
 'que fixa requisitos técnicos para instalações municipais multiespécie.',

 'A advertência vem da Irlanda, onde a instalação multiespécie existe por lei desde 1935 mas o poder de fixar '
 'padrões técnicos nunca foi exercido: a revisão governamental de 2017 concluiu que essas instalações estão '
 '«inadequately regulated». É aqui que o contributo desta área se torna indispensável — não na figura, mas na '
 'norma técnica: separação por espécie, quarentena, biossegurança e lotação em função dos recursos.',

 'Uma nota de precisão: o centro de recolha oficial e o «centro de agrupamento» definido no artigo 48.º são '
 'figuras distintas, com regimes e finalidades próprias.',

 'Ficamos ao dispor para desenvolver qualquer destes pontos.',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(12.5)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.italic = True; r.font.size = Pt(9.5)
for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Observacoes_04_Animais_errantes.docx')
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
