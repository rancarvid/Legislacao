# -*- coding: utf-8 -*-
"""Observações 04, v.5 — cinco pontos, só no âmbito do bem-estar dos animais de companhia."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Animais errantes e centros de recolha oficial'
SUB = 'Projeto de DL das atividades económicas, versão de 25.9.2026. À consideração.'

TEXTO = [
 'Sobre a questão de incluir neste diploma a recolha de animais errantes de todas as espécies, deixamos cinco '
 'observações, no âmbito do bem-estar dos animais de companhia. A montante, note-se que a Lei n.º 38/2026, de '
 '3 de agosto, já cometeu aos municípios o planeamento do alojamento de animais de qualquer espécie em '
 'situação de emergência, pelo que a incumbência aqui em causa acresceria a uma que já existe.',

 'Primeira. O artigo 59.º pressupõe um período de guarda do animal, mas não identifica a instalação. O n.º 3 '
 'manda notificar o detentor para pagar as despesas realizadas com «a retirada, captura, transporte, '
 'alojamento, alimentação, identificação e cuidados médico-veterinários», pelo que há alojamento e alimentação '
 'a cargo de alguém. Sendo a rede dos centros de recolha oficial a única resposta pública existente para '
 'recolha de animais errantes, a matéria toca diretamente a nossa área.',

 'Segunda. Se a resposta vier a ser o centro de recolha oficial, o regime atual não o habilita para estas '
 'espécies. O n.º 2 do artigo 1.º do Decreto-Lei n.º 276/2001 exclui do seu âmbito «as espécies de pecuária», '
 'e o artigo 2.º da Portaria n.º 146/2017 circunscreve o regime dos centros de recolha às espécies da Parte A '
 'do Anexo I do Regulamento (UE) 2016/429, ou seja, ao cão, ao gato e ao furão. Um equídeo alojado num CRO '
 'fica, hoje, fora de ambos.',

 'Terceira, e é o que nos parece mais importante ressalvar: faltam os critérios de bem-estar do alojamento. '
 'Uma norma técnica para estas espécies teria de fixar áreas e dimensões por espécie e fase de '
 'desenvolvimento, separação entre espécies incompatíveis, com isolamento visual, sonoro e olfativo entre '
 'presas e predadores, quarentena e biossegurança à entrada, uma vez que o animal chega sem identificação e '
 'sem origem conhecida, alimentação e abeberamento próprios da espécie, meios de contenção e de maneio seguro '
 'e lotação máxima determinada pelos recursos efetivos da instalação. É o contributo concreto que esta área '
 'pode dar.',

 'Quarta. Importa ter presente que a escolha da instalação arrasta o regime do animal. Entrando num centro de '
 'recolha oficial, passa a aplicar-se-lhe a Lei n.º 27/2016, que não distingue espécies, com a presunção de '
 'abandono, a esterilização obrigatória, o encaminhamento para adoção e a proibição do abate como forma de '
 'controlo da população. É consequência que convém ser deliberada e não incidental.',

 'Quinta. A infraestrutura já esteve prevista. Os Avisos de financiamento para construção e modernização de '
 'centros de recolha oficial fixaram, entre 2018 e 2025, um compartimento para outras espécies como requisito '
 'mínimo de um CRO municipal, com área mínima coberta de três metros quadrados por animal, primeiro sob o ICNF '
 'e depois sob a DGAV. O requisito não consta do Aviso em vigor e pode ser reposto por via administrativa. '
 'Deixamos, porém, a nota de que aquela área foi dimensionada para espécies de pequeno porte e é insuficiente '
 'para um equídeo, pelo que a reposição não deveria fazer-se por simples transcrição.',

 'Uma nota de precisão: o centro de recolha oficial e o «centro de agrupamento» definido no artigo 48.º são '
 'figuras distintas, com regimes e finalidades próprias.',

 'Ficamos ao dispor para desenvolver estes critérios, caso seja útil.',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(12.5)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.italic = True; r.font.size = Pt(9.5)
for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Observacoes_04_Animais_errantes_v5.docx')
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
