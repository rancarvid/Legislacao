# -*- coding: utf-8 -*-
"""Observações da DSBEA ao DL das atividades económicas. Ponto 2: animais de companhia e pecuária."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Animais de companhia e atividade pecuária'
SUB = 'Observações ao projeto de DL das atividades económicas (versão de 25.9.2026), à consideração.'

TEXTO = [
 'Na alínea c) do artigo 48.º do anexo, a definição de «atividade pecuária» exclui a apicultura. Na versão '
 'anterior excluía também os animais de companhia: a expressão foi eliminada em alteração registada de 21 de '
 'setembro de 2026, ficando a sequência «com exceção da apicultura eem:».',

 'A exclusão vinha do regime em vigor. A alínea a) do n.º 3 do artigo 1.º do Decreto-Lei n.º 81/2013 afasta '
 'do seu âmbito a apicultura, CAE 01491, e os animais de companhia, CAE 01493. No projeto subsiste a primeira '
 'e desaparece a segunda.',

 'A alínea b) do mesmo artigo mantém a referência à «produção pecuária de animais destinados a animais de '
 'companhia», vinda do artigo 2.º daquele decreto-lei. Sem a exclusão que a acompanhava, a detenção de '
 'leporídeos ou de pequenos roedores para fins de companhia pode ser lida como atividade pecuária, e como '
 'detenção caseira sujeita a registo no SNIRA. Acresce que o n.º 2 do artigo 1.º do Decreto-Lei n.º 276/2001 '
 'exclui «as espécies de pecuária», sendo o seu artigo 26.º que fixa as medidas mínimas das caixas para '
 'pequenos roedores e coelhos.',

 'Dá-se à consideração que se esclareça o alcance pretendido e que, não havendo intenção de alterar o '
 'enquadramento dos animais de companhia, se reponha a exclusão ou se adote critério que ressalve a detenção '
 'para fins exclusivos de companhia. Fica também à consideração a correção da sequência «eem».',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(12.5)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.italic = True; r.font.size = Pt(9.5)
for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Observacoes_02_Animais_companhia_pecuaria.docx')
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
