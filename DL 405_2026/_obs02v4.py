# -*- coding: utf-8 -*-
"""Observações 02, v.4 — dúvida de aplicação: qual a norma de bem-estar aplicável."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Animais de companhia e atividade pecuária'
SUB = 'Projeto de DL das atividades económicas, versão de 25.9.2026. À consideração.'

TEXTO = [
 'A propósito da questão já colocada neste articulado sobre quais as normas de bem-estar aplicáveis à '
 'atividade pecuária, deixamos uma dúvida de aplicação.',

 'A definição de atividade pecuária, na alínea c) do artigo 48.º, exclui a apicultura. Até 21 de setembro '
 'excluía também os animais de companhia, e essa parte foi retirada.',

 'A dúvida é esta: a detenção caseira de coelhos ou de pequenos roedores por lazer fica sujeita a que normas '
 'de bem-estar?',

 'A Portaria n.º 635/2009 regula a detenção e produção de leporídeos, mas assenta nas classes de '
 'licenciamento, e a detenção caseira é expressamente isenta. O artigo 51.º remete para o Decreto-Lei '
 'n.º 142/2006, que estabelece a identificação, o registo e a circulação. E o Decreto-Lei n.º 276/2001, cujo '
 'artigo 26.º fixa as medidas das caixas para pequenos roedores e coelhos, exclui do seu âmbito «as espécies '
 'de pecuária».',

 'Parece-nos útil que a resposta fique clara antes da consulta pública, uma vez que estes animais passam a ter '
 'registo sem que se identifique a norma que assegura o seu bem-estar. Ficamos ao dispor para ajudar nessa '
 'clarificação.',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(12.5)
par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.italic = True; r.font.size = Pt(9.5)
for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Observacoes_02_Animais_companhia_pecuaria_v4.docx')
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
