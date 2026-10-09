# -*- coding: utf-8 -*-
"""Observações da DSBEA ao DL das atividades económicas. Ponto 1: estabelecimentos de venda."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITULO = 'Estabelecimentos de comércio a retalho de animais de companhia'
SUB = ('Observações ao projeto de DL das atividades económicas (DL 405/XXV/2026, versão de 25.9.2026), '
       'que se dão à consideração.')

TEXTO = [
 'A atividade consta do n.º 4 do artigo 16.º do anexo, sujeita a comunicação prévia. Na linha '
 'correspondente do Anexo II figuram a câmara municipal como entidade coordenadora e a Direção-Geral da '
 'Defesa do Consumidor, Comércio e Serviços como entidade notificada, sem entidades públicas consultadas. '
 'Neste procedimento o título é emitido automaticamente com a apresentação do pedido, nos termos do n.º 3 '
 'do artigo 29.º.',

 'Hoje a atividade rege-se pelo Decreto-Lei n.º 10/2015, de 16 de janeiro, por remissão expressa do n.º 1 '
 'do artigo 3.º do Decreto-Lei n.º 276/2001, de 17 de outubro, sendo neste último que estão os requisitos '
 'de detenção dos animais destinados a venda. Revogado o Decreto-Lei n.º 10/2015 pela alínea j) do n.º 1 '
 'do artigo 8.º, a matéria transita para o presente diploma sem intervenção da DGAV, que a alínea x) do '
 'n.º 1 do artigo 2.º do Decreto-Lei n.º 276/2001 identifica como autoridade competente.',

 'Acresce o que se aproxima. A partir de 31 de agosto de 2028 estes estabelecimentos são «estabelecimentos '
 'de venda» na aceção da alínea q) do artigo 4.º do Regulamento (UE) 2026/1818, e o artigo 9.º obriga o '
 'operador a notificar a autoridade competente quanto a cada estabelecimento, com indicação da localização, '
 'do tipo, das espécies e da capacidade, e essa autoridade a manter um registo de estabelecimentos. Sem '
 'intervenção no licenciamento, a DGAV não terá como reunir a informação para constituir esse registo.',

 'Dá-se à consideração, em alternativa, a passagem da atividade ao procedimento de comunicação com prazo, '
 'no n.º 3 do artigo 16.º, com a DGAV como entidade pública consultada, única via que permite pronúncia '
 'prévia, podendo aproveitar-se a formulação que o n.º 2 do artigo 54.º do anexo já usa para a pecuária; '
 'ou, mantendo-se a comunicação prévia, a simples inclusão da DGAV entre as entidades notificadas da linha '
 'respetiva do Anexo II, que assegura apenas o conhecimento dos estabelecimentos, mas é de aplicação '
 'imediata e sem custo para o operador.',

 'Observa-se, por último, que o CAE 47762 abrange estabelecimentos que não detêm animais vivos, pelo que '
 'qualquer exigência adicional deverá distinguir essas situações.',
]

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.8)
    s.top_margin = s.bottom_margin = Cm(2.5)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(8)
n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(4)
r = par.add_run(TITULO); r.bold = True; r.font.size = Pt(12.5)

par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(16)
r = par.add_run(SUB); r.italic = True; r.font.size = Pt(9.5)

for t in TEXTO:
    doc.add_paragraph(t)

doc.save('/home/user/Legislacao/DL 405_2026/Observacoes_01_Estabelecimentos_de_venda.docx')
print('palavras total:', sum(len(t.split()) for t in TEXTO) + len(TITULO.split()) + len(SUB.split()))
