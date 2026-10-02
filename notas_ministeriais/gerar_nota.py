# -*- coding: utf-8 -*-
"""Gera as notas de apoio à atividade ministerial a partir de `dados_notas.py`.

    python3 notas_ministeriais/gerar_nota.py --verificar          # só valida, não escreve nada
    python3 notas_ministeriais/gerar_nota.py rgac                 # nota temática (8 linhas)
    python3 notas_ministeriais/gerar_nota.py --mensal             # nota mensal do Gabinete
    python3 notas_ministeriais/gerar_nota.py --atualizacao        # atualização à Diretora-Geral
    python3 notas_ministeriais/gerar_nota.py --mensal --mes 2026-09

`--mes` fixa o mês a que a nota se refere. Sem ele usa-se o mês corrente, que
não é o mesmo: a nota da primeira semana de outubro reporta setembro.

Formato conferido contra a nota do Gabinete de 28.08.2026: A4, margens 3 cm
(superior 4,25), Segoe UI, título a 14 pt negrito, tabela «Table Grid» com
coluna de rótulos de 4,18 cm a 12 pt negrito e sombreado D9D9D9 nas quatro
linhas do modelo do Gabinete e D0CECE nas quatro acrescentadas pela DGAV.
"""
import os
import re
import sys
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
import dados_notas as D  # noqa: E402

PRETO = RGBColor(0, 0, 0)
FUNDO = {"modelo": "D9D9D9", "dgav": "D0CECE"}
FIM_DE_FRASE = ".:;!?»)”"

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]

DIAS_FICHA_VELHA = 40        # a ficha é mensal; mais do que isto está por rever
DIAS_SEM_VARRIMENTO = 45     # tempo máximo sem confirmar legislação nova


# ───────────────────────── validação ─────────────────────────

def _dias_desde(iso):
    try:
        a, m, d = (int(x) for x in iso.split("-"))
    except (ValueError, AttributeError):
        return None
    return (date.today() - date(a, m, d)).days


def validar(slug, ficha):
    """Devolve (erros, avisos). Com erros não se gera; com avisos gera-se e diz-se."""
    erros = []
    avisos = []
    campos = [c for c, _, _ in D.ROTULOS]

    for campo in campos:
        if campo not in ficha:
            erros.append("falta o campo «%s»" % campo)
            continue
        if not ficha[campo]:
            erros.append(
                "o campo «%s» está vazio. Na nota de 28.08.2026 foi precisamente este o campo "
                "deixado em branco, e é o que o Senhor Ministro mais precisa de levar para uma "
                "audição. Preencher antes de gerar." % campo)

    if not ficha.get("tema"):
        erros.append("falta o tema.")
    if not ficha.get("unidade"):
        avisos.append("a unidade orgânica está vazia; a nota sai sem a identificar.")

    texto = " ".join(p for c in campos if c in ficha for p in ficha[c])

    # designações: as exceções são nomes próprios onde a forma «errada» está certa
    neutro = texto
    for excecao in getattr(D, "EXCECOES_DESIGNACAO", []):
        neutro = neutro.replace(excecao, "")
    for errada, certa in D.DESIGNACOES_PROIBIDAS:
        if errada in neutro:
            erros.append("designação errada: «%s». Usar «%s»." % (errada, certa))
    for expressao, razao in getattr(D, "DESIGNACOES_A_CONFIRMAR", []):
        if expressao in neutro:
            avisos.append("«%s»: %s" % (expressao, razao))

    # frases truncadas: a nota de agosto terminava com «que foi publicado dia»
    for campo in campos:
        for i, par in enumerate(ficha.get(campo, [])):
            p = par.rstrip()
            if p and p[-1] not in FIM_DE_FRASE:
                erros.append("«%s», parágrafo %d: termina sem pontuação final — frase truncada? «…%s»"
                             % (campo, i + 1, p[-45:]))
            if re.search(r"\b(publicado|publicada|remetido|remetida|aprovado|aprovada|previsto|"
                         r"prevista|entregue)\s+(no dia|dia|em|a)\s*[.,;]?\s*$", p):
                erros.append("«%s», parágrafo %d: data em falta." % (campo, i + 1))
            if len(p) > 700:
                avisos.append("«%s», parágrafo %d: %d caracteres. O Gabinete lê em pé; partir."
                              % (campo, i + 1, len(p)))

    # a confusão que esta casa não pode cometer: entrada em vigor ≠ aplicação
    for frase in re.split(r"(?<=[.;])\s+", texto):
        # só interessa «em vigor» amarrado a uma data de 2028 ou posterior, sem
        # que a palavra «aplicável» apareça pelo meio a desfazer o equívoco
        if re.search(r"em vigor(?:(?!aplic)[^.;]){0,40}?20(?:2[89]|3\d)", frase):
            avisos.append("«…%s»: datas de 2028 ou posteriores são de aplicação, não de entrada em "
                          "vigor. Confirmar contra o artigo 33.º." % frase[-90:])
        if re.search(r"já (se aplica|é aplicável|vigora)", frase) and "2026/1818" in frase:
            avisos.append("«…%s»: a regra geral do Regulamento só é aplicável a 31 de agosto de "
                          "2028." % frase[-90:])

    # a ficha é um documento vivo: tem de dizer quando foi mexida e o que mudou
    dias = _dias_desde(ficha.get("atualizado"))
    if dias is None:
        erros.append("falta a data de atualização, ou está mal escrita. Usar AAAA-MM-DD.")
    elif dias > DIAS_FICHA_VELHA:
        avisos.append("a ficha não é revista há %d dias. Confirmar antes de a mandar." % dias)
    registo = ficha.get("registo") or []
    if not registo:
        erros.append("o registo está vazio: toda a alteração à ficha deixa uma linha.")
    elif ficha.get("atualizado") and ficha["atualizado"] not in [d for d, _ in registo]:
        erros.append("o registo não tem linha de %s, a data de atualização da ficha."
                     % ficha["atualizado"])

    # antes de mandar uma nota, confirma-se que não saiu legislação nova
    varrimento = ficha.get("legislacao_verificada")
    if not varrimento:
        avisos.append("a ficha não diz quando se confirmou pela última vez que não há legislação "
                      "nova. Ver a receita de varrimento do Diário da República na skill.")
    else:
        dv = _dias_desde(varrimento[0])
        if dv is None:
            erros.append("a data do varrimento legislativo está mal escrita. Usar AAAA-MM-DD.")
        elif dv > DIAS_SEM_VARRIMENTO:
            avisos.append("o varrimento legislativo tem %d dias. Repetir antes de mandar." % dv)
    return erros, avisos


def _relatar(nome, erros, avisos):
    for a in avisos:
        print("  aviso  [%s] %s" % (nome, a))
    for e in erros:
        print("  ERRO   [%s] %s" % (nome, e))
    return not erros


# ───────────────────────── construção ─────────────────────────

def _sombrear(celula, cor):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), cor)
    celula._tc.get_or_add_tcPr().append(sh)


def _escrever(celula, paragrafos, tamanho=11, alinhar=WD_ALIGN_PARAGRAPH.JUSTIFY,
              negrito=False, fonte=None):
    celula.text = ""
    primeiro = celula.paragraphs[0]
    for i, txt in enumerate(paragrafos):
        p = primeiro if i == 0 else celula.add_paragraph()
        if txt.startswith("- "):
            txt = txt[2:]
            p.style = celula.part.document.styles["List Bullet"]
        p.alignment = alinhar
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        r = p.add_run(txt)
        r.font.size = Pt(tamanho)
        r.font.bold = negrito
        r.font.color.rgb = PRETO
        if fonte:
            r.font.name = fonte


def _documento():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = s.bottom_margin = Cm(3)
    s.top_margin = Cm(4.25)
    est = doc.styles["Normal"]
    est.font.name = "Segoe UI"
    est.font.size = Pt(11)
    est.font.color.rgb = PRETO
    return doc


def _titulo(doc, texto, tamanho=14, antes=0, depois=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    r = p.add_run(texto)
    r.font.size = Pt(tamanho)
    r.font.bold = True
    r.font.color.rgb = PRETO
    return p


def _corpo(doc, paragrafos, marca=True):
    for txt in paragrafos:
        p = doc.add_paragraph(style="List Bullet" if marca else None)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(txt[2:] if txt.startswith("- ") else txt)
        r.font.size = Pt(11)
        r.font.color.rgb = PRETO


def _cabecalho(doc, ficha):
    _titulo(doc, "Notas de apoio à atividade ministerial")
    doc.add_paragraph()
    p = doc.add_paragraph()
    for texto in ("Tema: ", ficha["tema"]):
        r = p.add_run(texto)
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = PRETO
    if ficha.get("unidade"):
        p = doc.add_paragraph()
        r = p.add_run(ficha["unidade"])
        r.font.size = Pt(11)
        r.font.color.rgb = PRETO


def _fixar_grelha(tb, larguras_cm):
    """O python-docx não escreve a grelha; sem isto o Word reparte as colunas
    em partes iguais e a nota sai diferente do modelo do Gabinete."""
    tblPr = tb._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tb._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tb._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for cm in larguras_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(cm * 567)))  # 1 cm = 567 twips
        grid.append(col)
    tb._tbl.insert(list(tb._tbl).index(tblPr) + 1, grid)


def _tabela(doc, ficha, campos):
    linhas = [(c, rot, grupo) for c, rot, grupo in D.ROTULOS if c in campos]
    tb = doc.add_table(rows=len(linhas), cols=2)
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.LEFT
    tb.autofit = False
    _fixar_grelha(tb, (4.18, 10.82))
    for i, (campo, rotulo, grupo) in enumerate(linhas):
        lin = tb.rows[i]
        lin.cells[0].width = Cm(4.18)
        lin.cells[1].width = Cm(10.82)
        _escrever(lin.cells[0], rotulo.split("\n"), tamanho=12,
                  alinhar=WD_ALIGN_PARAGRAPH.LEFT, negrito=True, fonte="Aptos")
        _sombrear(lin.cells[0], FUNDO[grupo])
        _escrever(lin.cells[1], ficha[campo])
    return tb


# ───────────────────────── mês de referência ─────────────────────────

def _mes_referencia(arg):
    if arg:
        a, m = (int(x) for x in arg.split("-"))
        return a, m
    hoje = date.today()
    return hoje.year, hoje.month


def _nome_mes(ano, mes):
    return "%s de %d" % (MESES[mes - 1], ano)


def _evolucoes(ficha, ano, mes):
    """As evoluções do mês são as linhas do registo desse mês. É o que obriga a
    manter o registo: sem ele a atualização mensal não tem conteúdo."""
    prefixo = "%04d-%02d" % (ano, mes)
    return [t for d, t in (ficha.get("registo") or []) if d.startswith(prefixo)]


# ───────────────────────── saídas ─────────────────────────

def _validar_todas(slugs):
    ok = True
    for slug in slugs:
        if not _relatar(slug, *validar(slug, D.FICHAS[slug])):
            ok = False
    return ok


def nota_tematica(slug):
    ficha = D.FICHAS[slug]
    if not _validar_todas([slug]):
        print("NÃO GERADO.")
        return None
    doc = _documento()
    _cabecalho(doc, ficha)
    _tabela(doc, ficha, [c for c, _, _ in D.ROTULOS])
    out = os.path.join(BASE, "%s_%s_Notas_apoio_atividade_ministerial.docx"
                       % (date.today().strftime("%d.%m.%Y"), slug.upper()))
    doc.save(out)
    return out


def nota_mensal(mes=None):
    """Nota mensal do Gabinete: só os pontos 2 e 4 do modelo, com o subcapítulo
    obrigatório de todas as medidas já tomadas."""
    if not _validar_todas(list(D.FICHAS)):
        print("NÃO GERADO.")
        return None
    ano, m = _mes_referencia(mes)
    doc = _documento()
    _titulo(doc, "Notas de apoio à atividade ministerial")
    _titulo(doc, "Nota mensal de temas em destaque — %s" % _nome_mes(ano, m))
    doc.add_paragraph()

    for ficha in D.FICHAS.values():
        _titulo(doc, ficha["tema"], tamanho=12, antes=12, depois=6)
        _tabela(doc, ficha, ["pontos_sensiveis", "contexto"])
        _titulo(doc, "Medidas já tomadas", tamanho=11, antes=8)
        _corpo(doc, ficha["acoes_medidas"])

    out = os.path.join(BASE, "%s_Nota_mensal_temas_destaque.docx"
                       % date.today().strftime("%d.%m.%Y"))
    doc.save(out)
    return out


def atualizacao_mensal(mes=None):
    """Atualização à Diretora-Geral, até ao último dia útil do mês: evoluções
    relevantes, medidas adotadas, riscos ou constrangimentos e matérias com
    possível impacto político, mediático ou parlamentar."""
    if not _validar_todas(list(D.FICHAS)):
        print("NÃO GERADO.")
        return None
    ano, m = _mes_referencia(mes)
    doc = _documento()
    _titulo(doc, "Atualização mensal dos temas acompanhados")
    _titulo(doc, _nome_mes(ano, m).capitalize())
    doc.add_paragraph()

    for slug, ficha in D.FICHAS.items():
        _titulo(doc, ficha["tema"], tamanho=12, antes=12, depois=2)
        if ficha.get("unidade"):
            _corpo(doc, [ficha["unidade"]], marca=False)
        blocos = [
            ("Evoluções relevantes", _evolucoes(ficha, ano, m) or
             ["Sem evoluções registadas no período."]),
            ("Medidas adotadas", ficha["acoes_medidas"]),
            ("Riscos ou constrangimentos", ficha["riscos"]),
            ("Matérias com possível impacto político, mediático ou parlamentar", ficha["impacto"]),
        ]
        for rotulo, conteudo in blocos:
            _titulo(doc, rotulo, tamanho=11, antes=8, depois=2)
            _corpo(doc, conteudo)

    out = os.path.join(BASE, "%s_Atualizacao_mensal_DG.docx" % date.today().strftime("%d.%m.%Y"))
    doc.save(out)
    return out


def _ajuda():
    print(__doc__)
    print("Fichas disponíveis: %s" % ", ".join(D.FICHAS))


if __name__ == "__main__":
    args = sys.argv[1:]
    mes = None
    if "--mes" in args:
        i = args.index("--mes")
        try:
            mes = args[i + 1]
            _mes_referencia(mes)
        except (IndexError, ValueError):
            print("--mes pede AAAA-MM, por exemplo --mes 2026-09.")
            sys.exit(1)
        del args[i:i + 2]

    if "--verificar" in args:
        args.remove("--verificar")
        alvos = [a for a in args if not a.startswith("--")] or list(D.FICHAS)
        desconhecidas = [a for a in alvos if a not in D.FICHAS]
        if desconhecidas:
            print("ficha desconhecida: %s" % ", ".join(desconhecidas))
            sys.exit(1)
        print("A verificar %d ficha(s), sem escrever nada." % len(alvos))
        sys.exit(0 if _validar_todas(alvos) else 1)

    if "--mensal" in args:
        r = nota_mensal(mes)
    elif "--atualizacao" in args:
        r = atualizacao_mensal(mes)
    elif args and not args[0].startswith("--"):
        slug = args[0]
        if slug not in D.FICHAS:
            print("ficha desconhecida: %s. Disponíveis: %s" % (slug, ", ".join(D.FICHAS)))
            sys.exit(1)
        r = nota_tematica(slug)
    else:
        _ajuda()
        sys.exit(0)
    if r:
        print("gerado:", os.path.basename(r))
    else:
        sys.exit(1)
