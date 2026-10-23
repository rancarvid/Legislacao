# -*- coding: utf-8 -*-
"""Anexo ao parecer PT1225FL: as três vias para parar a atividade (art.º 3.º-G do DL 276/2001,
art.º 89.º do CPA e art.º 48.º do RJCE) e a razão de o RJCE se aplicar às contraordenações do DL 276/2001.
Gera Anexo_Vias_Cautelares_PT1225FL.docx. Aptos 11, preto, sem travessões fora das citações."""
import os, re, zipfile, shutil
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Anexo_Vias_Cautelares_PT1225FL.docx")

T = [
    ("t", "Anexo ao parecer sobre o alojamento PT1225FL. Medidas cautelares: art.º 3.º-G do DL 276/2001, art.º 89.º do CPA e art.º 48.º do RJCE"),
    ("m", "Nota de acompanhamento, 21.10.2026. Regista a resposta à dúvida sobre a origem do art.º 48.º citado no parecer e sobre a sua relação com o DL 276/2001."),

    ("s", "1. A dúvida"),
    ("p", "O parecer cita o n.º 1 do art.º 48.º do RJCE como base para medidas cautelares sem audição prévia. A dúvida colocada foi se esse artigo faz parte do DL 276/2001."),

    ("s", "2. Resposta"),
    ("p", "Não faz. O art.º 48.º pertence ao Regime Jurídico das Contraordenações Económicas (RJCE), aprovado em anexo ao Decreto-Lei n.º 9/2021, de 29 de janeiro. Aplica-se às contraordenações do DL 276/2001 por remissão: o art.º 36.º do DL 9/2021 deu nova redação ao art.º 68.º do DL 276/2001, que passou a qualificar as suas contraordenações como «contraordenação económica grave, punível nos termos do Regime Jurídico das Contraordenações Económicas (RJCE)» (n.º 1) e «contraordenação económica muito grave, punível nos termos do RJCE» (n.º 2). A partir daí, tudo o que o DL 276/2001 não regula no plano contraordenacional vai buscar-se ao RJCE: medidas cautelares (art.º 48.º), apreensão de bens e seres vivos (art.º 49.º), destino antecipado dos bens apreendidos (art.º 50.º), sanções acessórias gerais (art.º 28.º), prescrição (art.º 36.º) e concurso com crime (art.º 27.º)."),
    ("p", "O DL 276/2001, por si, não contém qualquer medida cautelar nem norma de apreensão. A verificação no texto consolidado do DRE (versão à data de 29.1.2021, sem alterações posteriores) mostrou que nenhum artigo usa as expressões «cautelar», «apreensão» ou «sequestro». O que o diploma tem é: o art.º 3.º-G (suspensão de atividade e encerramento dos alojamentos, que é procedimento administrativo e não contraordenacional); o n.º 8 do art.º 19.º (recolha de animais quando esteja em causa a sua saúde e bem-estar); o art.º 67.º-A (acesso ao alojamento, com mandado judicial em caso de recusa); e a al. a) do art.º 69.º (perda dos animais a favor do Estado, como sanção acessória)."),
    ("q", "«1 — Quando esteja em causa o exercício ilegal de atividades, a existência de risco grave ou iminente para a saúde e a segurança das pessoas, animais ou bens ou para o ambiente, a violação grave de direitos dos consumidores ou quando se revele necessário para a instrução do processo de contraordenação, a autoridade administrativa competente pode determinar, sem audição prévia dos interessados, as seguintes medidas cautelares: a) A suspensão, total ou parcial, do exercício de atividades económicas; b) A cessação de práticas proibidas nos termos das disposições legais aplicáveis que se encontrem em curso; c) O encerramento, no todo ou em parte, de estabelecimentos ou de espaços que se encontrem adstritos ao apoio às atividades exercidas; d) A apreensão ou selagem de todos ou parte dos bens; e) A imposição das medidas que se mostrem adequadas à prevenção de danos à saúde e à segurança das pessoas, animais ou bens ou de danos ambientais, à reposição da situação anterior à infração e à minimização dos efeitos decorrentes da mesma.»"),
    ("r", "n.º 1 do art.º 48.º do RJCE, anexo ao Decreto-Lei n.º 9/2021, de 29 de janeiro."),
    ("q", "«1 - Compete à DGAV e aos órgãos de polícia criminal a instrução dos processos de contraordenação. 2 - Compete ao diretor-geral de Alimentação e Veterinária ou ao diretor do respetivo órgão de polícia criminal a aplicação das coimas e das sanções acessórias.»"),
    ("r", "art.º 70.º do Decreto-Lei n.º 276/2001, de 17 de outubro."),

    ("s", "3. Três vias para parar a atividade"),
    ("tb", [
        ["Via", "Base legal", "Natureza", "Quem decide", "Audiência prévia"],
        ["Suspensão ou encerramento do alojamento", "art.º 3.º-G do DL 276/2001", "Procedimento administrativo, regido pelo CPA; medida de polícia", "Diretor-geral, sob proposta da direção de serviços da região (n.º 2)", "Sim, salvo dispensa por urgência (al. a) do n.º 1 do art.º 124.º do CPA)"],
        ["Medida provisória", "art.º 89.º do CPA, dentro do procedimento do art.º 3.º-G", "Cautelar, acessória do procedimento; caduca com a decisão final ou nos termos do art.º 90.º", "Órgão competente para a decisão final: diretor-geral", "Não (n.º 2 do art.º 89.º); exige fundamentação e prazo"],
        ["Medida cautelar contraordenacional", "art.º 48.º do RJCE, dentro de um processo de contraordenação", "Cautelar, acessória do processo sancionatório; recurso sem efeito suspensivo (n.º 3 do art.º 52.º)", "Autoridade administrativa do processo: diretor-geral ou diretor do OPC instrutor (art.º 70.º do DL 276/2001)", "Não (n.º 1 do art.º 48.º)"],
    ]),

    ("s", "4. Por que razão o parecer prefere o art.º 89.º do CPA"),
    ("p", "A medida do art.º 48.º do RJCE só existe dentro de um processo de contraordenação já instaurado e tem de assentar em factos que constituam contraordenação autónoma. Os factos relativos aos maus-tratos e às condições de alojamento estão absorvidos pelo crime (art.º 27.º do RJCE: «Se o mesmo facto constituir simultaneamente crime e contraordenação económica o agente é punido a título de crime, sem prejuízo das sanções acessórias previstas para a contraordenação»). A medida teria de apoiar-se nos factos não absorvidos, sobretudo o «exercício ilegal de atividades» no Complexo B, que nunca teve comunicação prévia (al. a) do n.º 1 do art.º 68.º do DL 276/2001). É possível e é um fundamento limpo."),
    ("p", "Mas a medida do RJCE segue a sorte do processo de contraordenação, não se liga ao procedimento de encerramento e pressupõe que esse processo esteja formalmente aberto, o que ainda não aconteceu. A medida provisória do art.º 89.º do CPA nasce dentro do próprio procedimento de encerramento, caduca com a decisão final e não depende de mais nada. Por isso é a primeira escolha."),
    ("p", "O art.º 48.º do RJCE fica como alternativa ou complemento, a usar quando o processo de contraordenação pelos factos não absorvidos for instruído: por exemplo, para selar o Complexo B (al. d)) ou para impor a cessação da atividade nesse núcleo (al. b))."),

    ("s", "5. Síntese para a pasta de acompanhamento"),
    ("b", "O RJCE aplica-se às contraordenações do DL 276/2001 desde 28.7.2021, por força da redação dada ao art.º 68.º pelo art.º 36.º do DL 9/2021."),
    ("b", "O DL 276/2001 não tem medidas cautelares próprias; tem o art.º 3.º-G (administrativo), o n.º 8 do art.º 19.º (recolha), o art.º 67.º-A (acesso) e a al. a) do art.º 69.º (perda)."),
    ("b", "Para parar a atividade antes da decisão de encerramento, a via recomendada é a medida provisória do art.º 89.º do CPA; a via do art.º 48.º do RJCE exige processo de contraordenação aberto e factos não absorvidos pelo crime."),
]

prosa = " ".join(t for k, t in T if k in ("t", "m", "s", "p", "b") and isinstance(t, str))
prosa = re.sub(r"«[^»]*»", "«»", prosa)
assert "—" not in prosa and "–" not in prosa, "travessão fora de citação"

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.2)
st = d.styles["Normal"]; st.font.name = "Aptos"; st.font.size = Pt(11); st.font.color.rgb = RGBColor(0, 0, 0)
rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
if rf is None:
    rf = OxmlElement("w:rFonts"); rpr.append(rf)
for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
    rf.set(qn(a), "Aptos")
st.paragraph_format.space_before = Pt(0); st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.0


def add(txt, size=11, bold=False, italic=False, align=None, before=0, after=6, indent=0):
    p = d.add_paragraph(); r = p.add_run(txt)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    if align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


for kind, txt in T:
    if kind == "t":
        add(txt, 14, bold=True, after=4)
    elif kind == "m":
        add(txt, 10, after=8)
    elif kind == "s":
        add(txt, 11, bold=True, before=10, after=4)
    elif kind == "p":
        add(txt, align="j")
    elif kind == "b":
        p = add(txt, align="j", indent=0.6); p.paragraph_format.first_line_indent = Cm(-0.4)
    elif kind == "q":
        add(txt, 10, align="j", indent=1.0, after=2)
    elif kind == "r":
        add(txt, 9, italic=True, indent=1.0, after=8)
    elif kind == "tb":
        tbl = d.add_table(rows=len(txt), cols=len(txt[0])); tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(txt):
            for j, val in enumerate(row):
                c = tbl.cell(i, j); c.text = ""
                r = c.paragraphs[0].add_run(val); r.font.size = Pt(9.5); r.bold = (i == 0)
                c.paragraphs[0].paragraph_format.space_after = Pt(2)
        d.add_paragraph().paragraph_format.space_after = Pt(4)

d.save(OUT)
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if it.filename == "word/settings.xml":
            data = re.sub(r'<w:zoom(?![^>]*w:percent)([^>]*)/>', r'<w:zoom w:percent="100"\1/>', data.decode()).encode()
        zo.writestr(it, data)
shutil.move(tmp, OUT)
print(os.path.basename(OUT), "gerado")
