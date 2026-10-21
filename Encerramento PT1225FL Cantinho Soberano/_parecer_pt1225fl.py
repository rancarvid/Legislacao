# -*- coding: utf-8 -*-
"""Gera o parecer sobre a proposta de encerramento do alojamento PT1225FL (Cantinho Soberano, Amarante).

Produz Parecer_Encerramento_PT1225FL.docx (Aptos 11, preto, sem negrito no corpo, sem travessões
fora das citações) e Parecer_Encerramento_PT1225FL.md com o mesmo conteúdo.

Tipos de bloco:
  t  título            m  linha de identificação       h  capítulo
  s  secção            p  parágrafo                    b  ponto de lista
  q  citação verbatim  r  referência da citação        n  observação (opinião)
  tb tabela: ('tb', [cabeçalho, linha, linha, ...])
"""
import re, zipfile, shutil, sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DOCX = os.path.join(HERE, "Parecer_Encerramento_PT1225FL.docx")
OUT_MD = os.path.join(HERE, "Parecer_Encerramento_PT1225FL.md")

from _conteudo_pt1225fl import T  # noqa: E402

# ------------------------------------------------------------------ verificações
prosa = " ".join(t for k, t in T if k in ("t", "m", "h", "s", "p", "b", "n") and isinstance(t, str))
for ch, nome in (("—", "travessão longo"), ("–", "travessão curto")):
    if ch in prosa:
        i = prosa.index(ch)
        raise SystemExit(f"{nome} fora de citação: ...{prosa[max(0, i-80):i+80]}...")

# ------------------------------------------------------------------ Word
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.2)
st = d.styles["Normal"]
st.font.name = "Aptos"; st.font.size = Pt(11); st.font.color.rgb = RGBColor(0, 0, 0)
rf = st.element.get_or_add_rPr().find(qn("w:rFonts"))
if rf is None:
    rf = OxmlElement("w:rFonts"); st.element.get_or_add_rPr().append(rf)
for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
    rf.set(qn(a), "Aptos")
pf = st.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(6); pf.line_spacing = 1.0


def add(txt, size=11, bold=False, italic=False, align=None, before=0, after=6, indent=0):
    p = d.add_paragraph()
    r = p.add_run(txt)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    if align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def table(rows):
    tbl = d.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = ""
            r = c.paragraphs[0].add_run(val)
            r.font.size = Pt(9.5)
            r.bold = (i == 0)
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
    d.add_paragraph().paragraph_format.space_after = Pt(4)


for kind, txt in T:
    if kind == "t":
        add(txt, 14, bold=True, after=4)
    elif kind == "m":
        add(txt, 10, after=2)
    elif kind == "h":
        add(txt, 12, bold=True, before=14, after=6)
    elif kind == "s":
        add(txt, 11, bold=True, before=8, after=4)
    elif kind == "p":
        add(txt, align="j")
    elif kind == "b":
        p = add(txt, align="j", indent=0.6)
        p.paragraph_format.first_line_indent = Cm(-0.4)
    elif kind == "q":
        add(txt, 10, align="j", indent=1.0, after=2)
    elif kind == "r":
        add(txt, 9, italic=True, indent=1.0, after=8)
    elif kind == "n":
        add("Observação. " + txt, align="j", indent=0.6, after=8)
    elif kind == "tb":
        table(txt)
    else:
        raise SystemExit(f"tipo desconhecido: {kind}")

d.save(OUT_DOCX)
tmp = OUT_DOCX + ".tmp"
with zipfile.ZipFile(OUT_DOCX) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if it.filename == "word/settings.xml":
            data = re.sub(r'<w:zoom(?![^>]*w:percent)([^>]*)/>', r'<w:zoom w:percent="100"\1/>',
                          data.decode()).encode()
        zo.writestr(it, data)
shutil.move(tmp, OUT_DOCX)

# ------------------------------------------------------------------ Markdown
md = []
for kind, txt in T:
    if kind == "t":
        md.append(f"# {txt}\n")
    elif kind == "m":
        md.append(f"{txt}  ")
    elif kind == "h":
        md.append(f"\n## {txt}\n")
    elif kind == "s":
        md.append(f"\n### {txt}\n")
    elif kind == "p":
        md.append(f"{txt}\n")
    elif kind == "b":
        md.append(f"- {txt}")
    elif kind == "q":
        md.append(f"> {txt}")
    elif kind == "r":
        md.append(f"> *{txt}*\n")
    elif kind == "n":
        md.append(f"> **Observação.** {txt}\n")
    elif kind == "tb":
        hdr = txt[0]
        md.append("| " + " | ".join(hdr) + " |")
        md.append("|" + "---|" * len(hdr))
        for row in txt[1:]:
            md.append("| " + " | ".join(c.replace("|", "/") for c in row) + " |")
        md.append("")
with open(OUT_MD, "w", encoding="utf-8") as f:
    f.write("\n".join(md) + "\n")

print(os.path.basename(OUT_DOCX), "e", os.path.basename(OUT_MD), "gerados;", len(prosa.split()), "palavras de prosa")
