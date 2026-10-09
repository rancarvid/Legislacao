# -*- coding: utf-8 -*-
"""Quadro de casos-tipo — limites de detenção por fogo e lotação dos alojamentos.

Documento autónomo, em paisagem. As células trazem referências curtas; a prosa e
os textos legais conferidos no Diário da República estão no anexo
«Casos_Tipo_Limites_por_Fogo_2026-10-09_Anexo.docx», gerado por `_anexo_casos_tipo.py`.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = '/home/user/Legislacao/criador informal/Casos_Tipo_Limites_por_Fogo_2026-10-09.docx'

TINTA = RGBColor(0x1A, 0x1A, 0x1A)
CINZA = RGBColor(0x59, 0x59, 0x59)
CAB = RGBColor(0x1F, 0x3B, 0x5C)

COLUNAS = ['Caso', 'Limite de animais adultos', 'Quem autoriza', 'Sanção e instrução',
           'Remoção e encerramento', 'Urbanismo e vizinhança']
LARGURAS = [4.2, 4.6, 4.0, 4.4, 4.0, 4.4]

# Blocos repetidos. «Autorização municipal» e «instrução pelas DSAVR da DGAV»
# substituem as formulações provisórias do anexo, por decisão do utilizador.
AUT_6 = ('Até quatro, nada. Até seis, pedido do detentor e pareceres vinculativos do MVM '
         'e do delegado de saúde: autorização municipal ¹. Acima de seis, no fogo, não há')
SANC_314 = 'Al. c) do n.º 3 do art. 14.º do DL 314/2003 ⁴. Instrução: DSAVR. Coima: DGAV'
REM_314 = 'Câmara, após vistoria conjunta ⁵ (n.º 5). Mandado judicial (n.º 6)'
SANC_276 = ('Al. f) do n.º 1 do art. 68.º do DL 276/2001 ⁶. Instrução: DGAV e OPC. '
            'Coima: DGAV ou OPC (art. 70.º)')
ENC_3G = 'Despacho do diretor-geral (n.º 1 do art. 3.º-G). Execução e recolha: câmaras (n.º 6)'
REC_19 = 'Recolha pela DGAV com as câmaras e pelas polícias (n.º 8 do art. 19.º)'

LINHAS = [
    ['1. Fração autónoma em propriedade horizontal. Animais do agregado, incluindo '
     'varanda e terraço',
     'N.º 2: três cães ou quatro gatos, máximo de quatro; até seis com autorização. '
     'O condomínio pode fixar menos (n.º 3)',
     AUT_6 + '. Acrescem título constitutivo e deliberações',
     SANC_314,
     REM_314,
     'Uso diverso do fim e atos proibidos (als. c) e d) do n.º 2 do art. 1422.º CC). '
     'Ruído de vizinhança ²'],

    ['2. Andar em prédio sem propriedade horizontal. Animais do agregado',
     'N.º 2, como acima. O n.º 3 não se aplica',
     AUT_6,
     SANC_314,
     REM_314,
     'Resolução do arrendamento por higiene, sossego e boa vizinhança (al. a) do n.º 2 '
     'do art. 1083.º CC). Ruído de vizinhança ²'],

    ['3. Moradia, com ou sem logradouro, qualquer que seja a dimensão. Animais do '
     'agregado',
     'N.º 2, como acima. A área do logradouro não altera o limite',
     AUT_6,
     SANC_314 + '. Alcança o logradouro, como terreno anexo',
     REM_314,
     'Anexos para animais até 1/15 do logradouro, quando autorizados, podendo a câmara '
     'interditá-los em zonas de aglomeração (art. 115.º e § único do RGEU). Ruído de '
     'vizinhança ²'],

    ['4. Moradia com alojamento registado em instalações do art. 25.º, no logradouro',
     'No alojamento: capacidade declarada (al. h) do n.º 1 do art. 3.º-A), limitada pelo '
     'anexo III. Dentro de casa: n.º 2',
     'Mera comunicação prévia à DGAV (art. 3.º-A), sem vistoria prévia. Obras: RJUE',
     SANC_276 + '. Animais de casa: DL 314/2003',
     ENC_3G + '. Animais de casa: n.º 5',
     'PDM e título de utilização. Anexos até 1/15 do logradouro (art. 115.º RGEU). '
     'Atividade ruidosa permanente ²'],

    ['5. Fração autónoma ou andar com alojamento registado ³',
     'Fração habitacional: a hospedagem sem fins lucrativos está excluída das frações '
     '(al. p) do n.º 1 do art. 2.º), logo n.os 2 e 3. Fração comercial ou de serviços: '
     'não há fogo, vale a capacidade declarada',
     'Mera comunicação prévia à DGAV. Na fração habitacional, instalar o alojamento é '
     'alteração de uso',
     SANC_276,
     ENC_3G,
     'Uso diverso do fim proibido; silente o título, a alteração exige dois terços do '
     'valor do prédio (al. c) do n.º 2 e n.º 4 do art. 1422.º CC). Título de utilização. '
     'Ruído ²'],

    ['6. Prédio urbano sem fogo — loja, armazém, serviços — com alojamento registado',
     'O n.º 2 não tem unidade a que se aplicar. Só a capacidade declarada, limitada pelo '
     'anexo III',
     'Mera comunicação prévia à DGAV. Título de utilização compatível',
     SANC_276,
     ENC_3G,
     'PDM e título de utilização (RJUE). Atividade ruidosa permanente ²'],

    ['7. Prédio urbano com alojamento de facto, sem registo',
     'Dentro de casa: n.º 2. Nas instalações: sem capacidade declarada e atividade sem '
     'título',
     'Ninguém: falta a mera comunicação prévia',
     'Al. a) do n.º 1 do art. 68.º pela falta de título, podendo acrescer a al. f); e a '
     'al. c) do n.º 3 do art. 14.º do DL 314/2003 ⁴',
     REM_314 + '. ' + REC_19 + '. Art. 3.º-G sem registo: discutível *',
     'Operações urbanísticas não licenciadas (RJUE). Ruído ²'],

    ['8. Prédio misto. Animais no fogo da parte urbana',
     'N.º 2 quanto ao fogo; n.º 4 quanto ao prédio',
     AUT_6,
     SANC_314,
     REM_314,
     'Ruído de vizinhança ²'],

    ['9. Prédio misto. Animais no logradouro ou na parte rústica',
     'N.º 4: seis, excedíveis se a dimensão do terreno o permitir, sem teto fixo',
     'Sem autorização prévia. A dimensão do terreno é apreciada em caso de controlo',
     'Al. c) do n.º 3 do art. 14.º ⁴ no terreno anexo à habitação; na parte rústica '
     'afastada, cobertura duvidosa. Instrução: DSAVR',
     REM_314,
     'PDM. Ruído de vizinhança ou atividade ruidosa permanente, conforme o caso ²'],

    ['10. Prédio misto com alojamento registado',
     'No alojamento: anexo III. No fogo: n.º 2. Restantes espaços: n.º 4',
     'Mera comunicação prévia à DGAV. Obras: RJUE',
     SANC_276,
     ENC_3G,
     'PDM, conforme o solo seja rústico ou urbano. Atividade ruidosa permanente ²'],

    ['11. Prédio rústico sem habitação nem atividade',
     'N.º 4: seis, excedíveis se a dimensão do terreno o permitir',
     'Sem autorização prévia',
     'Sem tipo contraordenacional no DL 314/2003: não há habitação nem terreno anexo a ela',
     REM_314 + '. É o único remédio',
     'PDM. Ruído ²'],

    ['12. Prédio rústico com alojamento registado',
     'Capacidade declarada, limitada pelo anexo III',
     'Mera comunicação prévia à DGAV. Edificação sujeita ao regime do solo e ao RJUE',
     SANC_276,
     ENC_3G,
     'Regime do solo rústico (PDM). Atividade ruidosa permanente ²'],

    ['13. Prédio rústico ou misto com alojamento de facto, sem registo',
     'N.º 4 quanto ao prédio. Atividade sem título',
     'Ninguém: falta a mera comunicação prévia',
     'Als. a) e f) do n.º 1 do art. 68.º do DL 276/2001',
     REM_314 + '. ' + REC_19 + '. Art. 3.º-G sem registo: discutível *',
     'Operações urbanísticas não licenciadas (RJUE). Ruído ²'],
]

CABECALHO = ('Os casos organizam-se pelo lugar onde estão os animais, que é o critério que '
             'decide a norma aplicável, e não pela classificação matricial do prédio nem '
             'pela dimensão do terreno. Na coluna do limite e nas colunas da sanção e da '
             'remoção quanto aos animais do agregado, os artigos sem diploma são do '
             'Decreto-Lei n.º 314/2003; quanto aos alojamentos registados, do Decreto-Lei '
             'n.º 276/2001. MVM: médico veterinário municipal. DSAVR: direções de serviços '
             'de alimentação e veterinária regionais da DGAV. OPC: órgãos de polícia '
             'criminal. 9 de outubro de 2026.')

NOTAS = [
    'O n.º 1 do artigo 3.º do Decreto-Lei n.º 314/2003 — alojamento condicionado a boas '
    'condições e à ausência de riscos hígio-sanitários — aplica-se em todos os casos, com '
    'ou sem registo, e não se repete no quadro. Contam-se apenas os animais adultos, com '
    'um ano ou mais (als. f) e g) do art. 2.º). O registo de um alojamento não altera a '
    'contagem dos animais que vivem dentro do fogo.',

    '¹ A autorização é municipal: o n.º 2 do artigo 3.º não nomeia o órgão, mas o n.º 5 '
    'comete a execução do artigo às câmaras municipais e os pareceres vinculativos são do '
    'médico veterinário municipal e do delegado de saúde. Qual o órgão do município — '
    'câmara ou presidente — depende do regulamento de cada município; o Regulamento '
    'n.º 181/2025 do Município do Cartaxo prevê requerimento ao presidente da câmara e '
    'vistoria conjunta (art. 13.º).',

    '² Regulamento Geral do Ruído, aprovado pelo Decreto-Lei n.º 9/2007, de 17 de janeiro, '
    'na redação do Decreto-Lei n.º 278/2007, de 1 de agosto, que alterou apenas o '
    'artigo 4.º daquele decreto-lei e o artigo 15.º do Regulamento — nenhuma das normas '
    'aqui invocadas. Os animais do agregado produzem ruído de vizinhança, fiscalizado '
    'pelas autoridades policiais, que podem ordenar a cessação imediata entre as 23 e as '
    '7 horas (art. 24.º). O alojamento é atividade ruidosa permanente, sujeita a valores '
    'limite e ao critério de incomodidade (art. 13.º) e fiscalizada pela entidade '
    'licenciadora e pelas câmaras municipais (art. 26.º).',

    '³ O caso 5 não foi analisado no estudo. A construção é nova, feita a partir da '
    'alínea p) do n.º 1 do artigo 2.º do Decreto-Lei n.º 276/2001 e do artigo 1422.º do '
    'Código Civil, e distingue a fração destinada a habitação da fração destinada a '
    'comércio ou serviços.',

    '⁴ A alínea pune «a permanência de cães e gatos em habitações e terrenos anexos em '
    'desrespeito pelas condições previstas no artigo 3.º». Coima de 50 a 3740 euros, ou '
    'até 44 890 euros se o agente for pessoa coletiva (n.º 3 do art. 14.º).',

    '⁵ Vistoria conjunta do delegado de saúde e do médico veterinário municipal. A câmara '
    'notifica o detentor para retirar os animais para o canil ou gatil municipal no prazo '
    'que aquelas entidades fixarem, caso o detentor não opte por outro destino que reúna '
    'as condições legais. O mandado judicial é pedido pelo presidente da câmara.',

    '⁶ Contraordenação económica grave, punível nos termos do regime jurídico das '
    'contraordenações económicas.',

    '* Ponto por confirmar: se o artigo 3.º-G do Decreto-Lei n.º 276/2001, que permite ao '
    'diretor-geral suspender a atividade ou encerrar o alojamento, se aplica a um '
    'alojamento que nunca foi registado.',

    'A justificação de cada campo, com os textos legais conferidos no Diário da República, '
    'consta do anexo «Casos-tipo e análise», de 9 de outubro de 2026.',
]


def _unico(pai, tag):
    """Devolve o elemento, apagando duplicados — o python-docx já cria alguns."""
    achados = pai.findall(qn(tag))
    for extra in achados[1:]:
        pai.remove(extra)
    if achados:
        return achados[0]
    novo = OxmlElement(tag)
    pai.append(novo)
    return novo


def _grelha(tab, larguras_cm):
    """Fixa o layout e as larguras na grelha que já existe.

    Inserir um segundo w:tblGrid deixa o Word e o LibreOffice a ler o primeiro,
    com as larguras do autofit, e a última coluna absorve o resto da linha.
    """
    pr = tab._tbl.tblPr
    layout = _unico(pr, 'w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    largura = _unico(pr, 'w:tblW')
    largura.set(qn('w:type'), 'dxa')
    largura.set(qn('w:w'), str(int(sum(larguras_cm) * 567)))

    grid = _unico(tab._tbl, 'w:tblGrid')
    for col in list(grid):
        grid.remove(col)
    for cm in larguras_cm:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(int(cm * 567)))
        grid.append(col)


def _sombrear(celula, cor):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), cor)
    celula._tc.get_or_add_tcPr().append(sh)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(1.4)

n = doc.styles['Normal']
n.font.name = 'Calibri'
n.font.size = Pt(10)
n.font.color.rgb = TINTA
n.paragraph_format.space_after = Pt(6)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
r = p.add_run('Limites de detenção por fogo e lotação dos alojamentos: quadro de casos-tipo')
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = CAB

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r = p.add_run(CABECALHO)
r.italic = True
r.font.size = Pt(8)
r.font.color.rgb = CINZA

tab = doc.add_table(rows=1, cols=len(COLUNAS))
tab.style = 'Table Grid'
tab.alignment = WD_TABLE_ALIGNMENT.CENTER
tab.autofit = False
_grelha(tab, LARGURAS)

_trpr = tab.rows[0]._tr.get_or_add_trPr()
_cab = OxmlElement('w:tblHeader')   # repete o cabeçalho em cada página
_cab.set(qn('w:val'), 'true')
_trpr.append(_cab)

for i, (celula, texto) in enumerate(zip(tab.rows[0].cells, COLUNAS)):
    celula.width = Cm(LARGURAS[i])
    _sombrear(celula, 'E8EDF3')
    par = celula.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run(texto)
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = CAB

for linha in LINHAS:
    cels = tab.add_row().cells
    for i, (celula, texto) in enumerate(zip(cels, linha)):
        celula.width = Cm(LARGURAS[i])
        par = celula.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = par.add_run(texto)
        run.font.size = Pt(7)
        if i == 0:
            run.bold = True

doc.add_paragraph()
for i, nota in enumerate(NOTAS):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(3)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = par.add_run(nota)
    run.font.size = Pt(7.5)
    run.font.color.rgb = CINZA if i else TINTA

doc.save(OUT)
print('gravado:', OUT)
print('linhas:', len(LINHAS), '| largura total:', round(sum(LARGURAS), 1), 'cm de',
      round(29.7 - 3.6, 1), 'cm úteis')
