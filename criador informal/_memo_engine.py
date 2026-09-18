# -*- coding: utf-8 -*-
"""Motor de composicao do memorando (Word). Paleta pastel, tipografia Aptos."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT      = 'Aptos'
INK       = '2B2B2B'
HEAD      = '35506B'
SUB       = '5C7891'
MUTE      = '7C8894'
RULE      = 'D7DEE5'

CIT_BG, CIT_BAR   = 'F3F6F9', 'A9C2D8'
KEY_BG, KEY_BAR   = 'F7F4ED', 'D8C9AE'
NOTE_BG, NOTE_BAR = 'FAF2F1', 'E0BEBA'
TBL_HEAD          = 'E9EFF4'
TBL_ALT           = 'F8FAFC'
OK_BG             = 'EDF3ED'
NO_BG             = 'FAF0EF'


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), v)
    return e



# ------------------------------------------------- insercao respeitando o esquema
_ORD = {
 'rPr': ['w:rStyle','w:rFonts','w:b','w:bCs','w:i','w:iCs','w:caps','w:smallCaps','w:strike',
         'w:dstrike','w:outline','w:shadow','w:emboss','w:imprint','w:noProof','w:snapToGrid',
         'w:vanish','w:webHidden','w:color','w:spacing','w:w','w:kern','w:position','w:sz','w:szCs',
         'w:highlight','w:u','w:effect','w:bdr','w:shd','w:fitText','w:vertAlign','w:rtl','w:cs',
         'w:em','w:lang','w:eastAsianLayout','w:specVanish','w:oMath'],
 'pPr': ['w:pStyle','w:keepNext','w:keepLines','w:pageBreakBefore','w:framePr','w:widowControl',
         'w:numPr','w:suppressLineNumbers','w:pBdr','w:shd','w:tabs','w:suppressAutoHyphens',
         'w:kinsoku','w:wordWrap','w:overflowPunct','w:topLinePunct','w:autoSpaceDE','w:autoSpaceDN',
         'w:bidi','w:adjustRightInd','w:snapToGrid','w:spacing','w:ind','w:contextualSpacing',
         'w:mirrorIndents','w:suppressOverlap','w:jc','w:textDirection','w:textAlignment',
         'w:textboxTightWrap','w:outlineLvl','w:divId','w:cnfStyle','w:rPr','w:sectPr','w:pPrChange'],
 'tcPr': ['w:cnfStyle','w:tcW','w:gridSpan','w:hMerge','w:vMerge','w:tcBorders','w:shd','w:noWrap',
          'w:tcMar','w:textDirection','w:tcFitText','w:vAlign','w:hideMark'],
 'trPr': ['w:cnfStyle','w:divId','w:gridBefore','w:gridAfter','w:wBefore','w:wAfter',
          'w:cantSplit','w:trHeight','w:tblHeader','w:tblCellSpacing','w:jc','w:hidden'],
 'tblPr': ['w:tblStyle','w:tblpPr','w:tblOverlap','w:bidiVisual','w:tblStyleRowBandSize',
           'w:tblStyleColBandSize','w:tblW','w:jc','w:tblCellSpacing','w:tblInd','w:tblBorders',
           'w:shd','w:tblLayout','w:tblCellMar','w:tblLook','w:tblCaption','w:tblDescription'],
}


def put(parent, child, kind):
    """Insere child em parent na posicao exigida pelo esquema WordprocessingML."""
    tag = child.tag.split('}')[-1]
    order = _ORD[kind]
    name = 'w:' + tag
    idx = order.index(name) if name in order else len(order)
    existing = parent.find(qn(name))
    if existing is not None:
        parent.remove(existing)
    for el in parent:
        t = 'w:' + el.tag.split('}')[-1]
        j = order.index(t) if t in order else len(order)
        if j > idx:
            el.addprevious(child)
            return child
    parent.append(child)
    return child


def font(run, size=10, color=INK, bold=False, italic=False, name=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)
    return run


def spacing(p, before=0, after=6, line=1.22):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    return p


def cell_bg(cell, hexv):
    put(cell._tc.get_or_add_tcPr(), _el('w:shd', val='clear', color='auto', fill=hexv), 'tcPr')


def cell_borders(cell, spec):
    """spec: {'left': (sz_eighths, color), ...}; ausente = nil"""
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side)
        if side in spec:
            sz, col = spec[side]
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
            e.set(qn('w:space'), '0'); e.set(qn('w:color'), col)
        else:
            e.set(qn('w:val'), 'nil')
        b.append(e)
    put(tcpr, b, 'tcPr')


def cell_pad(cell, t=110, b=110, l=150, r=150):
    tcpr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for side, v in (('top', t), ('bottom', b), ('left', l), ('right', r)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa')
        m.append(e)
    put(tcpr, m, 'tcPr')


def no_table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side); e.set(qn('w:val'), 'nil')
        b.append(e)
    put(tblPr, b, 'tblPr')


def row_no_split(row):
    trPr = row._tr.get_or_add_trPr()
    put(trPr, OxmlElement('w:cantSplit'), 'trPr')


def row_repeat(row):
    trPr = row._tr.get_or_add_trPr()
    put(trPr, OxmlElement('w:cantSplit'), 'trPr')
    put(trPr, OxmlElement('w:tblHeader'), 'trPr')


def keep_together(p):
    p.paragraph_format.keep_with_next = True


def rule(doc, color=RULE, sz=6, space_after=10):
    p = doc.add_paragraph()
    spacing(p, 0, space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(2)
    ppr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr')
    e = OxmlElement('w:bottom')
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
    e.set(qn('w:space'), '1'); e.set(qn('w:color'), color)
    b.append(e); put(ppr, b, 'pPr')
    return p


# ---------------------------------------------------------------- blocos

def h1(doc, n, txt):
    p = doc.add_paragraph(); spacing(p, 18, 2); keep_together(p)
    font(p.add_run(f'{n}  '), 13.5, SUB, bold=True)
    font(p.add_run(txt), 13.5, HEAD, bold=True)
    r = rule(doc, RULE, 4, 8)
    keep_together(r)
    return p


def h2(doc, n, txt):
    p = doc.add_paragraph(); spacing(p, 13, 4); keep_together(p)
    font(p.add_run(f'{n}  '), 10.5, MUTE, bold=True)
    font(p.add_run(txt), 10.5, HEAD, bold=True)
    return p


def h3(doc, txt):
    p = doc.add_paragraph(); spacing(p, 10, 3); keep_together(p)
    font(p.add_run(txt), 9.5, SUB, bold=True)
    return p


def para(doc, txt, after=7, size=10):
    p = doc.add_paragraph(); spacing(p, 0, after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for seg, bold in _marks(txt):
        font(p.add_run(seg), size, INK, bold=bold)
    return p


def _marks(txt):
    """**negrito** minimalista"""
    out, i = [], 0
    while True:
        a = txt.find('**', i)
        if a < 0:
            out.append((txt[i:], False)); break
        b = txt.find('**', a + 2)
        if b < 0:
            out.append((txt[i:], False)); break
        if a > i:
            out.append((txt[i:a], False))
        out.append((txt[a + 2:b], True))
        i = b + 2
    return [(s, bl) for s, bl in out if s]


def bullets(doc, items, after=4, size=10, marker='—'):
    for it in items:
        p = doc.add_paragraph(); spacing(p, 0, after)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Cm(0.62); pf.first_line_indent = Cm(-0.62)
        font(p.add_run(marker + '\t'), size, MUTE)
        for seg, bold in _marks(it):
            font(p.add_run(seg), size, INK, bold=bold)


def numlist(doc, items, after=5, size=10):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph(); spacing(p, 0, after)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Cm(0.72); pf.first_line_indent = Cm(-0.72)
        font(p.add_run(f'{i}.\t'), size, SUB, bold=True)
        for seg, bold in _marks(it):
            font(p.add_run(seg), size, INK, bold=bold)


def _box(doc, lines, bg, bar, ref=None, size=9.5, italic=False, width=Cm(16.6)):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    no_table_borders(t)
    row_no_split(t.rows[0])
    c = t.cell(0, 0)
    c.width = width
    cell_bg(c, bg)
    cell_borders(c, {'left': (18, bar)})
    cell_pad(c, 130, 130, 200, 180)
    c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)
    for i, ln in enumerate(lines):
        p = c.add_paragraph(); spacing(p, 0, 4 if i < len(lines) - 1 else 0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for seg, bold in _marks(ln):
            font(p.add_run(seg), size, INK, bold=bold, italic=italic)
    if ref:
        p = c.add_paragraph(); spacing(p, 6, 0)
        font(p.add_run(ref), 7.5, MUTE)
    sp = doc.add_paragraph(); spacing(sp, 0, 9)
    return t


def citacao(doc, lines, ref):
    return _box(doc, lines, CIT_BG, CIT_BAR, ref=ref, size=9.5, italic=True)


def destaque(doc, lines):
    return _box(doc, lines, KEY_BG, KEY_BAR, size=10)


def nota(doc, lines):
    return _box(doc, lines, NOTE_BG, NOTE_BAR, size=9.5)


def enquadramento(doc, lines):
    return _box(doc, lines, 'F6F7F8', 'C9D2DA', size=9.2)


def tabela(doc, header, rows, widths, shades=None, size=8.8, head_size=8.6):
    """shades: lista (por linha) de dict {col_idx: hex} ou None"""
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    no_table_borders(t)
    hdr = t.rows[0]
    row_repeat(hdr)
    for j, htxt in enumerate(header):
        c = hdr.cells[j]; c.width = widths[j]
        cell_bg(c, TBL_HEAD)
        cell_borders(c, {'bottom': (8, CIT_BAR)})
        cell_pad(c, 90, 90, 120, 120)
        p = c.paragraphs[0]; spacing(p, 0, 0)
        font(p.add_run(htxt), head_size, HEAD, bold=True)
    for i, row in enumerate(rows):
        r = t.add_row()
        row_no_split(r)
        for j, val in enumerate(row):
            c = r.cells[j]; c.width = widths[j]
            bg = TBL_ALT if i % 2 == 0 else 'FFFFFF'
            if shades and shades[i] and j in shades[i]:
                bg = shades[i][j]
            cell_bg(c, bg)
            cell_borders(c, {'bottom': (4, RULE)})
            cell_pad(c, 85, 85, 120, 120)
            c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)
            for k, ln in enumerate(str(val).split('\n')):
                p = c.add_paragraph(); spacing(p, 0, 2 if k else 0, 1.14)
                for seg, bold in _marks(ln):
                    font(p.add_run(seg), size, INK, bold=bold)
    sp = doc.add_paragraph(); spacing(sp, 2, 9)
    return t


def novo_doc(titulo_corrente):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10)
    st.font.color.rgb = RGBColor.from_string(INK)
    rf = st.element.rPr.rFonts
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.22

    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(2.3), Cm(2.1)
    s.left_margin, s.right_margin = Cm(2.2), Cm(2.2)
    s.header_distance, s.footer_distance = Cm(1.1), Cm(1.1)

    hp = s.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    spacing(hp, 0, 0)
    font(hp.add_run(titulo_corrente), 7.5, MUTE)

    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    spacing(fp, 0, 0)
    r = fp.add_run(); font(r, 8, MUTE)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    rr = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    szel = OxmlElement('w:sz'); szel.set(qn('w:val'), '16'); rpr.append(szel)
    colel = OxmlElement('w:color'); colel.set(qn('w:val'), MUTE); rpr.append(colel)
    rr.append(rpr); fld.append(rr); fp._p.append(fld)
    return doc


def capa(doc, chapeu, titulo, subtitulo, meta):
    p = doc.add_paragraph(); spacing(p, 0, 3)
    font(p.add_run(chapeu.upper()), 8.5, SUB, bold=True)
    for run in p.runs:
        put(run._element.get_or_add_rPr(), _el('w:spacing', val='50'), 'rPr')
    p = doc.add_paragraph(); spacing(p, 0, 2)
    font(p.add_run(titulo), 21, HEAD, bold=True)
    p = doc.add_paragraph(); spacing(p, 0, 10)
    font(p.add_run(subtitulo), 11, SUB)
    rule(doc, CIT_BAR, 8, 6)
    p = doc.add_paragraph(); spacing(p, 0, 14)
    font(p.add_run(meta), 8.5, MUTE)
    return doc


def pagebreak(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph(); spacing(p, 0, 0)
    p.add_run().add_break(WD_BREAK.PAGE)
