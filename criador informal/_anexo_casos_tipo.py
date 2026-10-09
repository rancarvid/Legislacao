# -*- coding: utf-8 -*-
"""Gera o anexo de análise ao quadro de casos-tipo (314 vs alojamentos)."""
import re, zipfile, shutil
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = "Casos_Tipo_Limites_por_Fogo_2026-10-09_Anexo.docx"

# ------------------------------------------------------------------ quadro
COLS = ["Caso", "Limite de animais adultos", "Quem autoriza e como", "Sanção e quem instrui",
        "Remoção e encerramento", "Urbanismo e vizinhança"]
L_FOGO = "N.º 2: três cães ou quatro gatos, máximo quatro; até seis com autorização"
A_FOGO = "Até seis: pedido do detentor, pareceres vinculativos do médico veterinário municipal e do delegado de saúde; autorização municipal"
S_314 = "Al. c) do n.º 3 do art. 14.º do DL 314/2003; instrução regional (art. 16.º, n.º 2); decisão do diretor-geral"
R_314 = "Câmara, após vistoria conjunta, notifica a remoção (n.º 5); mandado judicial (n.º 6)"
S_276 = "Al. f) do n.º 1 do art. 68.º do DL 276/2001 (contraordenação económica grave); instrução DGAV e polícias; decisão do diretor-geral (art. 70.º)"
R_276 = "Suspensão ou encerramento por despacho do diretor-geral, execução e recolha pelas câmaras (art. 3.º-G)"
LINHAS = [
 ("1. Fração autónoma em propriedade horizontal; animais do agregado, incluindo varanda ou terraço",
  L_FOGO + "; o regulamento do condomínio pode fixar menos (n.º 3)", A_FOGO + "; regras do condomínio",
  S_314, R_314, "Título constitutivo e deliberações do condomínio (art. 1422.º CC); ruído de vizinhança (RGR)"),
 ("2. Andar em prédio sem propriedade horizontal; animais do agregado",
  L_FOGO, A_FOGO, S_314, R_314, "Contrato de arrendamento (al. a) do n.º 2 do art. 1083.º CC); ruído de vizinhança (RGR)"),
 ("3. Moradia, com ou sem logradouro, qualquer que seja a dimensão; animais do agregado",
  L_FOGO + "; a área do logradouro é irrelevante", A_FOGO, S_314 + "; alcança o logradouro", R_314,
  "Anexos para animais até 1/15 do logradouro (art. 115.º RGEU); ruído de vizinhança (RGR)"),
 ("4. Moradia com alojamento registado em instalações do art. 25.º no logradouro",
  "No alojamento: capacidade declarada, limitada pelo anexo III. Dentro de casa: n.º 2",
  "Mera comunicação prévia à DGAV; obras dos anexos sujeitas ao RJUE (câmara)",
  S_276 + "; DL 314/2003 quanto aos animais da casa", R_276 + "; n.º 5 quanto aos animais da casa",
  "PDM e título de utilização; anexos até 1/15 do logradouro (art. 115.º RGEU); atividade ruidosa permanente (RGR)"),
 ("5. Fração autónoma ou andar com alojamento registado",
  "Hospedagem sem fins lucrativos excluída em fração autónoma (al. p) do n.º 1 do art. 2.º DL 276/2001). Restantes: anexo III; animais do agregado: n.º 2",
  "Mera comunicação prévia à DGAV; uso da fração compatível com o título",
  S_276, R_276, "Uso diverso do fim proibido (al. c) do n.º 2 e n.º 4 do art. 1422.º CC); título de utilização; RGR"),
 ("6. Prédio urbano sem fogo (loja, armazém, serviços) com alojamento registado",
  "N.º 2 sem campo; capacidade declarada, limitada pelo anexo III", "Mera comunicação prévia à DGAV; título de utilização compatível (câmara)",
  S_276, R_276, "PDM e título de utilização (RJUE); atividade ruidosa permanente (RGR)"),
 ("7. Prédio urbano com alojamento de facto, sem registo",
  "Dentro de casa: n.º 2. Nas instalações: atividade sem título; n.º 1",
  "Ninguém autorizou: falta a mera comunicação prévia",
  "Al. a) e f) do n.º 1 do art. 68.º DL 276/2001; al. c) do n.º 3 do art. 14.º DL 314/2003 se houver violação do art. 3.º",
  R_314 + "; recolha pela DGAV com as câmaras (n.º 8 do art. 19.º DL 276/2001)",
  "Operações urbanísticas não licenciadas (RJUE); RGR"),
 ("8. Prédio misto; animais no fogo da parte urbana",
  "N.º 2 quanto ao fogo; n.º 4 quanto ao prédio", A_FOGO, S_314, R_314, "Ruído de vizinhança (RGR)"),
 ("9. Prédio misto; animais no logradouro ou na parte rústica",
  "N.º 4: seis, excedíveis se a dimensão do terreno o permitir", "Sem autorização prévia",
  "Al. c) do n.º 3 do art. 14.º DL 314/2003 no terreno anexo à habitação; na parte rústica afastada, cobertura duvidosa",
  R_314, "PDM; ruído de vizinhança ou atividade, conforme o caso"),
 ("10. Prédio misto com alojamento registado",
  "No alojamento: anexo III. No fogo: n.º 2. Restantes: n.º 4", "Mera comunicação prévia à DGAV; obras sujeitas ao RJUE",
  S_276, R_276, "PDM (solo rústico ou urbano); atividade ruidosa permanente (RGR)"),
 ("11. Prédio rústico sem habitação nem atividade",
  "N.º 4: seis, excedíveis se a dimensão do terreno o permitir", "Sem autorização prévia",
  "Sem tipo contraordenacional no DL 314/2003 (não há habitação)", R_314, "PDM; RGR"),
 ("12. Prédio rústico com alojamento registado",
  "Capacidade declarada, limitada pelo anexo III", "Mera comunicação prévia à DGAV; edificação em solo rústico sujeita ao PDM e ao RJUE",
  S_276, R_276, "Regime do solo rústico (PDM); atividade ruidosa permanente (RGR)"),
 ("13. Prédio rústico ou misto com alojamento de facto, sem registo",
  "N.º 4; atividade sem título; n.º 1", "Ninguém autorizou: falta a mera comunicação prévia",
  "Al. a) e f) do n.º 1 do art. 68.º DL 276/2001", R_314 + "; recolha pela DGAV com as câmaras (n.º 8 do art. 19.º DL 276/2001)",
  "Operações urbanísticas não licenciadas (RJUE); RGR"),
]

# ------------------------------------------------------------------ texto
T = []
def tt(x): T.append(("t", x))
def s(x): T.append(("s", x))
def p(x): T.append(("p", x))
def q(x, f): T.append(("q", "«" + x + "»")); T.append(("f", f))
def li(x): T.append(("l", x))

tt("Limites de detenção por fogo e alojamentos: quadro de casos-tipo e análise")
p("Este documento responde, para cada caso-tipo, a cinco questões: o limite de animais adultos, quem autoriza e como, a sanção e quem instrui, a remoção e o encerramento, e o enquadramento urbanístico e de vizinhança. Apresenta primeiro o quadro e depois a justificação de cada campo. Assenta na síntese «Limites de detenção por fogo e lotação dos alojamentos registados», no respetivo anexo e no estudo extenso, que se referem como «a síntese» e «o estudo».")
p("Os casos organizam-se pelo lugar onde estão os animais, que é o critério que decide a norma aplicável, e não pela classificação matricial do prédio nem pela dimensão do terreno.")

s("1. Regras gerais, comuns a todos os casos")
p("1.1. O dever de salubridade aplica-se sempre, em qualquer prédio e a qualquer detentor, com ou sem registo:")
q("O alojamento de cães e gatos em prédios urbanos, rústicos ou mistos, fica sempre condicionado à existência de boas condições do mesmo e ausência de riscos hígio-sanitários relativamente à conspurcação ambiental e doenças transmissíveis ao homem.",
  "N.º 1 do artigo 3.º do Decreto-Lei n.º 314/2003, de 17 de dezembro, Diário da República, I série-A, n.º 290, de 17.12.2003, p. 8445")
p("1.2. Contam-se apenas os animais adultos. Para efeitos do Decreto-Lei n.º 314/2003, é adulto o cão ou o gato com idade igual ou superior a um ano (alíneas f) e g) do artigo 2.º). As ninhadas e os animais jovens não entram no limite, sem prejuízo do n.º 1.")
p("1.3. O limite do n.º 2 conta por fogo, entendido como «uma parte ou a totalidade de um edifício, dotada de acesso independente, constituída por um ou mais compartimentos destinados à habitação e por espaços privativos complementares» (ficha n.º I-31 do anexo I do Decreto Regulamentar n.º 5/2019, de 27 de setembro, Diário da República, 1.ª série, n.º 186, de 27.9.2019, p. 44). As notas da mesma ficha distinguem o fogo da habitação, que integra também as dependências do fogo, como varandas, terraços, anexos, logradouros pavimentados, telheiros e alpendres. Os animais do agregado contam para o fogo a que pertencem, estejam dentro de casa ou nessas dependências.")
p("1.4. O registo de um alojamento não altera a contagem dos animais que vivem dentro do fogo. A lotação do alojamento só abrange os animais mantidos nas instalações do alojamento.")
p("1.5. A organização do quadro pelo lugar dos animais funde casos que a lei não distingue. A dimensão do logradouro não releva para o n.º 2 num prédio urbano, porque o logradouro integra o prédio urbano e não converte o conjunto em prédio misto (estudo, capítulo 10 e ponto 15.4). A varanda e o terraço são dependências do fogo e seguem a regra do fogo.")

s("2. Quadro")
T.append(("tabela", None))
p("Nota: o n.º 1 do artigo 3.º do Decreto-Lei n.º 314/2003 aplica-se em todos os casos e não se repete no quadro.")

s("3. Justificação por campo")

s("3.1. Limite de animais adultos")
q("Nos prédios urbanos podem ser alojados até três cães ou quatro gatos adultos por cada fogo, não podendo no total ser excedido o número de quatro animais, excepto se, a pedido do detentor, e mediante parecer vinculativo do médico veterinário municipal e do delegado de saúde, for autorizado alojamento até ao máximo de seis animais adultos, desde que se verifiquem todos os requisitos hígio-sanitários e de bem-estar animal legalmente exigidos.",
  "N.º 2 do artigo 3.º do Decreto-Lei n.º 314/2003, p. 8445")
q("No caso de fracções autónomas em regime de propriedade horizontal, o regulamento do condomínio pode estabelecer um limite de animais inferior ao previsto no número anterior.",
  "N.º 3 do mesmo artigo, p. 8445")
q("Nos prédios rústicos ou mistos podem ser alojados até seis animais adultos, podendo tal número ser excedido se a dimensão do terreno o permitir e desde que as condições de alojamento obedeçam aos requisitos estabelecidos no n.º 1.",
  "N.º 4 do mesmo artigo, p. 8445")
p("Casos 1 a 3. O limite é o do n.º 2, por fogo. Na fração autónoma em propriedade horizontal acresce o n.º 3, que permite ao regulamento do condomínio fixar um limite inferior. No andar de prédio que não está em propriedade horizontal, o n.º 3 não se aplica, e as limitações vêm do contrato de arrendamento, quando o haja. Na moradia, a área do logradouro não aumenta nem diminui o limite.")
p("Casos 4, 5, 6, 10 e 12. Os animais mantidos nas instalações do alojamento não contam para o n.º 2. A lotação do alojamento é a capacidade declarada na mera comunicação prévia («A capacidade máxima de animais e respetivas espécies a alojar», alínea h) do n.º 1 do artigo 3.º-A do Decreto-Lei n.º 276/2001, aditado pelo Decreto-Lei n.º 260/2012, Diário da República, 1.ª série, n.º 240, de 12.12.2012, p. 6972), limitada pelas superfícies mínimas do anexo III, que o n.º 1 do artigo 27.º torna obrigatórias (Diário da República, I série-A, n.º 241, de 17.10.2001, pp. 6577, 6585 e 6586). Os animais que o titular mantém dentro de casa, como animais do agregado, continuam sujeitos ao n.º 2 (síntese, ponto 8).")
p("Caso 5. Há uma restrição própria da fração autónoma: a hospedagem sem fins lucrativos exclui expressamente as frações autónomas.")
q("p) «Hospedagem sem fins lucrativos», alojamento, permanente ou temporário, de animais de companhia que não vise a obtenção de rendimentos, com excepção das referidas no n.º 3 do artigo 3.º do diploma que aprova o Plano Nacional de Luta e Vigilância da Raiva Animal e outras Zoonoses;",
  "Alínea p) do n.º 1 do artigo 2.º do Decreto-Lei n.º 276/2001, na redação do Decreto-Lei n.º 315/2003, Diário da República, I série-A, n.º 290, de 17.12.2003, p. 8450")
p("Uma associação não pode, portanto, registar um alojamento de hospedagem sem fins lucrativos numa fração autónoma: os animais aí mantidos ficam sujeitos ao n.º 2 e ao n.º 3. As outras atividades (criação, hotel) exigem as instalações individualizadas do artigo 25.º e um uso da fração compatível com o título, o que torna o caso excecional na prática (ver 3.5).")
p("Caso 6. Num prédio urbano sem fogo, por exemplo um edifício licenciado para comércio, serviços ou armazém, o n.º 2 não tem unidade a que se aplicar. Vale apenas a capacidade do alojamento.")
p("Casos 7 e 13. Sem registo, não há capacidade declarada. Os animais mantidos dentro de casa continuam sujeitos ao n.º 2 e, no prédio rústico ou misto, ao n.º 4. A atividade exercida sem título é ilícita em si mesma (ver 3.3), e o n.º 1 aplica-se sempre.")
p("Casos 8 a 11. No prédio misto, o n.º 4 vale para o prédio no seu conjunto e o n.º 2 tem campo quanto ao fogo da parte urbana (estudo, ponto 15.7). O n.º 4 fixa seis animais adultos, excedíveis se a dimensão do terreno o permitir, sem teto nem controlo prévio (estudo, ponto 15.2; acórdão do Tribunal Central Administrativo Sul de 4.2.2010, processo n.º 04784/09).")

s("3.2. Quem autoriza e como")
p("Casos 1 a 3 e 8. Dentro dos limites do n.º 2 (três cães ou quatro gatos, no máximo quatro animais) não há autorização. Acima disso e até seis, a lei exige pedido do detentor e parecer vinculativo do médico veterinário municipal e do delegado de saúde. Não diz que órgão autoriza. A hipótese mais sustentada é a autorização municipal, porque a execução do artigo 3.º é cometida às câmaras municipais (n.º 5) e é assim que os regulamentos municipais a organizam. O Regulamento n.º 181/2025 do Município do Cartaxo, por exemplo, prevê requerimento ao presidente da câmara e vistoria conjunta (artigo 13.º, Diário da República, 2.ª série, n.º 22, de 31.1.2025). Acima de seis, dentro do fogo, não há autorização possível.")
p("Caso 1. Acrescem as regras do condomínio: o regulamento pode fixar limite inferior (n.º 3 do artigo 3.º) e os condóminos não podem praticar «actos ou actividades que tenham sido proibidos no título constitutivo ou, posteriormente, por deliberação da assembleia de condóminos aprovada sem oposição» (alínea d) do n.º 2 do artigo 1422.º do Código Civil).")
p("Casos 4, 5, 6, 10 e 12. O título de acesso à atividade é a mera comunicação prévia à DGAV (alínea a) do n.º 1 do artigo 3.º e artigo 3.º-A do Decreto-Lei n.º 276/2001), sem vistoria prévia. A reprodução e criação de animais potencialmente perigosos depende de permissão administrativa (alínea b) do n.º 1 do artigo 3.º). Na hospedagem com fins lucrativos destinada à reprodução e criação, só é permitida a atividade com animais pertencentes ao titular da exploração, salvo o acolhimento temporário para acasalamento (n.ºs 3 e 4 do artigo 2.º). Independentemente do título de acesso, as edificações e a sua utilização dependem do município, nos termos do regime jurídico da urbanização e edificação.")
p("Casos 7 e 13. Ninguém autorizou a atividade: falta a mera comunicação prévia.")
p("Casos 9 e 11. O n.º 4 não prevê autorização: o número pode ser excedido se a dimensão do terreno o permitir, apreciação que cabe à Administração em caso de controlo, e não a título prévio.")

s("3.3. Sanção e quem instrui")
q("c) A permanência de cães e gatos em habitações e terrenos anexos em desrespeito pelas condições previstas no artigo 3.º;",
  "Alínea c) do n.º 3 do artigo 14.º do Decreto-Lei n.º 314/2003, p. 8448")
p("Casos 1 a 3, 8 e 9. A violação do artigo 3.º em habitações e terrenos anexos é contraordenação, «punível pelo director-geral de Veterinária», hoje diretor-geral de Alimentação e Veterinária, com coima de 50 a 3740 euros ou a 44 890 euros, consoante o agente seja pessoa singular ou coletiva (n.º 3 do artigo 14.º). A instrução compete «à DRA da área em que foi praticada a infracção» (n.º 2 do artigo 16.º). As direções regionais de agricultura foram extintas, e a entidade que lhes sucedeu nesta competência deve ser confirmada. No caso 9, a alínea alcança os terrenos anexos à habitação. Na parte rústica afastada da casa a cobertura é duvidosa (estudo, ponto 15.5).")
p("Caso 11. Não há tipo contraordenacional no Decreto-Lei n.º 314/2003, porque não há habitação nem terreno anexo a ela. Fica a remoção do n.º 5 (estudo, ponto 15.5).")
q("a) A falta da mera comunicação prévia ou da permissão administrativa previstas no n.º 1 do artigo 3.º; […] f) O alojamento de animais de companhia em desrespeito das condições fixadas no presente diploma;",
  "Alíneas a) e f) do n.º 1 do artigo 68.º do Decreto-Lei n.º 276/2001, na redação do Decreto-Lei n.º 9/2021, de 29 de janeiro")
p("Casos 4, 5, 6, 10 e 12. A violação das condições do alojamento é contraordenação económica grave, punível nos termos do regime jurídico das contraordenações económicas. A instrução compete à DGAV e aos órgãos de polícia criminal, e a aplicação das coimas ao diretor-geral de Alimentação e Veterinária ou ao diretor do órgão de polícia criminal (artigo 70.º). Quanto aos animais mantidos dentro de casa, vale o Decreto-Lei n.º 314/2003.")
p("Casos 7 e 13. A falta da mera comunicação prévia é contraordenação económica grave (alínea a) do n.º 1 do artigo 68.º), a que pode acrescer a alínea f) pelas condições do alojamento e, quanto a habitação e terrenos anexos, a alínea c) do n.º 3 do artigo 14.º do Decreto-Lei n.º 314/2003.")
p("Em todos os casos ficam ressalvados os crimes de maus-tratos e de abandono (artigos 387.º e 388.º do Código Penal).")

s("3.4. Remoção e encerramento")
q("Em caso de não cumprimento do disposto nos números anteriores, as câmaras municipais, após vistoria conjunta do delegado de saúde e do médico veterinário municipal, notificam o detentor para retirar os animais para o canil ou gatil municipal no prazo estabelecido por aquelas entidades, caso o detentor não opte por outro destino que reúna as condições estabelecidas pelo presente diploma.",
  "N.º 5 do artigo 3.º do Decreto-Lei n.º 314/2003, p. 8445")
p("Casos 1 a 3, 8, 9 e 11. A remoção é municipal, após vistoria conjunta, e pode ser apoiada por mandado judicial a pedido do presidente da câmara (n.º 6). É o único remédio disponível no caso 11.")
q("O diretor-geral de Alimentação e Veterinária pode, mediante despacho, determinar a suspensão da atividade ou o encerramento do alojamento, designadamente quando se verifique uma das seguintes situações: a) Existência de riscos higiossanitários que ponham em causa a saúde das pessoas e ou dos animais; b) Maus tratos aos animais; c) Existência de graves problemas de saúde e bem-estar dos animais; […]",
  "N.º 1 do artigo 3.º-G do Decreto-Lei n.º 276/2001, aditado pelo Decreto-Lei n.º 260/2012, p. 6975")
p("Casos 4, 5, 6, 10 e 12. Quem decide é o diretor-geral. Às câmaras compete executar a decisão, «nomeadamente proceder, quando necessário, à recolha dos animais» (n.º 6 do mesmo artigo, p. 6976). Quanto aos animais mantidos dentro de casa, mantém-se o n.º 5 do artigo 3.º do Decreto-Lei n.º 314/2003.")
p("Casos 7 e 13. Além do n.º 5, a DGAV, com a intervenção das câmaras municipais quando necessário, e as autoridades policiais devem proceder à recolha dos animais quando esteja em causa a sua saúde e bem-estar, podendo pedir mandado judicial para aceder a casas de habitação e terrenos privados (n.º 8 do artigo 19.º do Decreto-Lei n.º 276/2001). Se o artigo 3.º-G se aplica a um alojamento que nunca foi registado é questão discutível (estudo, ponto 15.7).")

s("3.5. Urbanismo e vizinhança")
p("Propriedade horizontal (casos 1 e 5). É vedado aos condóminos dar à fração «uso diverso do fim a que é destinada» (alínea c) do n.º 2 do artigo 1422.º do Código Civil), e, quando o título constitutivo não disponha sobre o fim da fração, a alteração do uso carece de autorização da assembleia por maioria de dois terços do valor total do prédio (n.º 4 do mesmo artigo). Instalar um alojamento numa fração destinada a habitação é alteração de uso.")
p("Arrendamento (caso 2). A violação de regras de higiene, sossego e boa vizinhança pode fundar a resolução do contrato (alínea a) do n.º 2 do artigo 1083.º do Código Civil). O Tribunal da Relação de Guimarães incluiu o artigo 3.º nessas regras (acórdão de 19.5.2022, processo n.º 119/20.1T8FAF.G1).")
q("Os anexos para alojamento de animais domésticos construídos nos logradouros dos prédios, quando expressamente autorizados, não poderão ocupar mais do que 1/15 da área destes logradouros.",
  "Artigo 115.º do Regulamento Geral das Edificações Urbanas, aprovado pelo Decreto-Lei n.º 38 382, de 7 de agosto de 1951")
p("Moradias e alojamentos em prédio urbano (casos 3 e 4). O artigo 115.º do Regulamento Geral das Edificações Urbanas limita os anexos para animais no logradouro a 1/15 da sua área, quando expressamente autorizados, e permite às câmaras interditá-los em zonas de aglomeração de habitações (§ único). Este limite de área é, para o alojamento instalado no logradouro de uma moradia, um controlo de dimensão independente do Decreto-Lei n.º 314/2003. O Regulamento está em vigor, por força do artigo 25.º do Decreto-Lei n.º 10/2024, na redação do Decreto-Lei n.º 108/2026.")
p("Uso do solo (casos 4, 6, 10 e 12). O alojamento tem de ser compatível com o plano diretor municipal e com o título de utilização da edificação. O Supremo Tribunal Administrativo declarou nulo o licenciamento de um canil para cem cães numa zona de uso dominantemente residencial (acórdão de 7.3.2006, processo n.º 0794/05). Em solo rústico, a edificação de instalações depende do regime do solo fixado no plano.")
p("Operações urbanísticas sem controlo (casos 7 e 13). Contra alojamentos de facto, o instrumento municipal que se mostrou eficaz foi o regime jurídico da urbanização e edificação, com coimas por operações urbanísticas não licenciadas (estudo, ponto 15.7, caso de Santo Tirso).")
p("Ruído. O Regulamento Geral do Ruído, aprovado pelo Decreto-Lei n.º 9/2007, de 17 de janeiro, distingue duas situações, que coincidem com a distinção entre fogo e alojamento:")
q("r) «Ruído de vizinhança» o ruído associado ao uso habitacional e às actividades que lhe são inerentes, produzido directamente por alguém ou por intermédio de outrem, por coisa à sua guarda ou animal colocado sob a sua responsabilidade, que, pela sua duração, repetição ou intensidade, seja susceptível de afectar a saúde pública ou a tranquilidade da vizinhança;",
  "Alínea r) do artigo 3.º do Regulamento Geral do Ruído, Diário da República, 1.ª série, n.º 12, de 17.1.2007, p. 391")
p("Os animais do agregado (casos 1 a 3, 8 e 9) produzem ruído de vizinhança. As autoridades policiais podem ordenar a sua cessação imediata entre as 23 e as 7 horas e fixar prazo entre as 7 e as 23 horas (artigo 24.º, p. 396), e são elas que fiscalizam (alínea f) do artigo 26.º).")
p("O alojamento (casos 4, 5, 6, 7, 10, 12 e 13) é uma «actividade ruidosa permanente», definida como «a actividade desenvolvida com carácter permanente, ainda que sazonal, que produza ruído nocivo ou incomodativo para quem habite ou permaneça em locais onde se fazem sentir os efeitos dessa fonte de ruído» (alínea a) do artigo 3.º, p. 390). Está sujeito aos valores limite e ao critério de incomodidade (artigo 13.º, pp. 393 e 394). Fiscalizam a entidade responsável pelo licenciamento ou autorização da atividade e as câmaras municipais (alíneas b) e d) do artigo 26.º). A citação faz-se pela redação originária do Regulamento. As alterações posteriores devem ser confirmadas.")
p("Em todos os casos, o vizinho pode opor-se a ruídos e cheiros que importem prejuízo substancial ou não resultem da utilização normal do prédio (artigo 1346.º do Código Civil).")

s("4. Pontos que ficam por confirmar")
li("A entidade que sucedeu às direções regionais de agricultura na instrução dos processos do n.º 3 do artigo 14.º do Decreto-Lei n.º 314/2003 (n.º 2 do artigo 16.º).")
li("O órgão municipal competente para a autorização até seis animais, que a lei não nomeia. A resposta depende do regulamento de cada município.")
li("A aplicação do artigo 3.º-G a alojamentos nunca registados.")
li("As alterações ao Regulamento Geral do Ruído posteriores a 2007 quanto às normas citadas.")
li("A fração autónoma com alojamento (caso 5) não foi analisada no estudo (capítulo 17). A hipótese aqui apresentada é nova e resulta da alínea p) do n.º 1 do artigo 2.º do Decreto-Lei n.º 276/2001 e do artigo 1422.º do Código Civil.")

# ------------------------------------------------------------------ verificações
txt = " ".join(x for k, x in T if x) + " ".join(" ".join(r) for r in LINHAS)
assert "—" not in txt and "–" not in txt, "travessão"

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = sec.bottom_margin = Cm(2.0)
st = d.styles["Normal"]; st.font.name = "Aptos"; st.font.size = Pt(11); st.font.color.rgb = RGBColor(0, 0, 0)
rf = st.element.get_or_add_rPr().find(qn("w:rFonts"))
for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), "Aptos")
pf = st.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(6); pf.line_spacing = 1.0

def tabela():
    tb = d.add_table(rows=1, cols=len(COLS))
    tb.style = "Table Grid"
    for i, c in enumerate(COLS):
        tb.rows[0].cells[i].text = c
    for linha in LINHAS:
        cells = tb.add_row().cells
        for i, v in enumerate(linha):
            cells[i].text = v
    for row in tb.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                for r in par.runs:
                    r.font.size = Pt(8); r.font.name = "Aptos"; r.font.color.rgb = RGBColor(0, 0, 0)
    larg = [Cm(3.2), Cm(3.0), Cm(2.8), Cm(3.0), Cm(2.6), Cm(2.4)]
    for row in tb.rows:
        for i, c in enumerate(row.cells):
            c.width = larg[i]

for kind, x in T:
    if kind == "tabela":
        tabela(); continue
    par = d.add_paragraph(x)
    f = par.paragraph_format
    if kind in ("p", "q", "l"): par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if kind == "t": f.space_after = Pt(12)
    if kind == "s": f.space_before = Pt(10); f.keep_with_next = True
    if kind == "q": f.left_indent = Cm(1); f.space_after = Pt(1)
    if kind == "f": f.left_indent = Cm(1)
    if kind == "l": f.left_indent = Cm(0.5)
d.save(OUT)
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if it.filename == "word/settings.xml":
            data = re.sub(r'<w:zoom(?![^>]*w:percent)([^>]*)/>', r'<w:zoom w:percent="100"\1/>', data.decode()).encode()
        zo.writestr(it, data)
shutil.move(tmp, OUT)
print(OUT, "palavras:", len(txt.split()))
