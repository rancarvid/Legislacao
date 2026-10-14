# -*- coding: utf-8 -*-
"""Síntese interna das reservas à proposta OMV/ANVETEM, em frases (máx. 25; máx. 60 palavras cada).

Três partes: A. recolha compulsiva; B. salvamento e socorro; C. tronco comum.
Gera dois Word na mesma pasta:
  Sintese_Observacoes_Proposta_OMV_ANVETEM.docx  (as frases)
  Apoio_Sintese_Observacoes.docx                 (que temas da análise v2.1 cada frase colige)
Executar a partir da raiz do repositório:
    python3 "Protocolo de retirada de alojamentos/_sintese_observacoes.py"
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PASTA = Path(__file__).resolve().parent
SAIDA_FRASES = PASTA / 'Sintese_Observacoes_Proposta_OMV_ANVETEM.docx'
SAIDA_APOIO = PASTA / 'Apoio_Sintese_Observacoes.docx'
ANALISE = PASTA / '_analise_proposta_omv_anvetem_v2.py'

AZUL = RGBColor(0x1F, 0x3B, 0x63)
CINZA = RGBColor(0x55, 0x55, 0x55)
MAX_PALAVRAS = 60
MAX_FRASES = 25

PARTES = [
    ('A', 'Parte A. Operações de recolha compulsiva (alojamentos, acumuladores e outras retiradas fora de emergência)'),
    ('B', 'Parte B. Operações de salvamento e socorro'),
    ('C', 'Parte C. Tronco comum'),
]

FRASES = [
# ---------------- A. Recolha compulsiva
('A1', 'Assenta só no DL 276/2001 (arts. 19.º e 3.º-G), frágil para cães e gatos, que o n.º 10 do art. 19.º remete '
       'para «legislação própria», e omite as bases mais fortes: Lei 92/95 (art. 1.º-A), CPP (arts. 178.º e 186.º), '
       'RJCE (arts. 48.º a 50.º), DL 314/2003 (art. 3.º) e DL 315/2009 (art. 39.º).'),
('A2', 'O âmbito (ponto 1.1 e tipologia A) omite os riscos higiossanitários e ambientais, fundamento expresso de '
       'suspensão, encerramento e remoção (DL 276/2001, art. 3.º-G, n.º 1, als. a) e d); DL 314/2003, art. 3.º, n.º 1), '
       'e não indica o art. 3.º-G como base para os alojamentos registados.'),
('A3', 'O ponto 1.5 só exclui a captura «de rotina» de errantes, termo vago para norma técnica e sem critério, e deixa '
       'dentro, sem os nomear, a remoção por excesso ou insalubridade (DL 314/2003, art. 3.º) e todo o DL 315/2009; a '
       'recolha após agressão (art. 14.º) devia sair. «Umas e outras» é ambíguo.'),
('A4', 'Os acumuladores ficam sem enquadramento: nada sobre a autoridade de saúde (vistoria conjunta do DL 314/2003, '
       'art. 3.º), a ação social (guia previsto na RCM 78/2021) ou a reincidência. Questão aos proponentes: que '
       'enquadramento propõem para estes casos e para a remoção por excesso ou insalubridade?'),
('A5', 'A DGAV assume o comando e o destino de todas as operações em que participe (pontos 1.2, 4.2 e 7.5), sem título '
       'legal: no processo penal dirige o Ministério Público e decide a autoridade judiciária (CPP, art. 186.º); na remoção por '
       'insalubridade decidem a câmara e o MVM (DL 314/2003, arts. 3.º e 9.º).'),
('A6', 'O CRO é o único destino previsto (ponto 7.1) e não há título para os animais que ficam no local nas operações '
       'faseadas; faltam o detentor como fiel depositário (RJCE, art. 49.º), o depositário idóneo (CPP, art. 178.º), '
       'abrigos de associações e lares de acolhimento, e a cedência voluntária no ato.'),
('A7', 'Animais apreendidos que entrem em CRO ficam sujeitos à presunção de abandono aos 15 dias e ao registo em nome '
       'do CRO (Lei 27/2016, art. 3.º; DL 82/2019, art. 11.º), podendo ser dados para adoção antes da decisão da '
       'autoridade; a proposta não o trata nem prevê estado «sob apreensão» no SIAC.'),
('A8', 'A triagem admite «eutanásia ou occisão» (ponto 6.2) com critério mais largo do que o legal (Lei 27/2016, art. '
       '3.º, n.º 6) e sem proteger a prova: havendo processo penal, faltam comunicação ao Ministério Público e exame ou '
       'necropsia (CPP, art. 159.º-A), também antes de eliminar cadáveres (ponto 6.7).'),
('A9', 'Exige autorização de transportador (ponto 9.1), que não se aplica a animais retirados por autoridades ou '
       'associações: o Regulamento (CE) 1/2005 exclui o transporte fora de atividade económica e o transporte direto '
       'para CAMV (art. 1.º, n.º 5). Basta o critério do DL 276/2001, art. 10.º, n.º 1.'),
# ---------------- B. Salvamento e socorro
('B1', 'Trata «busca, salvamento e socorro» como um só bloco (tipologia D). O salvamento retira o animal do perigo e '
       'não exige presença veterinária; o socorro presta cuidados e exige-a. Impor responsável clínico e triagem a todo '
       'o salvamento atrasa a operação e mobiliza veterinários onde não são indispensáveis.'),
('B2', 'O comando único da autoridade veterinária e o porta-voz único (princípios 1 e 5; pontos 4 e 8.1) contrariam o '
       'SIOPS: no socorro comanda o COS, que também informa a comunicação social (arts. 9.º e 10.º). À DGAV cabe a '
       'coordenação no CCON e nos CCOR e a direção técnica do setor animal.'),
('B3', 'Mandado judicial prévio (ponto 4.5) e observação individual antes da saída (ponto 5.2.4) são incompatíveis com '
       'uma emergência: a entrada justifica-se pelo perigo atual (Código Penal, art. 34.º) e a triagem faz-se depois, '
       'na zona de concentração de acolhimento de animais (SIOPS, art. 13.º, al. d)).'),
('B4', 'A comunicação prévia e o relatório à comissão municipal de proteção civil (pontos 4.3 e 10.2) não têm base na '
       'Lei 38/2026, que quanto à comissão só altera a sua composição; no socorro a operação nasce e é relatada na cadeia do '
       'SIOPS, e os setores funcionais (ponto 4.2) são definidos pelo COS.'),
('B5', 'A credenciação de entidades de socorro (ponto 9.1) ignora a figura legal própria: as organizações de '
       'voluntariado de proteção civil, reconhecidas pela ANEPC (Lei de Bases, art. 46.º-A; Portaria 91/2017), cuja '
       'portaria ainda só refere «pessoas e bens» e deve passar a incluir os animais.'),
('B6', 'Aplica ao socorro o modelo da apreensão (fiel depositário, decisão de destino, reserva), quando o animal '
       'continua a ser do detentor. Faltam a reunificação via SIAC e a passagem para o regime de retirada quando o '
       'socorro revela maus-tratos; a presunção de abandono aos 15 dias pode fazer perder animais de desalojados.'),
('B7', 'Omite os deveres da DGAV desde 8.8.2026 (representantes na Comissão Nacional, comissões distritais, CCON e '
       'CCOR), a eutanásia no terreno, hoje só admitida em CRO ou CAMV, e os planos de evacuação de CRO, abrigos e '
       'criadores; e põe a DGAV a planear o que é dever municipal (Lei 65/2007, art. 2.º).'),
# ---------------- C. Tronco comum
('C1', 'Um despacho só vincula os serviços da DGAV: não obriga câmaras, MVM (dependentes do presidente da câmara, DL '
       '116/98, art. 4.º), polícias, Ministério Público nem associações. Sem cláusula de que é norma interna, a exigência '
       'de documento prévio (ponto 4.1) pode ser invocada pela defesa contra atos praticados ao abrigo da lei.'),
('C2', 'Usa termos sem correspondência legal («resgate», «retirada»), troca «salvamento» por «resgate» ao citar a Lei '
       '38/2026 (ponto 9.1) e define a entidade coordenadora pelo DL 276/2001, que exclui as espécies pecuárias (art. '
       '1.º, n.º 2), misturando quatro tipologias com bases legais diferentes num só texto.'),
('C3', 'A responsabilidade clínica «indelegável» (ponto 6.1) é inviável para veterinário sem CAMV, porque a assistência '
       'exige instalações e a eutanásia só pode ser feita em CRO ou CAMV; não diz quem o designa, com que meios, nem '
       'quem o substitui, e remete para o art. 4.º do DL 276/2001, que trata do veterinário do alojamento.'),
('C4', 'O ponto 5.1.4 ignora o veterinário contratado pelo município, que gere o CRO sem poderes de autoridade (DL '
       '116/98, art. 2.º, n.º 3); «convocação» não tem sentido definido; e a nota dos proponentes erra ao dizer que o DL '
       '116/98 não refere o bem-estar animal, que consta do dever de colaboração do MVM (art. 3.º, n.º 1).'),
('C5', 'A credenciação, a lista anual e a suspensão ou exclusão de entidades (pontos 8.8, 9.3 e 9.5) criam efeitos '
       'externos sem norma habilitante, usam dois termos para a mesma sanção, não dizem quem decide nem o procedimento, '
       'e impõem seguro de responsabilidade civil sem fonte de financiamento.'),
('C6', 'As regras de reserva e as proibições de imagens e de angariação de fundos (pontos 8.2 a 8.5) restringem '
       'terceiros sem base legal e colidem com a legitimidade das associações zoófilas como assistentes (Lei 92/95, arts. '
       '9.º e 10.º); só valem como cláusula do termo de compromisso. O 8.5 repete o princípio 5.'),
('C7', 'Os encargos vão para uma «entidade responsável» que não é identificada (ponto 7.7), e a lei não tem regra de '
       'financiamento antecipado da custódia; a proposta atribui à DGAV planos, simulacros, inventários, listas e '
       'relatórios sem meios nem limiar de proporcionalidade.'),
('C8', 'Pormenores: posto de triagem sem condições de bem-estar e equipa sem enfermagem veterinária (pontos 5.2.2 e '
       '5.2.4); relatório distribuído sem critério de necessidade nem salvaguarda do segredo de justiça (10.2); '
       'avaliação anual só com OMV e ANVETEM (10.4); números variáveis e órgãos internos da OMV no texto (11.4).'),
('C9', 'Assenta no modelo vigente (a DGAV recolhe e decide o destino), que o projeto de RGAC inverte (câmaras recolhem, '
       'MVM decide; art. 91.º), e o RGAC mantém na apreensão o «retorno ao local de origem» sem ressalvar CPP nem RJCE '
       '(art. 139.º); procedimento e RGAC têm de convergir.'),
]

# ------------------------------------------------------------------ correspondência com a análise v2.1
# código: (secções da análise v2.1, notas manuscritas, bases legais, rótulos das recomendações)
COLIGE = {
    'A1': ('1.1; 4.2 (vias V1 a V7); 4.4, n); 3.3 (ponto 4.5)', '4',
           'DL 276/2001, arts. 3.º-G e 19.º; Lei 92/95, art. 1.º-A; CPP, arts. 178.º e 186.º; RJCE, arts. 48.º a 50.º; '
           'DL 314/2003, art. 3.º; DL 315/2009, art. 39.º', ['Completar a base legal']),
    'A2': ('4.1; 4.4, o)', '2, 8', 'DL 276/2001, art. 3.º-G, n.º 1, als. a) e d); DL 314/2003, art. 3.º, n.º 1',
           ['Âmbito completo']),
    'A3': ('4.4, u); 3.3 (ponto 1.5)', '9, 12; decisões de 14.10.2026',
           'Proposta, pontos 1.2, 1.5 e 11.3; DL 314/2003, arts. 3.º, 8.º e 9.º; DL 315/2009, arts. 14.º e 39.º',
           ['Ponto 1.5']),
    'A4': ('4.4, j); 4.4, u) (questão em aberto); 4.5, A7', 'decisão de 14.10.2026',
           'DL 314/2003, art. 3.º, n.º 5; RCM 78/2021, al. h)', ['Acumulação de animais', 'Ponto 1.5']),
    'A5': ('1.1; 4.4, a) e d)', '5, 16', 'CPP, art. 186.º; DL 314/2003, arts. 3.º e 9.º; RJCE, art. 50.º; '
           'DL 276/2001, art. 19.º, n.º 9', ['Matriz de liderança']),
    'A6': ('4.4, b), c) e d)', '17, 34', 'RJCE, art. 49.º; CPP, art. 178.º, n.º 2; Regulamento (UE) 2026/1818, '
           'art. 4.º, al. s)', ['Novos modelos']),
    'A7': ('4.4, e) e i); 4.5, A2', '—', 'Lei 27/2016, art. 3.º, n.º 1; DL 82/2019, arts. 8.º, n.ºs 4 e 5, e 11.º, '
           'n.º 5', ['SIAC', 'Orientação interpretativa']),
    'A8': ('4.4, f) e r)', '32, 33', 'Lei 27/2016, art. 3.º, n.º 6; CPP, art. 159.º-A; Regulamento (CE) '
           '1069/2009, arts. 8.º e 19.º', ['Eutanásia', 'Cadáveres']),
    'A9': ('4.4, s); 4.5, A9', '38', 'Regulamento (CE) 1/2005, art. 1.º, n.º 5; DL 265/2007; DL 276/2001, '
           'art. 10.º, n.º 1', ['Transporte']),
    'B1': ('Novo: não consta da análise v2.1 (observação de 14.10.2026)', '—',
           'Lei 38/2026 (LBPC, art. 4.º, n.º 2, al. e): «busca, salvamento, prestação de socorro e de assistência»); '
           'proposta, tipologia D e pontos 6.1 e 6.2', []),
    'B2': ('5.4, a)', '22', 'SIOPS, arts. 3.º, 4.º, 9.º, 10.º e 12.º', ['Reescrever a tipologia D']),
    'B3': ('5.4, j)', '20, 30', 'Código Penal, art. 34.º; SIOPS, art. 13.º, al. d)', ['Retirar do módulo B']),
    'B4': ('5.4, k)', '18, 19, 42', 'LBPC, art. 41.º, als. j) e k); SIOPS, arts. 10.º, n.º 1, als. g) e p), e 14.º',
           ['Retirar a comunicação prévia']),
    'B5': ('5.2; 5.4, h); 5.5, B7', '7, 15, 39', 'LBPC, art. 46.º-A; Portaria 91/2017, arts. 2.º, 3.º e 5.º',
           ['Organizações de voluntariado', 'Propor à ANEPC']),
    'B6': ('5.4, d), e) e i)', '—', 'DL 82/2019, art. 8.º, n.º 2; Lei 27/2016, art. 3.º, n.º 1',
           ['Protocolo de reunificação', 'Cláusula de passagem']),
    'B7': ('5.4, b), c), f) e g)', '—', 'LBPC, arts. 37.º e 39.º; SIOPS, arts. 3.º e 4.º (Lei 38/2026); Lei 65/2007, '
           'art. 2.º, n.º 2, al. e); Lei 27/2016, art. 3.º, n.º 6; DL 276/2001, art. 11.º',
           ['Despacho interno', 'Proposta de critérios', 'Partilha da lista', 'Orientação técnica sobre eutanásia']),
    'C1': ('3.1; 4.4, m)', '—', 'DL 116/98, art. 4.º', ['Aceitar o grupo', 'Cláusula de salvaguarda']),
    'C2': ('3.3; 3.4; 3.5', '1, 10, 11, 13, 21', 'DL 276/2001, art. 1.º, n.º 2; LBPC, art. 41.º, al. k)',
           ['Termos da lei']),
    'C3': ('4.4, p); 3.3 (ponto 6.1)', '6, 14, 31; esclarecimento de 14.10.2026 (CAMV)',
           'Lei 27/2016, art. 3.º, n.º 6; DL 276/2001, art. 4.º', ['Responsável clínico']),
    'C4': ('3.2; 3.3 (ponto 5.1.4); 3.6', '23, 24, 25, 26, 27', 'DL 116/98, arts. 2.º, n.ºs 3 e 5, 3.º, n.º 1, e '
           '4.º, n.ºs 3 e 4', ['Médico veterinário do município']),
    'C5': ('4.4, g); 3.3 (pontos 8.8 e 9.5)', '7, 37, 39, 40', 'DL 276/2001, arts. 3.º-A e 4.º; RCM 78/2021',
           ['Bolsa de entidades por protocolo', 'Bolsa de entidades: procedimento']),
    'C6': ('4.4, h)', '36', 'Lei 92/95, arts. 9.º e 10.º', ['Comunicação']),
    'C7': ('4.4, k) e l); 4.5, A3', '35', 'DL 276/2001, art. 19.º, n.º 3; Lei 27/2016, art. 3.º, n.ºs 9 e 10',
           ['Proporcionalidade']),
    'C8': ('4.4, q) e t); 5.4, l)', '28, 29, 41, 43, 44, 45', 'Proposta, pontos 5.2.2, 5.2.4, 10.2, 10.4 e 11.4',
           ['Posto de triagem', 'Relatório e avaliação', 'Retirar números']),
    'C9': ('6.1; memorando do RGAC (fichas G-01, G-02, S-02, E-01)', '—', 'Projeto de RGAC, arts. 91.º e 139.º',
           ['Levar ao RGAC as necessidades A', 'Levar ao RGAC as necessidades B']),
}

NAO_COLIGIDOS = [
    ('2', 'Descrição do documento', 'Descritivo; não é reserva.'),
    ('4.3 e 5.3', 'O que a proposta resolve bem', 'Excluído por instrução: só críticas.'),
    ('3.3 (pontos 2 e 11.1)', 'Designação «ZCAA» em vez de «zona de concentração de acolhimento de animais»',
     'Gralha de designação, sem efeito de fundo.'),
    ('3.3 (ponto 9.1)', 'Registo de associações zoófilas da RCM 78/2021 usado como requisito', 'Absorvido em C5 '
     '(efeitos externos sem norma habilitante); não referido em separado.'),
    ('3.3 (ponto 7.1)', '«designadamente o seu abate» no art. 19.º, n.º 9, do DL 276/2001, hoje limitado pela Lei '
     '27/2016', 'Desatualização da lei, não da proposta.'),
    ('3.3 (ponto 4.5)', 'Mandado judicial com base incompleta (art. 67.º-A; Lei 92/95; CPP; DL 315/2009, art. 30.º)',
     'Coligido só em parte, em A1 (bases omitidas).'),
    ('4.4, i)', 'Suspensão cautelar do acesso do detentor ao SIAC (DL 82/2019, art. 8.º, n.º 5)', 'Medida ao alcance '
     'da DGAV, não reserva à proposta.'),
    ('4.5, A4, A6 e A8', 'Destino antecipado no processo penal; decisão do MVM sem dependência hierárquica em '
     'bem-estar; destino após encerramento', 'Necessidades legislativas; em parte em C9.'),
    ('5.4, b)', 'Critérios técnicos da componente animal dos planos municipais; alojamento conjunto de pessoas e '
     'animais', 'Em parte em B7 (planeamento é municipal).'),
    ('5.4, c)', 'Correspondência entre DSAVR e CCOR', 'Questão interna de organização.'),
    ('5.4, h)', 'Voluntários espontâneos', 'Pormenor; cabe no módulo B quando for reescrito.'),
    ('6.2', 'Regulamento (UE) 2026/1818: quem recebe animais confiscados será «abrigo» a partir de 31.8.2028',
     'Horizonte de 2028; não é reserva imediata.'),
    ('7', 'Estrutura sugerida e próximos passos', 'Propositivo; a síntese só assinala reservas.'),
]


def palavras(t):
    return len(t.split())


def rotulos_recomendacoes():
    """Lê os rótulos das recomendações da análise v2.1, pela ordem, para obter os números R."""
    s = ANALISE.read_text(encoding='utf-8')
    rot = []
    for bloco in ('RECS_A = [', 'RECS_B = ['):
        i = s.index(bloco)
        j = s.index('\n]\n', i)
        rot += re.findall(r"^    \('([^']+)',", s[i:j], flags=re.M)
    return rot


def numero_r(rotulos, inicio):
    for k, r in enumerate(rotulos, start=1):
        if r.startswith(inicio):
            return 'R%d' % k
    raise KeyError(inicio)


def novo_doc(paisagem=False):
    doc = Document()
    z = doc.settings.element.find(qn('w:zoom'))
    if z is not None:
        z.set(qn('w:percent'), '100')
    for sec in doc.sections:
        if paisagem:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
            sec.left_margin = sec.right_margin = Cm(1.8)
        else:
            sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    return doc


def par(doc, txt, bold=False, size=10.5, cor=None, space=6, just=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    if just:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(txt)
    r.bold = bold
    r.font.size = Pt(size)
    if cor is not None:
        r.font.color.rgb = cor
    return p


def titulo(doc, txt, size=12.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(txt)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = AZUL


def tabela(doc, cab, linhas, larguras, size=8.5):
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, c in enumerate(cab):
        cell = t.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(c)
        r.bold = True
        r.font.size = Pt(size + 0.5)
        r.font.color.rgb = AZUL
    trpr = t.rows[0]._tr.get_or_add_trPr()
    trpr.append(trpr.makeelement(qn('w:tblHeader'), {qn('w:val'): 'true'}))
    for linha in linhas:
        cells = t.add_row().cells
        for i, c in enumerate(linha):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(c)
            r.font.size = Pt(size)
            if i == 0:
                r.bold = True
    for i, w in enumerate(larguras):
        t.columns[i].width = Cm(w)
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), larguras):
        gc.set(qn('w:w'), str(int(Cm(w).twips)))
    for row in t.rows:
        for i, w in enumerate(larguras):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ------------------------------------------------------------------ verificações
assert len(FRASES) <= MAX_FRASES, 'frases a mais: %d' % len(FRASES)
for cod, txt in FRASES:
    assert palavras(txt) <= MAX_PALAVRAS, '%s tem %d palavras' % (cod, palavras(txt))
    assert cod in COLIGE, 'sem correspondência: %s' % cod
assert set(COLIGE) == {c for c, _ in FRASES}
ROT = rotulos_recomendacoes()

# ------------------------------------------------------------------ documento 1: as frases
doc = novo_doc()
par(doc, 'DIREÇÃO-GERAL DE ALIMENTAÇÃO E VETERINÁRIA', bold=True, size=9.5, cor=CINZA, space=0)
par(doc, 'Direção de Serviços de Bem-Estar Animal', size=9.5, cor=CINZA, space=12)
p = doc.add_paragraph()
r = p.add_run('Proposta OMV/ANVETEM de procedimento-tipo de intervenção (v3.0, 26.8.2026)')
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = AZUL
par(doc, 'Síntese interna das reservas e incoerências. %d frases, cada uma com no máximo %d palavras. Base: análise '
         'v2.1, de 14.10.2026, e respetivo documento de apoio.' % (len(FRASES), MAX_PALAVRAS),
    size=9.5, cor=CINZA, space=8)
for letra, nome in PARTES:
    titulo(doc, nome)
    for cod, txt in FRASES:
        if cod[0] != letra:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        r0 = p.add_run(cod + '\t')
        r0.bold = True
        r0.font.color.rgb = AZUL
        p.add_run(txt)
doc.save(SAIDA_FRASES)
print('Gerado:', SAIDA_FRASES)

# ------------------------------------------------------------------ documento 2: apoio
doc = novo_doc(paisagem=True)
par(doc, 'DIREÇÃO-GERAL DE ALIMENTAÇÃO E VETERINÁRIA', bold=True, size=9.5, cor=CINZA, space=0)
par(doc, 'Direção de Serviços de Bem-Estar Animal', size=9.5, cor=CINZA, space=10)
p = doc.add_paragraph()
r = p.add_run('Apoio à síntese das reservas à proposta OMV/ANVETEM')
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = AZUL
par(doc, 'Para cada frase da síntese: palavras, secções da análise v2.1 que colige, notas manuscritas ao ficheiro '
         'anotado (numeração do documento de avaliação das notas), bases legais e recomendações da análise '
         'associadas. A numeração R é lida da própria análise v2.1. Documento de trabalho, 14.10.2026.',
    size=9.5, cor=CINZA, space=8, just=True)

titulo(doc, '1. Correspondência frase a frase')
linhas = []
for cod, txt in FRASES:
    secs, notas, bases, recs = COLIGE[cod]
    rs = '; '.join('%s (%s)' % (numero_r(ROT, x), next(r for r in ROT if r.startswith(x))) for x in recs) or '—'
    linhas.append([cod, str(palavras(txt)), secs, notas, bases, rs])
tabela(doc, ['Frase', 'Pal.', 'Secções da análise v2.1', 'Notas', 'Bases legais', 'Recomendações da análise'],
       linhas, [1.2, 1.0, 5.2, 2.8, 8.4, 7.4])

titulo(doc, '2. Cobertura por parte')
for letra, nome in PARTES:
    cods = [c for c, _ in FRASES if c[0] == letra]
    par(doc, '%s: %d frases (%s); média de %d palavras.' % (
        nome, len(cods), ', '.join(cods),
        round(sum(palavras(t) for c, t in FRASES if c[0] == letra) / len(cods))), size=10)

titulo(doc, '3. Temas da análise v2.1 não coligidos nas frases')
tabela(doc, ['Secção', 'Tema', 'Porquê'], [list(x) for x in NAO_COLIGIDOS], [3.2, 13.0, 9.8])

titulo(doc, '4. Notas manuscritas sem frase própria')
par(doc, 'Ficaram absorvidas noutras frases, sem menção autónoma: nota 3 (cenários 1 e 2, que estruturam as partes '
         'A e B); notas 24 e 25 (reduzidas, em C4, ao erro sobre o bem-estar); nota 35 (encargos, em C7). A frase B1 '
         'acrescenta uma observação que não constava da análise v2.1: a distinção entre salvamento e socorro quanto à '
         'presença veterinária.', size=10, just=True)
doc.save(SAIDA_APOIO)
print('Gerado:', SAIDA_APOIO)
