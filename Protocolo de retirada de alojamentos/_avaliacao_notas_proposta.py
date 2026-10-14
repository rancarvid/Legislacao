# -*- coding: utf-8 -*-
"""Avaliação das notas manuscritas à proposta OMV/ANVETEM (ficheiro anotado de 13.10.2026).

Para cada nota: transcrição, validade, se já constava da análise de 14.10.2026 e fundamento.
Gera: Avaliacao_Notas_Proposta_OMV_ANVETEM.docx (na mesma pasta).
Executar a partir da raiz do repositório:
    python3 "Protocolo de retirada de alojamentos/_avaliacao_notas_proposta.py"
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PASTA = Path(__file__).resolve().parent
SAIDA = PASTA / 'Avaliacao_Notas_Proposta_OMV_ANVETEM.docx'

AZUL = RGBColor(0x1F, 0x3B, 0x63)
CINZA = RGBColor(0x55, 0x55, 0x55)

doc = Document()
_zoom = doc.settings.element.find(qn('w:zoom'))
if _zoom is not None:
    _zoom.set(qn('w:percent'), '100')
for s in doc.sections:
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin = s.bottom_margin = Cm(1.8)

st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def p(txt='', bold=False, size=10.5, cor=None, space=6):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(space)
    r = par.add_run(txt)
    r.bold = bold
    r.font.size = Pt(size)
    if cor is not None:
        r.font.color.rgb = cor
    return par


def h1(txt):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(14)
    par.paragraph_format.space_after = Pt(6)
    par.paragraph_format.keep_with_next = True
    r = par.add_run(txt)
    r.bold = True
    r.font.size = Pt(13.5)
    r.font.color.rgb = AZUL


def bullet(txt, rotulo=None):
    par = doc.add_paragraph(style='List Bullet')
    par.paragraph_format.space_after = Pt(3)
    if rotulo:
        r0 = par.add_run(rotulo + ' ')
        r0.bold = True
    par.add_run(txt)


def tabela(cab, linhas, larguras, size=8.5):
    t = doc.add_table(rows=1, cols=len(cab))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cab):
        cell = t.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(c)
        r.bold = True
        r.font.size = Pt(size + 0.5)
        r.font.color.rgb = AZUL
    # Repetir o cabeçalho em cada página
    trpr = t.rows[0]._tr.get_or_add_trPr()
    el = trpr.makeelement(qn('w:tblHeader'), {qn('w:val'): 'true'})
    trpr.append(el)
    for linha in linhas:
        cells = t.add_row().cells
        for i, c in enumerate(linha):
            cells[i].text = ''
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = par.add_run(c)
            r.font.size = Pt(size)
            if i in (3, 4):
                r.bold = True
    # Larguras fixas: grelha, colunas e células (o LibreOffice ignora só as das células)
    t.autofit = False  # o python-docx escreve <w:tblLayout w:type="fixed"/> na posição certa
    for i, w in enumerate(larguras):
        t.columns[i].width = Cm(w)
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), larguras):
        gc.set(qn('w:w'), str(int(Cm(w).twips)))
    for row in t.rows:
        for i, w in enumerate(larguras):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# (n.º, página e ponto, transcrição, válida?, já na análise?, fundamento)
NOTAS = [
    ('1', 'p. 1, título', '«Resgate?» (seta para «retirada e apreensão»)', 'Sim', 'Não',
     '«Resgate» não é termo legal em nenhum dos cenários. No cenário 1 a lei fala em «recolha ou captura» (DL 276/2001, '
     'art. 19.º, n.º 8; Lei 92/95, art. 1.º-A) e em «apreensão» (CPP, art. 178.º; RJCE, art. 49.º). No cenário 2 a '
     'Lei 38/2026 fala em «busca, salvamento e socorro». A própria proposta troca «salvamento» por «resgate» no ponto '
     '9.1, ao citar a al. k) do art. 41.º da LBPC. Usar os termos legais.'),
    ('2', 'p. 1, título', '«Situações de riscos higiossanitários ou ambientais» (seta para «risco para a saúde e o '
     'bem-estar animal»)', 'Sim', 'Parcial',
     'O âmbito da proposta omite fundamentos que a lei prevê: «Existência de riscos higiossanitários» e «proteção do '
     'meio ambiente» (DL 276/2001, art. 3.º-G, n.º 1, als. a) e d)); «ausência de riscos hígio-sanitários '
     'relativamente à conspurcação ambiental» (DL 314/2003, art. 3.º, n.º 1). A análise citou estas normas no quadro '
     'das vias, mas não apontou que faltam no âmbito (ponto 1.1).'),
    ('3', 'p. 1, título', '«cenário 1» (maus-tratos ou risco) e «cenário 2» (socorro)', 'Sim', 'Sim',
     'Coincide com a divisão em âmbito A e âmbito B da análise.'),
    ('4', 'p. 1, síntese', '«314/2003? 315/2009?»', 'Sim', '314: Sim. 315: Não',
     'O DL 314/2003 estava na análise (via V5). O DL 315/2009 não estava, e é uma omissão. Tem regime próprio de '
     'retirada: recolha obrigatória para CRO do animal que ofenda uma pessoa (art. 14.º, n.º 1); destino e eutanásia '
     'decididos pelo MVM (art. 15.º); mandado pedido «ao tribunal cível da respetiva comarca» para aceder e remover '
     '(art. 30.º, n.º 3); apreensão provisória «nos termos do RJCE», com fiel depositário e «retorno ao local de '
     'origem» (art. 39.º, n.ºs 1, 7 e 8); perda e privação do direito de detenção (art. 30.º-A).'),
    ('5', 'p. 1, princípio 1', 'Comando único: «qual?»', 'Sim', 'Sim',
     'O princípio não diz qual autoridade. Na lei depende da via (análise, 4.4, a)).'),
    ('6', 'p. 1, princípio 3', 'Responsabilidade clínica: «qual?»', 'Sim', 'Parcial',
     'A definição diz «inscrito na OMV, designado para a operação», sem dizer por quem. A análise criticou só a '
     'remissão para o art. 4.º do DL 276/2001 (3.3).'),
    ('7', 'p. 1, princípio 6', '«registadas» (1) e «reconhecidas» (2)', 'Sim', 'Sim',
     'Registadas onde: mera comunicação prévia do alojamento (DL 276/2001, art. 3.º-A) ou registo de associações da '
     'RCM 78/2021 (ato programático, junto do ICNF). Reconhecidas por quem: município, para a comissão municipal '
     '(LBPC, art. 41.º, al. k)), ou ANEPC, como OVPC (ver nota 15). Análise, 3.3 e 4.4, g).'),
    ('8', 'p. 1, tipologia A', '«3.º G»', 'Sim', 'Parcial',
     'Para alojamentos registados, a base é o art. 3.º-G (suspensão e encerramento pelo diretor-geral; recolha pelas '
     'câmaras, n.º 6). A proposta só o cita na nota ao 5.1.1. A análise tinha-o na via V4, mas não apontou que '
     'falta na tipologia A.'),
    ('9', 'p. 1, tipologia A', '«+ 314/2003» e «+ situações a rever»', 'Sim', 'Sim',
     'DL 314/2003, art. 3.º, n.ºs 5 e 6 (análise, V5). «Situações a rever» é genérico (leitura da caligrafia a '
     'confirmar).'),
    ('10', 'p. 2, tipologia D', '«Socorro vs. Resgate»', 'Sim', 'Não', 'Ver nota 1.'),
    ('11', 'p. 3, ponto 1.1', 'Numeração 1 a 5 dos fundamentos (maus-tratos; risco para a saúde; bem-estar; saúde '
     'pública; socorro)', 'Sim', 'Sim',
     'São cinco fundamentos com bases legais diferentes. A análise organizou-os por via legal (4.2) e pediu módulos '
     'separados (3.4).'),
    ('12', 'p. 3, ponto 1.5', '«314/2003?» (recolha de errantes «regulada por legislação própria»)', 'Sim', 'Parcial',
     'A legislação própria é o DL 314/2003 (art. 8.º, captura pelas câmaras; art. 9.º, destino decidido pelo MVM), '
     'com a Lei 27/2016 e a Portaria 146/2017. O DL 276/2001, art. 19.º, n.º 10, diz o mesmo para cães e gatos '
     '(análise, 4.4, n)).'),
    ('13', 'p. 3, definição de entidade coordenadora', '«Socorro vs Resgate»', 'Sim', 'Parcial',
     'Terminologia: ver nota 1. Fundo: a Lei 38/2026 não define nenhuma entidade coordenadora; no socorro a '
     'responsabilidade é do COS (SIOPS, art. 9.º, n.º 1). A análise tratou o comando (5.4, a)), não a terminologia.'),
    ('14', 'p. 3, definição de responsável clínico', '«com que meios?»', 'Sim', 'Parcial',
     'A DGAV não tem meios clínicos próprios; o CAMV e a bolsa da OMV são voluntários e o P.A.R.A. está «em '
     'constituição». A análise tratou custos e carga da DGAV (4.4, k) e l)), não os meios clínicos.'),
    ('15', 'p. 3, definição de entidade participante', '«OVPC? organizações voluntárias de PC»', 'Sim', 'Não',
     'Nota importante. A LBPC já tem a figura: «Organizações de voluntariado de proteção civil», com especial dever '
     'de cooperação (art. 46.º-A, n.º 1, al. h)). São reconhecidas por despacho do presidente da ANEPC, nos termos da '
     'Portaria 91/2017 (art. 5.º, n.º 1), que admite entidades com atividades «conexas» à proteção civil (art. 2.º, '
     'n.º 2). É este o canal legal para as associações no cenário 2, não uma credenciação da DGAV. A Portaria fala de '
     '«pessoas e bens» (art. 3.º, n.º 1, al. b)) e não foi atualizada depois da Lei 38/2026.'),
    ('16', 'p. 3, ponto 3.1', '«planeamento prévio»: «por quem?»', 'Sim', 'Parcial',
     'Quem tem o título legal da operação (análise, 4.4, a)). No socorro, o planeamento é municipal (Lei 65/2007, '
     'art. 2.º, n.º 2, al. e)).'),
    ('17', 'p. 4, ponto 3.5', '«mais alojamentos disponíveis para acolher animais de companhia, exemplo: Abrigos AZ ou '
     'Lar de Acolhimento»', 'Sim', 'Parcial',
     'A lei vigente já o permite: depositários idóneos (CPP, art. 178.º, n.º 2); fiel depositário (RJCE, art. 49.º, '
     'n.º 3). O Regulamento (UE) 2026/1818 define abrigo e lar de acolhimento (art. 4.º, als. r) e s)) e o projeto de '
     'RGAC prevê a transmissão a associações e lares de acolhimento (art. 98.º, n.º 1). A análise falou de bolsa de '
     'entidades, sem nomear os lares.'),
    ('18', 'p. 4, ponto 4.2', 'Setores funcionais: «verificar se faz sentido estes elementos em ambos cenários»',
     'Sim', 'Parcial',
     'No cenário 2 os setores são definidos pelo COS (SIOPS, arts. 10.º, n.º 1, al. p), e 14.º). No cenário 1 fazem '
     'sentido.'),
    ('19', 'p. 4, ponto 4.3', '«Verificar se a lei da proteção civil e Lei 38/2026 permitem esta ordem de eventos»',
     'Sim', 'Parcial',
     'Verificado: não preveem. A Lei 38/2026 só põe na comissão municipal a autoridade sanitária veterinária '
     'concelhia e as entidades reconhecidas (LBPC, art. 41.º, als. j) e k)). Não cria dever de comunicação prévia. A '
     'operação de socorro nasce no SIOPS, sob o COS, e não num documento da DGAV.'),
    ('20', 'p. 4, ponto 4.5', 'Mandado prévio: «Incompatível com operações de resgate?»', 'Sim', 'Não',
     'No socorro, exigir mandado antes de entrar é incompatível com a urgência. A entrada para salvar o animal '
     'justifica-se pelo perigo atual (direito de necessidade, CP, art. 34.º) e faz-se numa operação sob o COS. O '
     'mandado continua necessário no cenário 1 para entrar em casa de habitação (CRP, art. 34.º).'),
    ('21', 'p. 4, ponto 4.6 e fundamento', '«276/2009 não se aplica a pecuária» (lido como 276/2001)', 'Sim', 'Não',
     'Confirmado: o DL 276/2001 exclui do seu âmbito «os touros de lide e as espécies de pecuária» (art. 1.º, n.º 2). A '
     'definição de entidade coordenadora remete para a al. x) do art. 2.º desse diploma, que por isso não serve as '
     'tipologias B e C. Reforça a separação por módulos (análise, 3.4).'),
    ('22', 'p. 4, fundamento', '«protege a entidade coordenadora…»: «só a protege neste cenário?»', 'Sim', 'Sim',
     'No cenário 2 a informação pública é do COS (SIOPS, art. 10.º, n.º 1, al. l)). Análise, 5.4, a).'),
    ('23', 'p. 4, ponto 5', '«E o que significa convocação?»', 'Sim', 'Parcial',
     'O termo não tem sentido jurídico definido aqui. O DL 116/98 fala em «convocação extraordinária por motivo '
     'urgente» no programa de contactos com os serviços regionais (art. 4.º, n.º 3). Para as associações é um '
     'convite. Deve dizer-se o que é: pedido escrito, com base legal, ao MVM e ao presidente da câmara.'),
    ('24', 'p. 4, nota dos proponentes, 1.ª parte', '«Confirmar» (DL 116/98 sem competências em bem-estar animal)',
     'Em parte', 'Sim',
     'Certo quanto aos poderes de autoridade sanitária, limitados à sanidade animal e à saúde pública (art. 2.º, '
     'n.º 3). Errado quanto ao dever de colaboração, que nomeia «saúde e bem-estar animal» (art. 3.º, n.º 1). Análise, '
     '3.2.'),
    ('25', 'p. 4, nota dos proponentes, 2.ª parte', '«Confirmar» (fundamento duplo; art. 3.º-G, n.º 6, e art. 19.º)',
     'Em parte', 'Sim',
     'As normas existem, mas atribuem competências às câmaras municipais, não ao MVM. O papel próprio do MVM está no '
     'art. 4.º, n.º 3 (CRO), no art. 19.º, n.º 9, e na al. x) do art. 2.º do DL 276/2001. Análise, 3.2.'),
    ('26', 'p. 5, ponto 5.1.4', '«médico veterinário municipal em funções»: «exclui veterinário contratado»', 'Sim',
     'Não',
     'Muitos municípios têm médico veterinário contratado e não MVM de carreira. Esse veterinário não tem os poderes '
     'de autoridade sanitária, que são conferidos «por inerência de cargo» (DL 116/98, art. 2.º, n.º 3), mas é quem '
     'gere o CRO. A proposta salta-o e vai para o MVM do concelho vizinho. O projeto de RGAC já usa «MVM ou '
     'médico-veterinário do município» (art. 91.º, n.º 8).'),
    ('27', 'p. 5, ponto 5.1.4', '«só a DGAV? Verificar 116/98»', 'Em parte', 'Não',
     'Verificado: a designação do substituto cabe à «autoridade sanitária veterinária nacional» (DL 116/98, art. 2.º, '
     'n.º 5), hoje a DGAV. Dois reparos: a norma prevê a «ausência ou impedimento» do MVM, não a falta de MVM no '
     'município; e o recurso final ao médico veterinário da DSAVR não está no DL 116/98 (vale só para competências '
     'da própria DGAV).'),
    ('28', 'p. 5, ponto 5.2.2', 'Posto de triagem «com condições de biossegurança»: «e bem-estar»', 'Sim', 'Não',
     'Correção simples e certa: o posto de triagem é um alojamento temporário e deve cumprir condições de bem-estar '
     '(DL 276/2001, arts. 7.º e seguintes).'),
    ('29', 'p. 5, pontos 5.2.3 e 5.2.4', '«Enfermagem vet.?»', 'Sim', 'Não',
     'A observação clínica é ato do médico veterinário, mas a equipa deve contar com enfermeiros ou auxiliares '
     'veterinários para contenção, maneio e registo. A proposta só dimensiona médicos veterinários.'),
    ('30', 'p. 5, ponto 5.2.4', 'Observação individual antes da saída: «Incompatível com resgate de emergência?»', 'Sim',
     'Parcial',
     'No socorro a prioridade é tirar o animal do perigo. A triagem faz-se na zona de concentração de acolhimento de '
     'animais, que a lei define para «acomodação, salvamento e triagem» (SIOPS, art. 13.º, al. d)). A análise disse '
     'que o modelo A não serve o B (5.4, d)), sem apontar este ponto.'),
    ('31', 'p. 5, ponto 6.1', '«quando não está disponível?» e «só vets com CAMV?»', 'Sim (1.ª parte)', 'Não',
     '1.ª parte: «indelegável» sem regra de substituição; numa operação longa o responsável clínico tem de ser '
     'substituível. 2.ª parte (leitura a confirmar): a definição só exige inscrição na OMV; a assistência urgente '
     '(6.2) depende de CAMV e a rede P.A.R.A. está em constituição.'),
    ('32', 'p. 6, ponto 6.6', '«Eutanásia?»', 'Sim', 'Não (fora do foco)',
     'A «decisão fundamentada» do DL 142/2006 corresponde, na tabela da própria proposta, a «retirada e destruição» '
     'de animais de produção sem rastreabilidade, isto é, occisão. Não se aplica a animais de companhia. Texto do DL '
     '142/2006 não conferido.'),
    ('33', 'p. 6, ponto 6.7', '«destino de animais de companhia?»', 'Sim', 'Parcial',
     'O ponto só remete para o Regulamento (CE) n.º 1069/2009. Para animais de companhia falta dizer o destino do '
     'cadáver (o regulamento trata-os como subprodutos e admite exceções para animais de companhia; redação a '
     'conferir) e, havendo processo penal, a necropsia antes da eliminação (CPP, art. 159.º-A). A análise tratou a '
     'prova na eutanásia (4.4, f)).'),
    ('34', 'p. 6, ponto 7.1', '«não pode ser este o único destino»', 'Sim', 'Sim',
     'Além do CRO: depositário idóneo (CPP, art. 178.º, n.º 2), detentor como fiel depositário (RJCE, art. 49.º, n.º 3), '
     'outro destino escolhido pelo detentor (DL 314/2003, art. 3.º, n.º 5). Análise, 4.4, b) e d).'),
    ('35', 'p. 6, ponto 7.7', 'Sublinhado em «à entidade responsável»', 'Sim', 'Sim',
     'A entidade não é identificada. Análise, 4.4, k).'),
    ('36', 'p. 7, ponto 8.5', '«Outra vez?»', 'Sim', 'Parcial',
     'Repete o princípio 5 e os pontos 8.1 a 8.4. A análise criticou a linguagem (4.4, h)), não a repetição.'),
    ('37', 'p. 7, ponto 8.8', '«quem determina?»', 'Sim', 'Parcial',
     'O 8.8 fala em «exclusão» e remete para o 9.5, que fala em «suspensão» decidida pela DGAV. Dois termos para o '
     'mesmo efeito, e nada se diz quando a coordenadora não é a DGAV. A análise tratou a falta de base legal (4.4, g)).'),
    ('38', 'p. 7, ponto 9.1, transporte', '«verificar que critérios se aplicam aos animais de companhia»', 'Sim', 'Não',
     'A autorização de transportador vem do DL 265/2007, que executa o Regulamento (CE) n.º 1/2005. Esse regulamento '
     'não se aplica ao transporte que não seja feito no âmbito de uma atividade económica (art. 1.º, n.º 5; redação a '
     'conferir), como o de animais retirados por autoridades ou associações. Para animais de companhia, o critério '
     'vigente é o do DL 276/2001, art. 10.º, n.º 1 (veículos e contentores apropriados à espécie e ao número).'),
    ('39', 'p. 7, ponto 9.2, seguro', '«as entidades ficam com encargo adicional?»', 'Sim', 'Parcial',
     'Sim. O seguro dos voluntários já é obrigatório para as OVPC (Portaria 91/2017, art. 5.º, n.º 4, al. l)); o seguro '
     'de responsabilidade civil para resgate, maneio e transporte é custo novo, sem fonte de financiamento. A análise '
     'disse que os requisitos podem afastar capacidade (4.4, g)).'),
    ('40', 'p. 7, ponto 9.3', '«atualizada anualmente! qual procedimento?»', 'Sim', 'Parcial',
     'Não há procedimento: quem pede, quem verifica, em que prazo, com que recurso. A análise tratou a base legal da '
     'lista (4.4, g)).'),
    ('41', 'p. 8, ponto 10.2', 'Relatório ao MVM e à câmara: «Porquê?»', 'Em parte', 'Não',
     'Há razão para os enviar (CRO, art. 3.º-G, n.º 6, encargos), mas não sempre nem por inteiro. Havendo processo '
     'penal, a distribuição tem de respeitar o segredo de justiça e a proteção de dados. Deve ser por necessidade.'),
    ('42', 'p. 8, ponto 10.2', 'Comissão municipal de proteção civil nas operações de socorro: «Verificar»', 'Sim', 'Não',
     'Verificado: a Lei 38/2026 não prevê relatórios à comissão municipal. No socorro, o relato segue a cadeia do '
     'SIOPS (COS e comando de emergência e proteção civil, art. 10.º, n.º 1, al. g)).'),
    ('43', 'p. 8, ponto 10.4', 'Avaliação anual só com OMV e ANVETEM: «só?»', 'Sim', 'Parcial',
     'Faltam a ANMP, as forças de segurança, a ANEPC e as associações. A análise propôs protocolo com ANMP e forças de '
     'segurança (3.1), não a avaliação conjunta.'),
    ('44', 'p. 8, ponto 11.4', '«Atualmente…»: «precisa atualização periódica pelo que o n.º 106 é irrelevante neste '
     'documento»', 'Sim', 'Não',
     'Um texto normativo não deve fixar números que mudam. A análise usou o número como dado (5.3), sem esta crítica.'),
    ('45', 'p. 8, ponto 11.4', 'Comissão Técnica de Emergência, Catástrofe e Socorro Animal: «que comissão é esta?»',
     'Sim', 'Não',
     'É um órgão interno da OMV, sem estatuto legal no procedimento. Pedir mandato, composição e papel.'),
]

# ───────────────────────────── Capa ─────────────────────────────
p('DIREÇÃO-GERAL DE ALIMENTAÇÃO E VETERINÁRIA', bold=True, size=9.5, cor=CINZA, space=0)
p('Direção de Serviços de Bem-Estar Animal', size=9.5, cor=CINZA, space=10)
par = doc.add_paragraph()
r = par.add_run('AVALIAÇÃO DAS NOTAS À PROPOSTA OMV / ANVETEM')
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = AZUL
p('Ficheiro anotado «Proposta de Procedimento Integrado_DGAV_OMV_ANVETE_261013_001303.pdf» (notas manuscritas sobre '
  'a v3.0, de 26.8.2026). Documento de trabalho, 14.10.2026. Complementa a «Análise da proposta OMV/ANVETEM» da mesma '
  'data.', size=9.5, cor=CINZA, space=10)

h1('1. Síntese')
validas = sum(1 for n in NOTAS if n[3].startswith('Sim'))
parte = sum(1 for n in NOTAS if n[3].startswith('Em parte'))
ja = sum(1 for n in NOTAS if n[4] == 'Sim')
parc = sum(1 for n in NOTAS if n[4].startswith('Parcial') or n[4].startswith('314'))
novas = len(NOTAS) - ja - parc
p('Foram transcritas %d notas. %d são válidas e %d são válidas em parte; nenhuma é infundada. Quanto à análise de '
  '14.10.2026: %d já lá estavam, %d estavam em parte e %d são novas.' % (len(NOTAS), validas, parte, ja, parc, novas))
p('As notas novas mais importantes, por ordem de peso:')
bullet('o DL 315/2009 tem regime próprio de recolha, apreensão, mandado e destino (cães que agridem, cães perigosos), '
       'e a análise não o tinha (nota 4);', rotulo='DL 315/2009:')
bullet('as organizações de voluntariado de proteção civil da LBPC e da Portaria 91/2017 são o canal legal para as '
       'associações no socorro, em vez de uma credenciação da DGAV (nota 15);', rotulo='OVPC:')
bullet('«resgate» não é termo legal; a lei diz «recolha», «captura», «apreensão» e «busca, salvamento e socorro» '
       '(notas 1, 10 e 13);', rotulo='Terminologia:')
bullet('o DL 276/2001 exclui as espécies de pecuária, pelo que a remissão para a sua al. x) não serve as tipologias B e '
       'C (nota 21);', rotulo='Âmbito do DL 276/2001:')
bullet('mandado prévio e observação individual antes da saída não cabem no socorro (notas 20 e 30); a comunicação '
       'prévia e o relatório à comissão municipal não têm base na Lei 38/2026 (notas 19 e 42);',
       rotulo='Socorro:')
bullet('o médico veterinário contratado pelo município fica de fora (nota 26); o transporte de animais de companhia '
       'retirados não segue o regime do transportador autorizado (nota 38).', rotulo='Outras:')

h1('2. O que as notas corrigem na análise de 14.10.2026')
bullet('Omissão do DL 315/2009 no quadro das vias legais (4.2). Corrigido na análise: nova via V7.')
bullet('Na via V5 (DL 314/2003), o destino dos animais recolhidos é decidido pelo MVM, após exame clínico, com '
       'permanência mínima de oito dias (DL 314/2003, art. 9.º, n.º 1). A análise dizia «regime do CRO». Corrigido.')
bullet('O «retorno ao local de origem» que a análise criticou no projeto de RGAC (art. 139.º, n.º 8) já existe na lei '
       'vigente, no DL 315/2009, art. 39.º, n.º 8. E esse artigo já remete a apreensão para o RJCE (n.º 1), o que o '
       'RGAC deixa cair. A ficha S-02 do memorando do RGAC foi completada com isto.')
p('As restantes notas novas (OVPC, terminologia, âmbito do DL 276/2001, socorro, veterinário contratado, transporte e '
  'as de menor peso) não foram integradas no texto da análise. Podem sê-lo numa versão 2, se for útil.')

h1('3. Quadro das notas')
p('Colunas: «Válida?» (Sim, Em parte, Não); «Já na análise?» (Sim, Parcial, Não), com a secção da análise de '
  '14.10.2026 indicada no fundamento. As leituras da caligrafia que ficaram incertas estão assinaladas.', size=9.5,
  cor=CINZA)
tabela(['N.º', 'Página e ponto', 'Nota (transcrição)', 'Válida?', 'Já na análise?', 'Verificação e fundamento'],
       NOTAS, [0.9, 2.6, 5.2, 1.6, 1.9, 13.9])

h1('4. Fontes conferidas para esta avaliação')
bullet('Lei de Bases da Proteção Civil, na republicação da Lei 80/2015, art. 46.º-A (DR, 1.ª série, n.º 149, de '
       '3.8.2015).')
bullet('Portaria 91/2017, de 2 de março (organizações de voluntariado de proteção civil), arts. 1.º a 5.º (DR, 1.ª '
       'série, n.º 44).')
bullet('DL 315/2009, versão consolidada, arts. 14.º, 15.º, 30.º, 30.º-A e 39.º (pasta «Legislação vigente»).')
bullet('DL 314/2003, versão consolidada, arts. 3.º, 8.º e 9.º (pasta «Legislação vigente»).')
bullet('DL 276/2001, versão consolidada, arts. 1.º, 3.º-G e 10.º; DL 116/98, arts. 2.º a 4.º; Lei 38/2026; SIOPS '
       '(fontes já usadas na análise).')
bullet('Não conferidos nesta sessão: Regulamento (CE) n.º 1/2005, art. 1.º, n.º 5, e Regulamento (CE) n.º 1069/2009 '
       '(EUR-Lex inacessível a partir do contentor); DL 142/2006, art. 23.º, n.º 7.')

doc.save(SAIDA)
print('Gerado:', SAIDA)
